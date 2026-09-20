"""Diagnostic script: preview every layout in Company_Template.pptx, export to PDF/PNG, and analyze decoration scores."""

from __future__ import annotations

import os
import pathlib
import re
import sys
from pathlib import Path

# Ensure root directory is on sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import pptx
from pptx.util import Inches, Pt
import pymupdf


def sanitize_filename(name: str) -> str:
    """Remove invalid filesystem characters from string."""
    return re.sub(r'[\\/*?:"<>|]', '_', name).strip()


def convert_pptx_to_pdf(pptx_path: Path, pdf_path: Path) -> bool:
    """Convert PPTX to PDF using MS PowerPoint COM (Windows) or soffice fallback."""
    pdf_path.parent.mkdir(parents=True, exist_ok=True)
    abs_pptx = str(pptx_path.resolve())
    abs_pdf = str(pdf_path.resolve())

    try:
        import win32com.client
        ppt_app = win32com.client.Dispatch("PowerPoint.Application")
        pres = ppt_app.Presentations.Open(abs_pptx, True, False, False)
        pres.SaveAs(abs_pdf, 32)  # 32 = ppSaveAsPDF
        pres.Close()
        ppt_app.Quit()
        return True
    except Exception as exc:
        print(f"[Warning] PowerPoint COM conversion failed: {exc}")

    import shutil
    import subprocess
    soffice_cmd = shutil.which("soffice") or shutil.which("libreoffice")
    if soffice_cmd:
        try:
            subprocess.run([soffice_cmd, "--headless", "--convert-to", "pdf", abs_pptx, "--outdir", str(pdf_path.parent)], check=True)
            return True
        except Exception as exc:
            print(f"[Warning] soffice conversion failed: {exc}")

    return False


def main() -> None:
    """Preview all layouts in Company_Template.pptx and analyze layout design features."""
    tmpl_path = Path("data/sample_templates/Company_Template.pptx")
    out_dir = Path("data/outputs/preview")
    layouts_png_dir = out_dir / "layouts"
    layouts_png_dir.mkdir(parents=True, exist_ok=True)

    if not tmpl_path.exists():
        print(f"Error: Missing template '{tmpl_path}'")
        sys.exit(1)

    prs = pptx.Presentation(tmpl_path)
    preview_prs = pptx.Presentation(tmpl_path)

    # Clean out existing slides in preview_prs
    for sldId in list(preview_prs.slides._sldIdLst):
        rId = sldId.rId
        preview_prs.part.drop_rel(rId)
        preview_prs.slides._sldIdLst.remove(sldId)

    layout_info_list = []

    # Iterate layouts and populate dummy content
    for idx, layout in enumerate(prs.slide_layouts):
        slide = preview_prs.slides.add_slide(preview_prs.slide_layouts[idx])

        # Fill placeholders with dummy text
        ph_summary = []
        for ph in slide.placeholders:
            ph_type = str(ph.placeholder_format.type).split(".")[-1].lower()
            ph_summary.append(f"{ph.name} ({ph_type})")
            if shape_has_text(ph):
                try:
                    ph.text_frame.text = f"[{ph.name} - Layout {idx}: {layout.name}]"
                except Exception:
                    pass

        # Count non-placeholder shapes on layout & slide master
        layout_non_ph = [s for s in layout.shapes if not s.is_placeholder]
        master_non_ph = [s for s in layout.slide_master.shapes if not s.is_placeholder]
        total_non_ph = len(layout_non_ph) + len(master_non_ph)

        # Estimate decoration score
        decoration_score = len(layout_non_ph) * 10 + len(master_non_ph) * 2
        bg_type = "Design Elements / Master Graphic" if total_non_ph > 0 else "Plain White Background"

        layout_info_list.append({
            "index": idx,
            "name": layout.name,
            "placeholders": len(slide.placeholders),
            "ph_desc": ", ".join(ph_summary[:3]) + ("..." if len(ph_summary) > 3 else ""),
            "non_ph_shapes": total_non_ph,
            "bg_type": bg_type,
            "decoration_score": decoration_score,
        })

    # Save preview deck
    preview_pptx_path = out_dir / "all_layouts_preview.pptx"
    preview_prs.save(preview_pptx_path)
    print(f"Saved layout preview PPTX to {preview_pptx_path}")

    # Convert preview PPTX to PDF and PNGs
    preview_pdf_path = out_dir / "all_layouts_preview.pdf"
    if convert_pptx_to_pdf(preview_pptx_path, preview_pdf_path):
        print(f"Converted PPTX to PDF: {preview_pdf_path}")
        doc_pdf = pymupdf.open(preview_pdf_path)
        for page_idx, page in enumerate(doc_pdf):
            pix = page.get_pixmap(dpi=150)
            safe_name = sanitize_filename(layout_info_list[page_idx]['name']).replace(" ", "_")
            png_file = layouts_png_dir / f"layout_{page_idx}_{safe_name}.png"
            pix.save(str(png_file))
        doc_pdf.close()
        print(f"Rendered {len(layout_info_list)} layout PNG previews into {layouts_png_dir}")

    # Print Diagnostic Table
    print("\n" + "=" * 90)
    print(f"{'Idx':<4} | {'Layout Name':<28} | {'PHs':<4} | {'Non-PH Shapes':<14} | {'Dec. Score':<10} | {'Design Status':<22}")
    print("=" * 90)

    for item in layout_info_list:
        print(
            f"{item['index']:<4} | "
            f"{item['name'][:28]:<28} | "
            f"{item['placeholders']:<4} | "
            f"{item['non_ph_shapes']:<14} | "
            f"{item['decoration_score']:<10} | "
            f"{item['bg_type']:<22}"
        )

    print("=" * 90)


def shape_has_text(shape: Any) -> bool:
    try:
        return shape.has_text_frame
    except Exception:
        return False


if __name__ == "__main__":
    main()
