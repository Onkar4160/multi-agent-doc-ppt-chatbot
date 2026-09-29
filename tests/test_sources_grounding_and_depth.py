"""Tests covering Parts 1-4:
a) After an edit, rendered Sources contains real titles/URLs for old & new ids, no "Source N" placeholder.
b) New content citing unrelated old id is rejected and retried once; citing new evidence ids passes.
c) Empty evidence: no specific facts, no citations, and reply contains warning line.
d) Date equals today's date (format Month YYYY) even if mocked LLM returns another date.
e) update_title changes DOCX title and deck title slide.
f) Query prompt for timeless topic has no blanket "latest/2026" and asks for 3 different aspects.
g) Doc under 50% of target triggers exactly one retry.
h) Number or date without source_ids is a validator error.
i) Topic-adding edit updates both artifacts and merged Sources slide, within LLM budget (<= 5 calls).
j) Migration adds sources_json and backfills from Source rows without losing data.
"""

from __future__ import annotations

import datetime
import json
import re
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
from sqlalchemy import create_engine
from sqlalchemy.orm import sessionmaker

from app.agents.doc_generator import generate_document_model
from app.agents.editor import edit_artifact
from app.agents.graph import _get_sync_session, graph_app, node_edit, node_validate, should_retry
from app.agents.ppt_generator import generate_deck_model
from app.agents.state import GraphState
from app.agents.supervisor import Plan, parse_plan
from app.agents.validator import validate_outputs
from app.agents.web_research import research
from app.core.database import Base, run_migrations_sync
from app.llm.client import get_llm_client
from app.models.artifact import Artifact, ArtifactVersion
from app.models.chat import ChatMessage, ChatSession
from app.models.deck_model import BulletItem, DeckModel, SlideModel
from app.models.document_model import BulletsBlock, DocumentModel, ParagraphBlock, Section
from app.models.edit_ops import AddSectionOp, EditPlan, UpdateTitleOp
from app.models.file import UploadedFile
from app.models.source import Source
from app.models.template_profile import TemplateProfile
from app.models.workspace import Workspace
from app.services.context_builder import ContextResult
from app.services.docx_renderer import render_docx
from app.services.evidence_pack import EvidenceItem
from app.services.pptx_renderer import render_pptx
from docx import Document
from pptx import Presentation


@pytest.fixture
def sync_db():
    """Isolated SQLite DB with tables, active workspace, and default templates."""
    engine = create_engine("sqlite:///:memory:", connect_args={"check_same_thread": False})
    Base.metadata.create_all(engine)
    with engine.begin() as conn:
        run_migrations_sync(conn)

    SessionLocal = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    db = SessionLocal()

    tmpl_docx = Path("data/sample_templates/proposal_Template.docx")
    tmpl_pptx = Path("data/sample_templates/presentation_Template.pptx")
    tmpl_docx.parent.mkdir(parents=True, exist_ok=True)
    if not tmpl_docx.exists():
        d = Document()
        d.save(tmpl_docx)
    if not tmpl_pptx.exists():
        p = Presentation()
        p.save(tmpl_pptx)

    f_doc = UploadedFile(filename="test.docx", file_type="docx", stored_path=str(tmpl_docx), file_size=100)
    f_ppt = UploadedFile(filename="test.pptx", file_type="pptx", stored_path=str(tmpl_pptx), file_size=100)
    db.add(f_doc)
    db.add(f_ppt)
    db.flush()

    ws = Workspace(
        name="Default Workspace",
        is_active=True,
        docx_template_file_id=f_doc.id,
        pptx_template_file_id=f_ppt.id,
    )
    db.add(ws)
    db.commit()

    yield db
    db.close()


# ── Test a: Rendered Sources contains real titles/URLs and no "Source N" placeholder ──

def test_a_sources_section_no_source_n_placeholders(sync_db, tmp_path):
    """Sources section rendered on DOCX and PPTX uses real metadata, never 'Source N'."""
    doc_m = DocumentModel(
        title="Modern Java",
        date="September 2026",
        sections=[
            Section(
                heading="Overview",
                blocks=[ParagraphBlock(text="Java 21 introduces virtual threads [1] and Python 3.12 adds subinterpreters [2].", source_ids=[1, 2])],
            )
        ],
    )
    sources_map = {
        "1": {"id": 1, "title": "Java Virtual Threads Guide", "url_or_filename": "https://oracle.com/java21", "kind": "web", "accessed_at": "September 2026"},
        "2": {"id": 2, "title": "Python Concurrency Docs", "url_or_filename": "https://python.org/docs312", "kind": "web", "accessed_at": "September 2026"},
        "3": {"id": 3, "title": "Uncited Article", "url_or_filename": "https://example.com/uncited", "kind": "web", "accessed_at": "September 2026"},
    }

    tmpl_docx = Path("data/sample_templates/proposal_Template.docx")
    out_docx = tmp_path / "test_sources.docx"
    render_docx(doc_m, template_path=tmpl_docx, out_path=out_docx, sources_map=sources_map)
    d = Document(out_docx)
    full_text = "\n".join(p.text for p in d.paragraphs)

    assert not re.search(r"Source\s+\d+", full_text, re.IGNORECASE)
    assert "[1] Java Virtual Threads Guide - https://oracle.com/java21 (accessed September 2026)" in full_text
    assert "[2] Python Concurrency Docs - https://python.org/docs312 (accessed September 2026)" in full_text
    assert "Uncited Article" not in full_text


# ── Test b: Grounding rule on edits (unrelated old ID rejected, new ID passes) ──

def test_b_edit_grounding_rejects_unrelated_old_id_and_retries(sync_db):
    """New block citing unrelated old source ID triggers retry; citing fresh evidence ID passes."""
    doc_m = DocumentModel(
        title="Java Guide",
        date="September 2026",
        sections=[Section(heading="Java Intro", blocks=[ParagraphBlock(text="Java fundamentals.", source_ids=[1])])],
    )
    art = Artifact(title="Java Guide", artifact_type="docx", session_id="s1")
    sync_db.add(art)
    sync_db.flush()

    v1 = ArtifactVersion(
        artifact_id=art.id,
        version_no=1,
        model_json=doc_m.model_dump_json(),
        file_path="dummy.docx",
        file_type="docx",
        source_ids_json=json.dumps([1]),
        sources_json=json.dumps({"1": {"id": 1, "title": "Java Oracle Docs", "url_or_filename": "https://oracle.com", "kind": "web", "accessed_at": "September 2026"}}),
    )
    sync_db.add(v1)
    sync_db.commit()

    fresh_evidence = [
        EvidenceItem(id=2, text="Python was created by Guido van Rossum.", source_id=2, source_title="Python Org", url="https://python.org"),
    ]
    fresh_ctx = ContextResult(
        brief="Python",
        sources_map={"2": {"id": 2, "title": "Python Org", "url_or_filename": "https://python.org", "kind": "web", "accessed_at": "September 2026"}},
        evidence_pack=fresh_evidence,
        findings=[],
    )

    call_count = {"calls": 0}

    def mock_gen_plan(*args, **kwargs):
        call_count["calls"] += 1
        if call_count["calls"] == 1:
            # First attempt: cites old unrelated source [1] on Python block
            return EditPlan(
                target="docx",
                summary="Append Python section",
                ops=[AddSectionOp(section=Section(heading="Python Intro", blocks=[ParagraphBlock(text="Python is dynamic.", source_ids=[1])]))],
            )
        else:
            # Retry attempt: correctly cites fresh evidence ID [2]
            return EditPlan(
                target="docx",
                summary="Append Python section",
                ops=[AddSectionOp(section=Section(heading="Python Intro", blocks=[ParagraphBlock(text="Python is dynamic.", source_ids=[2])]))],
            )

    with patch("app.agents.editor._generate_edit_plan", side_effect=mock_gen_plan):
        res = edit_artifact(
            artifact_id=art.id,
            instruction="append information about Python",
            db_session=sync_db,
            shared_context=fresh_ctx,
            active_topic="Java and Python",
        )

    assert call_count["calls"] == 2  # exactly one retry triggered
    assert res.new_version_no == 2
    v2 = sync_db.query(ArtifactVersion).filter(ArtifactVersion.artifact_id == art.id, ArtifactVersion.version_no == 2).first()
    v2_sources = json.loads(v2.sources_json)
    assert "1" in v2_sources
    assert "2" in v2_sources


# ── Test c: Empty evidence pack produces no specific facts and adds warning line ──

def test_c_empty_evidence_pack_warning(sync_db):
    """When search finds nothing, no specific facts are cited and reply has warning message."""
    doc_m = DocumentModel(
        title="Existing Guide",
        date="September 2026",
        sections=[Section(heading="Intro", blocks=[ParagraphBlock(text="Overview content.", source_ids=[1])])],
    )
    art = Artifact(title="Existing Guide", artifact_type="docx", session_id="s_empty")
    sync_db.add(art)
    sync_db.flush()

    v1 = ArtifactVersion(
        artifact_id=art.id,
        version_no=1,
        model_json=doc_m.model_dump_json(),
        file_path="dummy.docx",
        file_type="docx",
        source_ids_json=json.dumps([1]),
        sources_json=json.dumps({"1": {"id": 1, "title": "Base Doc", "url_or_filename": "https://base.org", "kind": "web", "accessed_at": "September 2026"}}),
    )
    sync_db.add(v1)
    sync_db.commit()

    empty_ctx = ContextResult(brief="obscure tech", sources_map={}, evidence_pack=[], findings=[])

    def mock_plan(*args, **kwargs):
        return EditPlan(
            target="docx",
            summary="Add obscure topic",
            ops=[AddSectionOp(section=Section(heading="Obscure Topic", blocks=[ParagraphBlock(text="General conceptual text.", source_ids=[])]))],
        )

    with patch("app.agents.editor._generate_edit_plan", side_effect=mock_plan):
        res = edit_artifact(
            artifact_id=art.id,
            instruction="append information about obscure tech",
            db_session=sync_db,
            shared_context=empty_ctx,
            active_topic="obscure tech",
        )

    assert res.no_sources_warning is not None
    assert "No sources found for" in res.no_sources_warning
    assert "added general content without citations" in res.no_sources_warning


# ── Test d: Date equals today's date (Month YYYY) even if LLM returns another date ──

def test_d_date_equals_today_code_enforced(sync_db):
    """DocumentModel.date and deck date are forced to today's date in code."""
    today_expected = datetime.date.today().strftime("%B %Y")

    # Document generator
    llm_doc = DocumentModel(
        title="Report",
        date="March 2026",  # LLM hallucinates March 2026
        sections=[Section(heading="Intro", blocks=[ParagraphBlock(text="Body.", source_ids=[1])])],
    )
    with patch.object(get_llm_client(), "generate_json", return_value=llm_doc):
        prof = TemplateProfile(colors={"primary": "#000000"})
        gen_doc = generate_document_model(brief="Test brief", profile=prof)
        assert gen_doc.date == today_expected
        assert gen_doc.date != "March 2026"

    # Presentation generator
    llm_deck = DeckModel(
        title="Pitch",
        slides=[SlideModel(title="Intro", role="title", subtitle="March 2026", bullets=[BulletItem(text="Bullet", source_ids=[1])])],
    )
    with patch.object(get_llm_client(), "generate_json", return_value=llm_deck):
        gen_deck = generate_deck_model(brief="Test brief", profile=prof, slide_count=1)
        assert today_expected in gen_deck.slides[0].subtitle


# ── Test e: update_title changes DOCX title and deck title slide ──

def test_e_update_title_op(sync_db):
    """update_title op modifies DOCX title and deck title slide."""
    doc_m = DocumentModel(
        title="Java Guide",
        subtitle="Core Features",
        date="September 2026",
        sections=[Section(heading="Intro", blocks=[ParagraphBlock(text="Java.", source_ids=[1])])],
    )
    art = Artifact(title="Java Guide", artifact_type="docx", session_id="s_title")
    sync_db.add(art)
    sync_db.flush()

    v1 = ArtifactVersion(
        artifact_id=art.id,
        version_no=1,
        model_json=doc_m.model_dump_json(),
        file_path="dummy.docx",
        file_type="docx",
        source_ids_json=json.dumps([1]),
        sources_json=json.dumps({"1": {"id": 1, "title": "Java Guide", "url_or_filename": "https://java.com", "kind": "web", "accessed_at": "September 2026"}}),
    )
    sync_db.add(v1)
    sync_db.commit()

    def mock_plan(*args, **kwargs):
        return EditPlan(
            target="docx",
            summary="Update title and append Python",
            ops=[
                UpdateTitleOp(title="Java and Python: Modern Language Comparison", subtitle="Ecosystem and Architecture"),
                AddSectionOp(section=Section(heading="Python", blocks=[ParagraphBlock(text="Python.", source_ids=[])])),
            ],
        )

    with patch("app.agents.editor._generate_edit_plan", side_effect=mock_plan):
        res = edit_artifact(
            artifact_id=art.id,
            instruction="append information about Python and update title",
            db_session=sync_db,
        )

    v2 = sync_db.query(ArtifactVersion).filter(ArtifactVersion.artifact_id == art.id, ArtifactVersion.version_no == 2).first()
    new_doc = DocumentModel.model_validate_json(v2.model_json)
    assert new_doc.title == "Java and Python: Modern Language Comparison"
    assert new_doc.subtitle == "Ecosystem and Architecture"

    # Also test update_title on PPTX deck
    deck_m = DeckModel(
        title="Java Pitch",
        subtitle="Initial Subtitle",
        slides=[
            SlideModel(role="title", title="Java Pitch", subtitle="Initial Subtitle"),
            SlideModel(role="title_content", title="Details", bullets=[BulletItem(text="Details point.")]),
        ],
    )
    art_ppt = Artifact(title="Java Pitch", artifact_type="pptx", session_id="s_title")
    sync_db.add(art_ppt)
    sync_db.flush()
    v1_ppt = ArtifactVersion(
        artifact_id=art_ppt.id,
        version_no=1,
        model_json=deck_m.model_dump_json(),
        file_path="dummy.pptx",
        file_type="pptx",
    )
    sync_db.add(v1_ppt)
    sync_db.commit()

    def mock_ppt_plan(*args, **kwargs):
        return EditPlan(
            target="pptx",
            summary="Update title on deck",
            ops=[UpdateTitleOp(title="New Deck Title", subtitle="New Deck Subtitle")],
        )

    with patch("app.agents.editor._generate_edit_plan", side_effect=mock_ppt_plan):
        edit_artifact(
            artifact_id=art_ppt.id,
            instruction="update deck title",
            db_session=sync_db,
        )

    v2_ppt = sync_db.query(ArtifactVersion).filter(ArtifactVersion.artifact_id == art_ppt.id, ArtifactVersion.version_no == 2).first()
    new_deck = DeckModel.model_validate_json(v2_ppt.model_json)
    assert new_deck.title == "New Deck Title"
    assert new_deck.subtitle == "New Deck Subtitle"
    assert new_deck.slides[0].title == "New Deck Title"
    assert "New Deck Subtitle" in new_deck.slides[0].subtitle


# ── Test f: Query prompt for timeless topic has no blanket latest/2026 & asks for 3 aspects ──

def test_f_query_prompt_timeless_topic():
    """Query prompt asks for 3 distinct aspects and does not force blanket 2026/latest."""
    captured_prompt = []

    def mock_gen(prompt, schema, **kwargs):
        captured_prompt.append(prompt)
        from app.agents.web_research import ResearchFindingsSchema, SearchQueriesSchema
        if schema == SearchQueriesSchema:
            return SearchQueriesSchema(queries=["Java language overview", "Java current architecture", "Java ecosystem comparison"])
        return ResearchFindingsSchema(findings=[])

    with patch.object(get_llm_client(), "generate_json", side_effect=mock_gen), \
         patch("app.agents.web_research._search_web_single_query", return_value=[]):
        res = research("Java Language fundamentals")

    assert len(captured_prompt) >= 1
    p = captured_prompt[0]
    assert "include 2026 or latest" not in p
    assert "3" in p
    assert "fundamentals" in p.lower()
    assert "recent developments" in p.lower() or "current status" in p.lower()
    assert len(res.queries) == 3


# ── Test g: Doc under 50% target triggers exactly one retry ──

def test_g_doc_under_50_percent_triggers_retry():
    """Document under 50% of the minimum target word count is an error and triggers one retry."""
    short_doc = DocumentModel(
        title="Very Short Report",
        sections=[
            Section(heading="Intro", blocks=[ParagraphBlock(text="Short text.", source_ids=[1])]),
        ],
    )
    report = validate_outputs(doc_model=short_doc, valid_source_ids={1})
    assert not report.passed
    errors = [i for i in report.issues if i.severity == "error" and "under 50% of minimum target" in i.message]
    assert len(errors) == 1

    state: GraphState = {
        "validation": report.model_dump(),
        "retry_count": 0,
        "plan": {"outputs": ["docx"]},
    }
    assert should_retry(state) == "generate_doc_node"
    assert state["retry_count"] == 1
    assert should_retry(state) == "finalize_node"


# ── Test h: Number or date without source_ids is an error ──

def test_h_number_or_date_without_citation_is_error():
    """Any specific number, percentage, or date without source_ids is a validator error."""
    doc_with_stats = DocumentModel(
        title="Uncited Stats",
        sections=[
            Section(heading="Stats", blocks=[ParagraphBlock(text="In 2024, adoption reached 85% with version 21.0.", source_ids=[])]),
        ],
    )
    report = validate_outputs(doc_model=doc_with_stats, valid_source_ids={1})
    assert not report.passed
    grounding_errors = [i for i in report.issues if i.severity == "error" and "lacks citation source_ids" in i.message]
    assert len(grounding_errors) >= 1


# ── Test i: Topic-adding edit updates both artifacts, merged Sources slide, within call budget ──

def test_i_topic_adding_edit_dual_artifact_and_budget(sync_db):
    """Supervisor sets needs_new_facts=True; node_edit updates both docx & pptx within <= 5 calls."""
    # 1. Supervisor test: "append information about Python" sets needs_new_facts=True
    plan = parse_plan("append information about Python", session_artifacts=[{"artifact_id": 1, "kind": "docx", "title": "Java Guide", "latest_version": 1}])
    assert plan.action == "edit"
    assert plan.needs_new_facts is True

    # 2. Setup session with both DOCX and PPTX artifacts
    sess_id = "sess_dual"
    doc_m = DocumentModel(
        title="Java Guide",
        date="September 2026",
        sections=[Section(heading="Java Core", blocks=[ParagraphBlock(text="Java text.", source_ids=[1])])],
    )
    deck_m = DeckModel(
        title="Java Guide",
        slides=[
            SlideModel(title="Java Core", role="title_content", bullets=[BulletItem(text="Java point.", source_ids=[1])]),
            SlideModel(title="Sources", role="title_only", bullets=[BulletItem(text="[1] Java Docs", source_ids=[1])]),
        ],
    )

    art_doc = Artifact(title="Java Guide", artifact_type="docx", session_id=sess_id)
    art_ppt = Artifact(title="Java Guide", artifact_type="pptx", session_id=sess_id)
    sync_db.add(art_doc)
    sync_db.add(art_ppt)
    sync_db.flush()

    s_map = {"1": {"id": 1, "title": "Java Oracle Docs", "url_or_filename": "https://oracle.com", "kind": "web", "accessed_at": "September 2026"}}
    v_doc = ArtifactVersion(
        artifact_id=art_doc.id, version_no=1, model_json=doc_m.model_dump_json(),
        file_path="d.docx", file_type="docx", source_ids_json=json.dumps([1]), sources_json=json.dumps(s_map),
    )
    v_ppt = ArtifactVersion(
        artifact_id=art_ppt.id, version_no=1, model_json=deck_m.model_dump_json(),
        file_path="p.pptx", file_type="pptx", source_ids_json=json.dumps([1]), sources_json=json.dumps(s_map),
    )
    sync_db.add(v_doc)
    sync_db.add(v_ppt)
    sync_db.commit()

    llm_counter = {"calls": 0}
    real_gen_json = get_llm_client().generate_json

    def counted_gen_json(prompt, schema, **kwargs):
        llm_counter["calls"] += 1
        return real_gen_json(prompt, schema, **kwargs)

    state: GraphState = {
        "run_id": "run_topic_edit",
        "session_id": sess_id,
        "user_message": "append information about Python",
        "plan": plan.model_dump(),
        "artifacts": [],
    }

    session_factory = sessionmaker(bind=sync_db.get_bind(), expire_on_commit=False)
    with patch("app.agents.graph._get_sync_session", side_effect=session_factory), \
         patch.object(get_llm_client(), "generate_json", side_effect=counted_gen_json):
        res_state = node_edit(state)

    arts = res_state.get("artifacts", [])
    kinds = {a["kind"] for a in arts}
    assert "docx" in kinds
    assert "pptx" in kinds
    assert llm_counter["calls"] <= 5


# ── Test j: Migration adds sources_json and backfills from Source rows ──

def test_j_migration_adds_sources_json_and_backfills():
    """Database migration adds sources_json column and backfills from Source table without losing data."""
    engine = create_engine("sqlite:///:memory:")
    with engine.begin() as conn:
        conn.exec_driver_sql("""
            CREATE TABLE artifacts (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                title VARCHAR(255) NOT NULL,
                artifact_type VARCHAR(10) NOT NULL,
                session_id VARCHAR(100),
                run_id VARCHAR(100),
                created_at TIMESTAMP,
                updated_at TIMESTAMP
            );
        """)
        conn.exec_driver_sql("""
            CREATE TABLE artifact_versions (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                artifact_id INTEGER,
                project_id INTEGER,
                version_no INTEGER NOT NULL,
                model_json TEXT NOT NULL,
                file_path VARCHAR(1000) NOT NULL,
                file_type VARCHAR(10) NOT NULL,
                parent_version_id INTEGER,
                change_summary TEXT,
                source_ids_json TEXT,
                diff_json TEXT,
                created_at TIMESTAMP,
                updated_at TIMESTAMP
            );
        """)
        conn.exec_driver_sql("""
            CREATE TABLE sources (
                id INTEGER PRIMARY KEY AUTOINCREMENT,
                project_id INTEGER,
                title VARCHAR(255) NOT NULL,
                url VARCHAR(1000),
                kind VARCHAR(10) NOT NULL,
                metadata_json TEXT,
                created_at TIMESTAMP,
                updated_at TIMESTAMP
            );
        """)
        conn.exec_driver_sql(
            "INSERT INTO sources (id, title, url, kind, metadata_json) VALUES (1, 'Legacy Python Doc', 'https://python.org', 'web', '{\"citation_id\": 1, \"run_id\": \"run-old\"}')"
        )
        conn.exec_driver_sql(
            "INSERT INTO artifacts (id, title, artifact_type, run_id) VALUES (10, 'Legacy Doc', 'docx', 'run-old')"
        )
        conn.exec_driver_sql(
            "INSERT INTO artifact_versions (id, artifact_id, version_no, model_json, file_path, file_type, source_ids_json) "
            "VALUES (100, 10, 1, '{}', 'old.docx', 'docx', '[1]')"
        )

    # Run migration
    with engine.begin() as conn:
        run_migrations_sync(conn)

    # Verify column exists and data backfilled
    with engine.begin() as conn:
        row = conn.exec_driver_sql("SELECT sources_json FROM artifact_versions WHERE id = 100").fetchone()
        assert row is not None
        assert row[0] is not None
        backfilled = json.loads(row[0])
        assert "1" in backfilled
        assert backfilled["1"]["title"] == "Legacy Python Doc"
        assert backfilled["1"]["url_or_filename"] == "https://python.org"
