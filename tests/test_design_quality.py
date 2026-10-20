"""Expert-level slide polish on any template (ADR-037): copy fitting, template colours, photos in the template's own
frames, text clear of artwork, measurement with the fonts PowerPoint uses."""

from __future__ import annotations

import json
from pathlib import Path

import pptx
import pytest
from fastapi.testclient import TestClient
from PIL import Image

from app.models.blocks import TimelineBlock, TimelineEvent
from app.models.deck_model import BulletItem, DeckModel, SlideBulletsBlock, SlideModel
from app.services.analyzers.slide_design import analyze_slide_design
from app.services.copyfit import format_number, group_digits, smart_trim, strip_citations
from app.services.images import PhotoAsset
from app.services.layout.colors import tokens_from_template
from app.services.pptx_renderer import render_pptx
from app.services.slide_polish import polish_deck

GREEN = Path("data/sample_templates/Green Cream Simple Aesthetic Watercolor Presentation.pptx")
PRES = Path("data/sample_templates/presentation_Template.pptx")


def test_headlines_end_at_a_clause_boundary_never_with_an_ellipsis():
    long = ("Sachin Tendulkar holds the world records for the most international runs (34,357), "
            "most Test runs (15,921) and most centuries")
    short = smart_trim(long, 12)
    assert short == "Sachin Tendulkar holds the world records for the most international runs"
    assert "…" not in short and "..." not in short
    assert smart_trim("Only 34% of leaders are reimagining their business", 12) == "Only 34% of leaders are reimagining their business"


def test_citations_once_and_numbers_readable():
    assert strip_citations("Significant investment in AI startups [2]. [2]") == "Significant investment in AI startups"
    assert group_digits("34357 runs in 2026") == "34,357 runs in 2026"
    assert format_number(15_700_000_000_000) == "15.7 trillion" and format_number(2026) == "2026"


def test_polish_turns_a_timeline_of_one_date_into_points_and_is_idempotent():
    deck = DeckModel(title="T", slides=[SlideModel(archetype="timeline", title="What happens by 2026", blocks=[
        TimelineBlock(events=[TimelineEvent(date="By 2026", label="Agents"), TimelineEvent(date="By 2026", label="Multimodal"),
                              TimelineEvent(date="As of 2026", label="Robotics")])])])
    once = polish_deck(deck.model_copy(deep=True))
    assert once.slides[0].archetype == "bullets" and once.slides[0].blocks[0].type == "bullets"
    assert polish_deck(once.model_copy(deep=True)) == once


def test_template_accent_comes_from_its_own_artwork_not_the_theme_default_blue():
    design = analyze_slide_design(pptx.Presentation(PRES))
    assert tokens_from_template(design).primary.upper() not in ("4472C4", "156082")


def test_sample_photographs_are_found_in_picture_filled_shapes():
    design = analyze_slide_design(pptx.Presentation(GREEN))
    assert [b.index for b in design.slide_bases if b.sample_photos] == [4, 7, 8]


def _photo(tmp_path: Path) -> PhotoAsset:
    path = tmp_path / "p.jpg"
    Image.effect_mandelbrot((1600, 1000), (-2, -1.2, 1, 1.2), 60).convert("RGB").save(path)
    return PhotoAsset(path=path, credit="Photo: test", landing_url="https://example.org", alt_text="topic photo")


def _bullets(n: int) -> list:
    return [SlideBulletsBlock(items=[BulletItem(text=f"Point {i} about the subject", source_ids=[1]) for i in range(n)])]


def test_topic_photo_replaces_the_sample_photo_inside_the_template_frame(tmp_path):
    deck = DeckModel(title="Gen AI", slides=[SlideModel(role="title", title="Gen AI")] + [
        SlideModel(archetype="bullets", title=f"Slide {i} headline", blocks=_bullets(3)) for i in range(1, 6)])
    out = tmp_path / "d.pptx"
    photo = _photo(tmp_path)
    render_pptx(deck, GREEN, None, out, photos={i: photo for i in range(1, 6)})
    prs = pptx.Presentation(out)
    swapped = [sh for s in prs.slides for sh in s.shapes if sh.name == "content:image"]
    assert swapped, "no photo placed"
    for slide in prs.slides:
        images = [sh for sh in slide.shapes if sh.name == "content:image"]
        bodies = [sh for sh in slide.shapes if sh.name == "content:body"]
        for img in images:
            for body in bodies:  # text never runs under the photo
                assert int(body.left) + int(body.width) <= int(img.left) or int(img.left) + int(img.width) <= int(body.left)


def test_cover_title_stays_on_the_readable_panel(tmp_path):
    deck = DeckModel(title="Generative AI in 2026", slides=[
        SlideModel(role="title", title="Generative AI in 2026: From Pilots to Production", subtitle="An outlook")])
    out = tmp_path / "d.pptx"
    render_pptx(deck, PRES, None, out)
    title = next(sh for sh in pptx.Presentation(out).slides[0].shapes if sh.name == "content:title")
    assert (int(title.left) + int(title.width)) / 914400 <= 10.3  # the yellow panel ends at 10.17 in


def test_embedded_template_fonts_are_used_for_measurement():
    from app.services.layout.fonts import has_exact_metrics, register_embedded_fonts, resolve_font

    register_embedded_fonts(PRES)
    assert has_exact_metrics("Montserrat Bold")
    assert "montserrat" in str(resolve_font("Montserrat Bold", True)).lower()


def _login(api: TestClient) -> None:
    token = api.post("/auth/login", json={"username": "demo", "password": "demo123"}).json()["access_token"]
    api.headers.update({"Authorization": f"Bearer {token}"})


def test_use_my_template_with_nothing_attached_uses_the_default_and_says_so():
    from app.agents.graph import _get_sync_session
    from app.main import app
    from app.models.run import AgentRun

    with TestClient(app) as api:
        _login(api)
        res = api.post("/chat", json={"message": "create a presentation on solar microgrids using my uploaded template"}).json()
    db = _get_sync_session()
    run = db.query(AgentRun).filter(AgentRun.run_id == res["run_id"]).first()
    result = json.loads(run.result_json)
    db.close()
    assert [a["kind"] for a in result["artifacts"]] == ["pptx"]
    assert "default template" in result["reply"]
    assert "You can ask for changes" in result["reply"]
