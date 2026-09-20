"""Automated unit tests for DOCX and PPTX renderers asserting quality rules, font minimums, and placeholder cleanup."""

from pathlib import Path
import docx
import pptx
import pytest

from app.models.deck_model import BulletItem, DeckModel, SlideModel
from app.models.document_model import (
    BulletsBlock,
    DocumentModel,
    ParagraphBlock,
    Section,
    TableBlock,
)
from app.agents.doc_analyzer import analyze_document
from app.agents.ppt_analyzer import analyze_presentation
from app.services.docx_renderer import render_docx
from app.services.pptx_renderer import render_pptx

DOCX_TEMPLATE = Path("data/sample_templates/Company_Proposal.docx")
PPTX_TEMPLATE = Path("data/sample_templates/Company_Template.pptx")


def test_docx_renderer_quality_and_layout(tmp_path: Path):
    """Test rendering DocumentModel to DOCX file verifying keep_with_next and styles."""
    assert DOCX_TEMPLATE.exists()
    profile = analyze_document(DOCX_TEMPLATE)

    doc_model = DocumentModel(
        title="Generative AI Transformation Proposal",
        subtitle="Technical & Commercial Engagement Architecture",
        client_name="Apex Global Enterprise",
        date="September 2026",
        sections=[
            Section(
                heading="1. Introduction & Context",
                level=1,
                blocks=[
                    ParagraphBlock(text="NexaWorks AI Solutions provides enterprise generative AI engineering.", source_ids=[1]),
                    BulletsBlock(items=["Multi-agent workflow orchestration", "Pinecone hybrid vector search"], source_ids=[1, 2]),
                    TableBlock(headers=["Phase", "Cost"], rows=[["Phase 1", "INR 480k"], ["Phase 2", "INR 720k"]], source_ids=[3]),
                ]
            )
        ]
    )

    sources_map = {
        1: {"title": "Source One", "url": "https://source1.com", "snippet": "Snippet 1"},
        2: {"title": "Source Two", "url": "https://source2.com", "snippet": "Snippet 2"},
        3: {"title": "Source Three", "url": "https://source3.com", "snippet": "Snippet 3"},
    }

    out_file = tmp_path / "test_rendered.docx"
    render_docx(doc_model, DOCX_TEMPLATE, profile, out_file, sources_map=sources_map)

    assert out_file.exists()
    assert out_file.stat().st_size > 0

    # Reopen rendered DOCX and verify headings and keep_with_next
    doc = docx.Document(out_file)
    headings = [p for p in doc.paragraphs if p.style and p.style.name.startswith("Heading")]
    assert len(headings) >= 2  # Section heading + Sources heading

    for h in headings:
        assert h.paragraph_format.keep_with_next is True

    # Assert Sources section exists
    assert any("Sources & References" in p.text for p in doc.paragraphs)


def test_pptx_renderer_quality_and_font_minimums(tmp_path: Path):
    """Test rendering DeckModel to PPTX file verifying font minimums, placeholder cleanup, and layout rotation."""
    assert PPTX_TEMPLATE.exists()
    profile = analyze_presentation(PPTX_TEMPLATE)

    slides_list = [
        SlideModel(role="title", title="Generative AI Transformation Strategy", subtitle="NexaWorks AI Solutions for Indian Enterprise", source_ids=[1]),
        SlideModel(role="section_header", title="Executive Context & Industry Trends", subtitle="Rapid adoption of agentic workflows across Indian enterprise sectors", source_ids=[1]),
        SlideModel(role="title_content", title="Enterprise Market Trends 2026", bullets=[
            BulletItem(text="72% of mid-size Indian enterprises plan agentic workflow adoption by Q4 2026.", level=0, source_ids=[1]),
            BulletItem(text="Manual proposal creation consumes 4.2 hours per client pitch deck.", level=0, source_ids=[1]),
        ]),
        SlideModel(role="two_content", title="Manual Workflows vs NexaWorks AI", left=[
            BulletItem(text="Manual 4.2 hours proposal SLA", level=0, source_ids=[1]),
        ], right=[
            BulletItem(text="Automated 15-minute generation SLA", level=0, source_ids=[1]),
        ]),
        SlideModel(role="section_header", title="NexaWorks Architecture & Capabilities", subtitle="Turnkey multi-agent automation using LangGraph and Pinecone"),
        SlideModel(role="title_content", title="Multi-Agent System Capabilities", bullets=[
            BulletItem(text="LangGraph supervisor agent orchestrates specialized worker sub-agents.", level=0),
            BulletItem(text="Pinecone serverless hybrid vector search achieves 99.4% retrieval precision.", level=0),
        ]),
        SlideModel(role="title_content", title="Document & PPT Template Engineering", bullets=[
            BulletItem(text="Extracts font maps, color palettes, and margin tokens from uploaded templates.", level=0),
        ]),
        SlideModel(role="section_header", title="Proven Enterprise Case Studies", subtitle="Demonstrated SLA reductions across Finance, Healthcare, and Supply Chain"),
        SlideModel(role="title_content", title="Financial Services & Healthcare Impact", bullets=[
            BulletItem(text="80% reduction in turnaround time for investment memo processing.", level=0),
        ]),
        SlideModel(role="title_content", title="Supply Chain RFP Automation Impact", bullets=[
            BulletItem(text="Proposal turnaround time reduced from 7 days to under 24 hours.", level=0),
        ]),
        SlideModel(role="two_content", title="Flexible Engagement Structures", left=[
            BulletItem(text="Turnkey Fixed-Price Sprints", level=0),
        ], right=[
            BulletItem(text="Time & Materials Rate Card", level=0),
        ]),
        SlideModel(role="title_only", title="Initiate Enterprise Transformation", subtitle="Contact NexaWorks AI Solutions"),
    ]

    deck_model = DeckModel(title="12 Slide Quality Test Deck", slides=slides_list)

    sources_map = {
        1: {"title": "Source One", "url": "https://source1.com", "snippet": "Snippet 1"},
    }

    out_file = tmp_path / "test_rendered.pptx"
    render_pptx(deck_model, PPTX_TEMPLATE, profile, out_file, sources_map=sources_map)

    assert out_file.exists()
    assert out_file.stat().st_size > 0

    prs = pptx.Presentation(out_file)

    # 1. Assert slide count (12 deck slides + 1 Sources slide = 13 slides total)
    assert len(prs.slides) == 13

    # 2. Assert layout choices used across 12 slides are NOT all the same
    used_layout_names = {s.slide_layout.name for s in prs.slides}
    assert len(used_layout_names) > 1, f"Layouts used should be varied, found only: {used_layout_names}"

    # 3. Assert no empty text placeholders left on any slide
    for slide_idx, slide in enumerate(prs.slides, start=1):
        for shape in slide.placeholders:
            if shape.has_text_frame:
                txt = shape.text_frame.text.strip()
                assert len(txt) > 0, f"Slide {slide_idx} has empty text placeholder '{shape.name}'"

    # 4. Assert font sizes are >= minimums
    cover_slide = prs.slides[0]
    cover_title_p = cover_slide.shapes.title.text_frame.paragraphs[0]
    assert cover_title_p.font.size.pt >= 28.0, f"Cover title font size {cover_title_p.font.size.pt} is below min 28pt"

    for slide_idx, slide in enumerate(list(prs.slides)[1:], start=2):
        for shape in slide.shapes:
            if shape.has_text_frame:
                for p in shape.text_frame.paragraphs:
                    if p.font.size:
                        assert p.font.size.pt >= 14.0, f"Slide {slide_idx} paragraph font size {p.font.size.pt} is below min 14pt"
