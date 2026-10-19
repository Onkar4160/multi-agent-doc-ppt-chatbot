"""Stage 3 foundation: upload safety, errors, logging, LLM budget, ownership, migrations, contracts, CLI."""

import io
import json
import logging
import uuid
import zipfile
from pathlib import Path
from unittest.mock import patch

import pytest
from fastapi.testclient import TestClient
from pydantic import ValidationError


# ── Upload safety (NFR-6) ───────────────────────────────────────────────────


def _ooxml(main_part: str, extra: dict[str, bytes] | None = None) -> bytes:
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("[Content_Types].xml", "<Types/>")
        zf.writestr(main_part, "<root/>")
        for name, data in (extra or {}).items():
            zf.writestr(name, data)
    return buf.getvalue()


def test_upload_checks_judge_content_not_names():
    from app.core.errors import UnsupportedFileError
    from app.core.file_safety import check_upload

    assert check_upload("deck.pptx", _ooxml("ppt/presentation.xml"), 25) == ".pptx"
    assert check_upload("notes.csv", b"a,b\n1,2\n", 25) == ".csv"
    assert check_upload("paper.pdf", b"%PDF-1.7\n...", 25) == ".pdf"
    cases = [
        ("deck.pptx", _ooxml("word/document.xml"), "required parts"),          # a Word file renamed .pptx
        ("deck.pptx", b"PK\x03\x04not really a zip", "cannot be opened"),
        ("deck.pptm", _ooxml("ppt/presentation.xml"), "Macro-enabled"),
        ("deck.pptx", _ooxml("ppt/presentation.xml", {"ppt/vbaProject.bin": b"x"}), "macros"),
        ("paper.pdf", b"<html>", "not a PDF"),
        ("photo.png", b"\xff\xd8\xff jpeg bytes", "not a PNG"),
        ("data.csv", b"a,b\x00\x01", "binary data"),
        ("tool.exe", b"MZ", "Unsupported file type"),
        ("big.txt", b"x" * (2 * 1024 * 1024), "larger than"),
    ]
    for name, data, expected in cases:
        with pytest.raises(UnsupportedFileError) as info:
            check_upload(name, data, 1)
        assert expected in info.value.message, (name, info.value.message)


def test_zip_bomb_is_refused():
    from app.core.errors import UnsupportedFileError
    from app.core.file_safety import check_upload

    bomb = _ooxml("xl/workbook.xml", {"xl/huge.xml": b"0" * (60 * 1024 * 1024)})  # compresses to ~60 KB
    with pytest.raises(UnsupportedFileError, match="unsafe size"):
        check_upload("book.xlsx", bomb, 25)


# ── Errors, logging, budget ─────────────────────────────────────────────────


def test_app_errors_become_plain_replies():
    from app.core.errors import BudgetExceeded, TemplateError
    from app.llm.errors import friendly_error

    text = friendly_error(TemplateError("the template has no slides to copy"), "read the template")
    assert text == "I couldn't read the template: the template has no slides to copy. Nothing was delivered. Say “try again” to retry."
    assert "LLM_MAX_CALLS_PER_RUN" in friendly_error(BudgetExceeded("budget spent"), "build the slides")


def test_log_records_carry_the_run_id(caplog):
    from app.core.logging import RunIdFilter, run_context

    logger = logging.getLogger("test.runid")
    caplog.handler.addFilter(RunIdFilter())
    with caplog.at_level(logging.INFO, logger="test.runid"):
        with run_context("run_abc123"):
            logger.info("inside")
        logger.info("outside")
    ids = [r.run_id for r in caplog.records if r.name == "test.runid"]
    assert ids == ["run_abc123", "-"]


def _scripted_client(outcomes):
    from app.llm.client import LLMCallStats, LLMClient

    c = LLMClient.__new__(LLMClient)
    c._primary, c._fallbacks, c._max_retries = "m1", ["m2"], 3
    c._stats, c._last_record, c._sleep = LLMCallStats(), None, lambda s: None
    it = iter(outcomes)

    def single_call(model, prompt, **kwargs):
        result = next(it)
        if isinstance(result, Exception):
            raise result
        return result

    c._single_call = single_call
    return c


def test_run_budget_stops_further_calls_and_usage_is_counted():
    from app.core.errors import BudgetExceeded
    from app.llm import client as client_mod
    from app.llm.usage import track_run

    client_mod._unavailable.clear()
    c = _scripted_client(["one", "two", "three"])
    with track_run(budget=2) as usage:
        assert c._call_with_retry("a") == "one"
        assert c._call_with_retry("b") == "two"
        with pytest.raises(BudgetExceeded):
            c._call_with_retry("c")
    assert usage.calls == 2 and usage.by_model == {"m1": 2}
    with track_run(budget=None) as unlimited:  # outside a budget nothing is limited
        assert c._call_with_retry("d") == "three"
    assert unlimited.calls == 1


def test_failed_attempts_count_against_the_budget():
    from app.llm import client as client_mod
    from app.llm.usage import track_run

    class Busy(Exception):
        code = 503

    client_mod._unavailable.clear()
    c = _scripted_client([Busy(), "ok"])
    with track_run(budget=5) as usage:
        assert c._call_with_retry("x") == "ok"
    assert usage.calls == 2


def test_refused_api_key_stops_the_chain_with_a_clear_reply():
    """A 403 applies to the whole Google project: trying other models only spends budget (seen live 2026-10-18)."""
    from app.llm import client as client_mod
    from app.llm.errors import LLMAuthError, friendly_error
    from app.llm.usage import track_run

    class Denied(Exception):
        code = 403
        details = {"error": {"code": 403, "message": "Your project has been denied access. Please contact support."}}

    client_mod._unavailable.clear()
    c = _scripted_client([Denied(), "never reached"])
    with track_run(budget=10) as usage, pytest.raises(LLMAuthError) as info:
        c._call_with_retry("x")
    assert usage.calls == 1
    text = friendly_error(info.value, "build the presentation")
    assert "refused the API key" in text and "denied access" in text and "GEMINI_API_KEY" in text


# ── Contracts (DATA_MODELS) ─────────────────────────────────────────────────


def test_interpretation_invariants():
    from app.models.interpretation import ChangeRequest, ChangeScope, Interpretation, LengthSpec

    with pytest.raises(ValidationError, match="change request"):
        Interpretation(kind="modify_artifact", confidence=0.9)
    with pytest.raises(ValidationError, match="clarification"):
        Interpretation(kind="new_task", confidence=0.3)
    with pytest.raises(ValidationError, match="at least one target"):
        ChangeScope(level="unit")
    with pytest.raises(ValidationError, match="style change"):
        ChangeRequest(operation="restyle", scope=ChangeScope(level="style"), instruction="bigger headings")
    with pytest.raises(ValidationError, match="explicit length"):
        LengthSpec(unit="auto", explicit=True)
    ok = Interpretation.model_validate({
        "kind": "modify_artifact", "confidence": 0.85, "objective": "EV deck",
        "change": {"operation": "restyle", "scope": {"level": "style"}, "instruction": "headings 20pt",
                   "style_changes": [{"path": "heading_1.size_pt", "value": 20}]},
    })
    assert Interpretation.model_validate_json(ok.model_dump_json()) == ok


def test_dossier_traceability_invariants():
    from app.models.evidence import Claim, ResearchDossier

    sources = [{"id": 1, "kind": "paper", "title": "Cold climate heat pumps"}]
    evidence = [{"id": "e1", "source_id": 1, "statement": "COP stays above 2 at -15 C", "value": 2.0, "as_of": "2023"}]
    with pytest.raises(ValidationError, match="supported but cites no evidence"):
        Claim(id="c1", statement="x", status="supported")
    with pytest.raises(ValidationError, match="unknown source"):
        ResearchDossier(sources=sources, evidence=[{**evidence[0], "source_id": 9}])
    with pytest.raises(ValidationError, match="unknown evidence"):
        ResearchDossier(sources=sources, evidence=evidence, claims=[{"id": "c1", "statement": "x", "status": "supported", "evidence_ids": ["e9"]}])
    dossier = ResearchDossier(sources=sources, evidence=evidence, claims=[
        {"id": "c1", "statement": "Heat pumps work at -15 C", "status": "supported", "evidence_ids": ["e1"]},
        {"id": "c2", "statement": "Prices will fall", "status": "estimate"},
        {"id": "c3", "statement": "Everyone loves them", "status": "unsupported"},
    ])
    assert dossier.coverage_pct() == 66.7


def test_quality_report_derives_counts_and_verdict():
    from app.models.quality import QualityReport

    report = QualityReport(findings=[
        {"check_id": "pptx.text_overflow", "tier": 2, "severity": "critical", "status": "corrected", "message": "fixed"},
        {"check_id": "pptx.font_floor", "tier": 2, "severity": "major", "status": "unresolved", "message": "12pt"},
    ], not_performed=[{"check_id": "common.convert", "tier": 4, "reason": "LibreOffice not installed"}])
    assert report.counts == {"detected": 2, "corrected": 1, "unresolved": 1, "not_performed": 1}
    assert report.verdict == "pass_with_warnings"
    assert QualityReport(findings=[{"check_id": "x", "tier": 1, "severity": "critical", "message": "m"}]).verdict == "fail"
    assert QualityReport().verdict == "pass"


def test_content_models_carry_schema_version_and_old_json_still_loads():
    from app.models.deck_model import DeckModel

    old = '{"title": "Old deck", "slides": [{"role": "title", "title": "Hello"}]}'
    deck = DeckModel.model_validate_json(old)
    assert deck.schema_version == 1 and deck.slides[0].title == "Hello"


def test_every_env_example_key_is_read_by_the_settings():
    """A documented variable the config ignores is a silent misconfiguration (UPLOAD_MAX_MB was one)."""
    import re

    from app.core.config import Settings

    known = set()
    for name, field in Settings.model_fields.items():
        alias = field.validation_alias
        if alias is None:
            known.add(name.upper())
        elif isinstance(alias, str):
            known.add(alias)
        else:
            known.update(alias.choices)
    documented = [line.split("=")[0] for line in Path(".env.example").read_text(encoding="utf-8").splitlines() if re.match(r"^[A-Z_]+=", line)]
    assert documented and not [k for k in documented if k not in known]


# ── Persistence ─────────────────────────────────────────────────────────────


def test_migrations_add_new_columns_to_an_old_database(tmp_path):
    from sqlalchemy import create_engine, text

    from app.core.database import ADDED_COLUMNS, run_migrations_sync

    engine = create_engine(f"sqlite:///{(tmp_path / 'old.db').as_posix()}")
    with engine.begin() as conn:
        conn.execute(text("CREATE TABLE chat_sessions (id INTEGER PRIMARY KEY, session_id VARCHAR(100), title VARCHAR(200))"))
        conn.execute(text("CREATE TABLE artifact_versions (id INTEGER PRIMARY KEY, model_json TEXT, file_path TEXT, file_type TEXT)"))
        run_migrations_sync(conn)
        run_migrations_sync(conn)  # idempotent
        cols = {row[1] for row in conn.execute(text("PRAGMA table_info(artifact_versions)"))}
        assert {"quality_json", "dossier_id"} <= cols
        assert "owner" in {row[1] for row in conn.execute(text("PRAGMA table_info(chat_sessions)"))}
    assert ("tasks", "owner", "VARCHAR(100)") not in ADDED_COLUMNS  # new tables are created whole, not migrated


@pytest.fixture(name="api")
def api_fixture():
    from app.main import app

    with TestClient(app) as client:
        yield client


def _login(api, username: str, password: str = "pw-12345") -> dict:
    from app.agents.graph import _get_sync_session
    from app.core.security import hash_password
    from app.models.user import User

    if username != "demo":
        db = _get_sync_session()
        if not db.query(User).filter(User.username == username).first():
            db.add(User(username=username, hashed_password=hash_password(password)))
            db.commit()
        db.close()
    else:
        password = "demo123"
    token = api.post("/auth/login", json={"username": username, "password": password}).json()["access_token"]
    return {"Authorization": f"Bearer {token}"}


def test_interrupted_runs_are_failed_with_a_reply_at_startup(api):
    import asyncio

    from app.agents.graph import _get_sync_session
    from app.core.database import INTERRUPTED_REPLY, fail_interrupted_runs, get_session_factory
    from app.models.chat import ChatMessage, ChatSession
    from app.models.run import AgentRun

    sid, rid = f"sess_{uuid.uuid4().hex[:8]}", f"run_{uuid.uuid4().hex[:8]}"
    db = _get_sync_session()
    db.add(ChatSession(session_id=sid, title="t"))
    db.add(ChatMessage(session_id=sid, role="user", content="Make a deck on wind power", run_id=rid))
    db.add(AgentRun(run_id=rid, session_id=sid, status="running"))
    db.commit()
    db.close()

    async def run():
        async with get_session_factory()() as session:
            return await fail_interrupted_runs(session)

    assert asyncio.run(run()) >= 1
    db = _get_sync_session()
    run_rec = db.query(AgentRun).filter(AgentRun.run_id == rid).first()
    reply = db.query(ChatMessage).filter(ChatMessage.run_id == rid, ChatMessage.role == "assistant").first()
    db.close()
    assert run_rec.status == "failed" and json.loads(run_rec.result_json)["request"] == "Make a deck on wind power"
    assert reply.content == INTERRUPTED_REPLY


# ── Ownership (FR-UX-3) ─────────────────────────────────────────────────────


def test_other_users_records_look_missing(api, tmp_path):
    from app.agents.graph import _get_sync_session
    from app.models.artifact import Artifact
    from app.models.deck_model import DeckModel
    from app.services.versioning import create_version_sync

    alice, bob = _login(api, "alice"), _login(api, "bob")
    with patch("app.api.chat.graph_app.invoke", side_effect=lambda s: {**s, "reply": "ok"}):
        res = api.post("/chat", json={"message": "Alice's private deck"}, headers=alice).json()
    sid, rid = res["session_id"], res["run_id"]

    out = tmp_path / "deck.pptx"
    out.write_bytes(b"pptx")
    db = _get_sync_session()
    art = Artifact(artifact_type="pptx", title="Alice deck", session_id=sid, run_id=rid)
    db.add(art)
    db.flush()
    create_version_sync(db=db, artifact_id=art.id, model_json=DeckModel(title="Alice deck").model_dump_json(), source_file_path=out)
    db.commit()
    art_id = art.id
    db.close()

    assert api.get(f"/sessions/{sid}/messages", headers=alice).status_code == 200
    assert api.get(f"/runs/{rid}", headers=alice).status_code == 200
    for path in (f"/sessions/{sid}/messages", f"/sessions/{sid}/artifacts", f"/runs/{rid}", f"/runs/{rid}/trace",
                 f"/artifacts/{art_id}/versions", f"/artifacts/{art_id}/download"):
        assert api.get(path, headers=bob).status_code == 404, path
    assert sid not in [s["session_id"] for s in api.get("/sessions", headers=bob).json()]
    assert sid in [s["session_id"] for s in api.get("/sessions", headers=alice).json()]
    assert art_id not in [a["id"] for a in api.get("/artifacts", headers=bob).json()]
    # Bob cannot post into Alice's session either.
    assert api.post("/chat", json={"message": "hi", "session_id": sid}, headers=bob).status_code == 404
    # Browser download links carry the JWT as a query parameter.
    token = alice["Authorization"].split()[1]
    assert api.get(f"/artifacts/{art_id}/download", params={"token": token}).status_code == 200
    assert api.get(f"/artifacts/{art_id}/download").status_code == 401


def test_uploads_are_owned(api):
    import docx

    alice, bob = _login(api, "alice"), _login(api, "bob")
    d = docx.Document()
    d.add_paragraph(f"owned {uuid.uuid4()}")
    buf = io.BytesIO()
    d.save(buf)
    up = api.post("/files/upload", files={"file": ("brief.docx", buf.getvalue(), "application/octet-stream")}, headers=alice).json()
    assert up["id"] in [f["id"] for f in api.get("/files", headers=alice).json()]
    assert up["id"] not in [f["id"] for f in api.get("/files", headers=bob).json()]
    assert api.get(f"/files/{up['id']}/profile", headers=bob).status_code == 404
    bad = api.post("/files/upload", files={"file": ("macro.docm", b"PK", "application/octet-stream")}, headers=alice)
    assert bad.status_code == 400 and "Macro-enabled" in bad.json()["detail"]


def test_run_result_records_llm_usage(api):
    headers = _login(api, "demo")
    with patch("app.api.chat.graph_app.invoke", side_effect=lambda s: {**s, "reply": "done"}):
        res = api.post("/chat", json={"message": "usage please"}, headers=headers).json()
    from app.agents.graph import _get_sync_session
    from app.models.run import AgentRun

    db = _get_sync_session()
    run = db.query(AgentRun).filter(AgentRun.run_id == res["run_id"]).first()
    db.close()
    usage = json.loads(run.result_json)["usage"]
    assert usage["budget"] >= 1 and usage["calls"] == 0


# ── CLI (FR-UX-5) ───────────────────────────────────────────────────────────


def test_cli_analyze_template_and_health(capsys):
    from app.cli import main

    assert main(["analyze-template", "data/sample_templates/Green Cream Simple Aesthetic Watercolor Presentation.pptx"]) == 0
    summary = json.loads(capsys.readouterr().out)
    assert summary["design_mode"] == "slides" and summary["bases"][0] == "cover" and summary["roles"]
    assert main(["health"]) == 0
    assert json.loads(capsys.readouterr().out)["mock_llm"] is True
    assert main(["analyze-template", "missing.pptx"]) == 2


def test_cli_generate_runs_the_workflow(tmp_path, capsys):
    from app.cli import main

    code = main(["generate", "Make a 4-slide deck on community solar", "--out", str(tmp_path / "out")])
    out = capsys.readouterr().out
    result = json.loads(out[out.index("{"):])  # logs go to stderr; stdout carries the JSON summary
    assert code == 0, result
    assert result["status"] == "done" and result["usage"]["budget"] >= 1
    assert any(Path(f).suffix == ".pptx" and Path(f).exists() for f in result["files"])
