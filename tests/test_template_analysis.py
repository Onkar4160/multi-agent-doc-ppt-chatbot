"""Stage 6 — template intelligence (FR-TPL-1..7).

Manifest snapshots pin what the analysers extract from every sample template; hand labels (made by inspecting the
templates, independently of the classifier) measure element-role and layout-archetype accuracy. Set
``UPDATE_SNAPSHOTS=1`` to rewrite the snapshots after a deliberate analyser change (and bump ANALYZER_VERSION).
"""

from __future__ import annotations

import json
import os
from pathlib import Path

import pytest

from app.models.template_profile import ANALYZER_VERSION, TemplateProfile, TemplateProfileRecord
from app.services.analyzers.summary import NoLLM, manifest_summary
from tests.fixtures.build_fixtures import as_template_package, build_pptx_with_master_logo, build_xlsx_template

SAMPLES = Path("data/sample_templates")
MANIFESTS = Path("tests/fixtures/manifests")
PPTX_SAMPLES = ["Company_Template.pptx", "Green Cream Simple Aesthetic Watercolor Presentation.pptx", "presentation_Template.pptx"]
DOCX_SAMPLES = ["Company_Proposal.docx", "proposal_Template.docx"]
HAND = json.loads((MANIFESTS / "hand_labels.json").read_text(encoding="utf-8"))


def _analyze(path: Path) -> TemplateProfile:
    suffix = path.suffix.lower()
    if suffix in (".pptx", ".potx"):
        from app.agents.ppt_analyzer import analyze_presentation

        return analyze_presentation(path, llm_client=NoLLM())
    if suffix in (".docx", ".dotx"):
        from app.agents.doc_analyzer import analyze_document

        return analyze_document(path, llm_client=NoLLM())
    from app.agents.xlsx_analyzer import analyze_workbook

    return analyze_workbook(path)


@pytest.fixture(scope="module")
def profiles() -> dict[str, TemplateProfile]:
    return {name: _analyze(SAMPLES / name) for name in PPTX_SAMPLES + DOCX_SAMPLES}


def _check_snapshot(name: str, summary: dict) -> None:
    snap = MANIFESTS / f"{Path(name).stem.replace(' ', '_')}.json"
    text = json.dumps(summary, indent=2, sort_keys=True, default=str) + "\n"
    if os.getenv("UPDATE_SNAPSHOTS") == "1" or not snap.exists():
        snap.write_text(text, encoding="utf-8")
        if os.getenv("UPDATE_SNAPSHOTS") != "1":
            pytest.skip(f"snapshot {snap.name} created; rerun to compare")
    assert json.loads(text) == json.loads(snap.read_text(encoding="utf-8")), (
        f"{name}: manifest changed. If deliberate, bump ANALYZER_VERSION and rerun with UPDATE_SNAPSHOTS=1"
    )


@pytest.mark.parametrize("name", PPTX_SAMPLES + DOCX_SAMPLES)
def test_sample_manifest_snapshot(profiles, name):
    _check_snapshot(name, manifest_summary(profiles[name]))


def test_xlsx_manifest_snapshot(tmp_path):
    path = build_xlsx_template(tmp_path / "Revenue_Model.xlsx")
    summary = manifest_summary(_analyze(path))
    _check_snapshot("Revenue_Model.xlsx", summary)


def test_element_role_accuracy_against_hand_labels(profiles):
    total = correct = 0
    misses = []
    for name, slides in HAND["roles"].items():
        inv = {s.index: {e.shape_index: e.role for e in s.elements} for s in profiles[name].ppt_style.inventory}
        first = min(inv)  # inventory index base (0 or 1)
        for slide_no, shapes in slides.items():
            got = inv[int(slide_no) - 1 + first]
            for idx, role in shapes.items():
                total += 1
                if got.get(int(idx)) == role:
                    correct += 1
                else:
                    misses.append(f"{name} s{slide_no}#{idx}: want {role}, got {got.get(int(idx))}")
    assert total >= 50
    assert correct / total >= 0.9, f"role accuracy {correct}/{total}; misses: {misses}"


def test_layout_archetype_accuracy_against_hand_labels(profiles):
    total = correct = 0
    misses = []
    for name, labels in HAND["layout_archetypes"].items():
        got = {lay.name: lay.archetype for lay in profiles[name].ppt_style.layouts}
        for layout, want in labels.items():
            total += 1
            if got.get(layout) == want:
                correct += 1
            else:
                misses.append(f"{layout}: want {want}, got {got.get(layout)}")
    assert correct / total >= 0.9, f"archetype accuracy {correct}/{total}; misses: {misses}"


def test_every_element_has_role_and_reason(profiles):
    for name in PPTX_SAMPLES:
        for slide in profiles[name].ppt_style.inventory:
            for e in slide.elements:
                assert e.role in {"structural", "branding", "decorative", "placeholder", "sample_content"}
                assert e.reason, f"{name} slide {slide.index} shape {e.shape_index} has no reason"


def test_design_tokens_present_for_every_sample(profiles):
    for name, p in profiles.items():
        t = p.design_tokens
        assert t is not None and t.heading_font and t.body_font and t.primary, name
        assert p.sha256 and len(p.sha256) == 64


def test_docx_manifest_reads_layout_rules(profiles):
    company = profiles["Company_Proposal.docx"].doc_style
    assert company.page_size == "Letter"
    sec = company.manifest.sections[0]
    assert "NexaWorks" in sec.header_text and sec.has_page_number
    body = profiles["proposal_Template.docx"].doc_style.manifest.styles_by_role["body"]
    assert body.font == "Cambria" and body.size_pt == 11  # inherited from docDefaults / Normal


def test_xlsx_manifest_structure(tmp_path):
    xs = _analyze(build_xlsx_template(tmp_path / "m.xlsx")).xlsx_style
    roles = {s.name: s.role for s in xs.sheets}
    assert roles == {"README": "readme", "Assumptions": "assumptions", "Data": "data", "Calculations": "calculations",
                     "Dashboard": "report", "Sources": "sources"}
    by = {s.name: s for s in xs.sheets}
    assert set(by["Calculations"].references_sheets) == {"Assumptions", "Data"}
    assert by["Data"].tables[0]["name"] == "Customers" and by["Data"].freeze_panes == "A2"
    assert by["Dashboard"].charts[0]["type"] == "bar"
    assert {"growth", "churn"} <= set(xs.defined_names)


def test_master_logo_is_branding(tmp_path):
    p = _analyze(build_pptx_with_master_logo(tmp_path / "logo.pptx"))
    brands = [b for b in p.ppt_style.branding if b.kind == "picture"]
    assert brands, "logo on the slide master not detected as branding"
    assert brands[0].role == "branding"


@pytest.mark.parametrize("src,dest", [
    ("Company_Template.pptx", "Company_Template.potx"),
    ("Company_Proposal.docx", "Company_Proposal.dotx"),
])
def test_template_packages_analyse_like_regular_files(profiles, tmp_path, src, dest):
    pkg = as_template_package(SAMPLES / src, tmp_path / dest)
    a, b = manifest_summary(profiles[src]), manifest_summary(_analyze(pkg))
    a.pop("file"), b.pop("file")
    assert a == b


def test_xltx_analyses_like_xlsx(tmp_path):
    x = build_xlsx_template(tmp_path / "m.xlsx")
    t = as_template_package(x, tmp_path / "m.xltx")
    a, b = manifest_summary(_analyze(x)), manifest_summary(_analyze(t))
    a.pop("file"), b.pop("file")
    assert a == b


def test_template_packages_render(tmp_path):
    from app.models.deck_model import DeckModel, SlideModel
    from app.models.document_model import DocumentModel, ParagraphBlock, Section
    from app.services.docx_renderer import render_docx
    from app.services.pptx_renderer import render_pptx

    potx = as_template_package(SAMPLES / "Company_Template.pptx", tmp_path / "t.potx")
    out = tmp_path / "deck.pptx"
    render_pptx(DeckModel(title="Deck", slides=[SlideModel(title="Intro", bullet_points=[])]), potx, TemplateProfile(), out)
    import pptx

    assert len(pptx.Presentation(str(out)).slides) >= 1

    dotx = as_template_package(SAMPLES / "Company_Proposal.docx", tmp_path / "t.dotx")
    out_doc = tmp_path / "doc.docx"
    model = DocumentModel(title="Doc", sections=[Section(heading="Heading", level=1, blocks=[ParagraphBlock(text="Body")])])
    render_docx(model, dotx, TemplateProfile(), out_doc)
    import docx

    assert any("Body" in p.text for p in docx.Document(str(out_doc)).paragraphs)


@pytest.fixture()
def sync_db():
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker

    from app.core.database import Base
    import app.models  # noqa: F401

    engine = create_engine("sqlite:///:memory:")
    Base.metadata.create_all(engine)
    db = sessionmaker(bind=engine)()
    yield db
    db.close()


def test_cached_profile_hits_by_content_and_recomputes_on_version(sync_db, tmp_path, monkeypatch):
    import app.services.template_cache as tc

    calls = []
    real = tc.analyze_template

    def counting(path, file_id=None):
        calls.append(Path(path).name)
        return _analyze(Path(path))

    monkeypatch.setattr(tc, "analyze_template", counting)
    x1 = build_xlsx_template(tmp_path / "a.xlsx")
    x2 = tmp_path / "renamed.xlsx"
    x2.write_bytes(x1.read_bytes())

    p1 = tc.cached_profile(x1, sync_db, file_id=1)
    p2 = tc.cached_profile(x2, sync_db, file_id=2)  # same bytes, other name -> cache hit
    assert calls == ["a.xlsx"]
    assert p2.file_id == 2 and p2.sha256 == p1.sha256
    assert sync_db.query(TemplateProfileRecord).count() == 1

    monkeypatch.setattr(tc, "ANALYZER_VERSION", ANALYZER_VERSION + 1)
    tc.cached_profile(x2, sync_db)
    assert calls == ["a.xlsx", "renamed.xlsx"]
    assert real is not counting


def test_cli_analyze_template_xlsx(tmp_path, capsys):
    from app.cli import main

    path = build_xlsx_template(tmp_path / "m.xlsx")
    assert main(["analyze-template", str(path)]) == 0
    out = json.loads(capsys.readouterr().out)
    assert out["type"] == "xlsx" and len(out["sheets"]) == 6


def test_uploaded_template_packages_resolve_as_templates():
    from app.models.execution_plan import ArtifactType, TemplateMode
    from app.models.file import UploadedFile
    from app.services.template_resolution import resolve_templates

    files = [UploadedFile(id=7, filename="t.potx", file_type="potx", file_size=1, stored_path="t.potx", purpose="TEMPLATE"),
             UploadedFile(id=8, filename="t.dotx", file_type="dotx", file_size=1, stored_path="t.dotx", purpose="TEMPLATE")]
    res = resolve_templates(template_mode=TemplateMode.UPLOADED, artifact_types=[ArtifactType.PPTX, ArtifactType.DOCX],
                            file_ids=[7, 8], uploaded_files=files, user_instruction="Use the uploaded templates.")
    assert res.is_valid and res.pptx_template_file_id == 7 and res.docx_template_file_id == 8
