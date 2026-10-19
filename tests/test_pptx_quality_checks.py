"""Rendered-deck quality checks and the one "shorten" re-plan round (Step 3 of the port)."""

from pathlib import Path
from unittest.mock import patch

import pptx
from pptx.dml.color import RGBColor
from pptx.util import Inches, Pt

from app.agents.validator import ValidationIssue, fit_problems, merge_report, validate_pptx_file
from app.models.deck_model import BulletItem, DeckModel, SlideModel

GREEN = Path("data/sample_templates/Green Cream Simple Aesthetic Watercolor Presentation.pptx")


def _deck_with(tmp_path: Path, build) -> Path:
    prs = pptx.Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    build(prs, slide)
    path = tmp_path / "checked.pptx"
    prs.save(path)
    return path


def _box(slide, name, text, *, x=1.0, y=1.0, w=4.0, h=1.0, size=18, color="1F2328"):
    box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    box.name = name
    tf = box.text_frame
    tf.word_wrap = True
    run = tf.paragraphs[0].add_run()
    run.text = text
    run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    return box


def test_measured_overflow_is_an_error_and_becomes_a_shorten_instruction(tmp_path):
    path = _deck_with(tmp_path, lambda prs, s: _box(s, "content:body", "word " * 200, h=0.6))
    issues = validate_pptx_file(path)
    assert any(i.severity == "error" and "overflows its box" in i.message for i in issues)
    fit = fit_problems(issues)
    assert len(fit) == 1 and fit[0].startswith("Slide 1: text does not fit") and "Shorten" in fit[0]


def test_font_floor_contrast_and_bounds(tmp_path):
    def build(prs, s):
        _box(s, "content:body", "Tiny text", size=8)
        _box(s, "content:lead", "Pale text", y=2.5, color="F4F4F4")
        _box(s, "content:column-1", "Off the slide", x=9.5, y=4.0)
        _box(s, "Template artwork", "Bleeds on purpose", x=9.5, y=5.5)

    messages = [i.message for i in validate_pptx_file(_deck_with(tmp_path, build))]
    assert any("'content:body' uses 8pt, below the 14.0pt minimum" in m for m in messages)
    assert any("'content:lead' has contrast" in m for m in messages)
    assert any("'content:column-1' extends beyond the slide" in m for m in messages)
    assert not any("Template artwork" in m and "beyond" in m for m in messages)  # only our own shapes


def test_template_sample_text_is_flagged():
    issues = validate_pptx_file(GREEN, GREEN)  # the template itself is all sample text
    assert any("template sample text remains" in i.message for i in issues)


def test_rendered_sample_templates_pass(tmp_path):
    from app.services.pptx_renderer import render_pptx

    deck = DeckModel(title="Heat pumps in cold climates", slides=[
        SlideModel(role="title", title="Heat pumps in cold climates", subtitle="What works below -15 C"),
        SlideModel(role="title_content", title="Performance holds up", bullets=[
            BulletItem(text="Modern units keep useful output at low temperatures"),
            BulletItem(text="Defrost cycles cost a few percent of seasonal energy"),
            BulletItem(text="Sizing to the design temperature matters most"),
        ]),
    ])
    for tmpl in sorted(Path("data/sample_templates").glob("*.pptx")):
        out = tmp_path / f"{tmpl.stem}.pptx"
        render_pptx(deck, tmpl, None, out)
        errors = [i.message for i in validate_pptx_file(out, tmpl) if i.severity == "error"]
        assert errors == [], f"{tmpl.name}: {errors}"


def test_merge_report_recomputes_pass_and_score():
    base = {"passed": True, "score": 100.0, "issues": []}
    merged = merge_report(base, [ValidationIssue(severity="error", where="pptx:slide_2", message="x")])
    assert merged["passed"] is False and merged["score"] == 80.0


def test_one_shorten_round_when_text_does_not_fit(tmp_path):
    from app.agents import graph

    long_deck = DeckModel(title="T", slides=[SlideModel(role="title_content", title="Long", bullets=[BulletItem(text="a " * 40)])])
    short_deck = DeckModel(title="T", slides=[SlideModel(role="title_content", title="Short", bullets=[BulletItem(text="a b c")])])
    overflow = [ValidationIssue(severity="error", where="pptx:slide_1", message="Slide 1: text in 'content:body' overflows its box by 0.40 in.")]
    feedback = []

    def fake_generate(state, fb=""):
        feedback.append(fb)
        return short_deck

    state = {"validation": {"passed": True, "score": 100.0, "issues": []}}
    with patch.object(graph, "render_pptx") as render, \
         patch.object(graph, "validate_pptx_file", side_effect=[overflow, []]), \
         patch.object(graph, "_generate_deck", side_effect=fake_generate):
        result = graph._render_checked_deck(state, long_deck, GREEN, None, tmp_path / "d.pptx", {})
    assert result is short_deck and render.call_count == 2
    assert len(feedback) == 1 and "SHORTEN" in feedback[0] and "Slide 1" in feedback[0]
    assert state["validation"]["passed"] is True and state["deck_model"]["slides"][0]["title"] == "Short"

    # Still too long after shortening: no second round, the problem is reported.
    state = {"validation": {"passed": True, "score": 100.0, "issues": []}}
    with patch.object(graph, "render_pptx"), \
         patch.object(graph, "validate_pptx_file", side_effect=[overflow, overflow]), \
         patch.object(graph, "_generate_deck", return_value=short_deck) as gen:
        graph._render_checked_deck(state, long_deck, GREEN, None, tmp_path / "d.pptx", {})
    assert gen.call_count == 1 and state["validation"]["passed"] is False
