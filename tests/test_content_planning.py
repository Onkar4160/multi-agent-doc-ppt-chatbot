"""Stage 7 — content planning and synthesis (FR-PLAN-1..7, AC-5, AC-7).

Offline tests use the deterministic planners and writers; a scripted model stands in for the LLM to exercise the
real-mode paths (plan revision round, writer conformance, document length revision).
"""

from __future__ import annotations

from pathlib import Path

import pytest

from app.agents.doc_generator import SectionBatch, generate_document_model
from app.agents.planner import offline_deck_plan, offline_document_plan, plan_deck, plan_workbook
from app.agents.ppt_generator import generate_deck_model
from app.agents.offline_writer import write_workbook
from app.models.content_plan import DeckPlan, LengthTarget, SlidePlan
from app.models.deck_model import DeckModel, SlideModel
from app.models.document_model import ParagraphBlock, Section
from app.models.template_profile import TemplateProfile
from app.services.ir_compat import legacy_view
from app.services.length_rules import deck_length, document_length, within
from app.services.plan_inputs import ClaimRef
from app.services.plan_validation import (
    document_words, validate_deck_content, validate_deck_plan, validate_document_plan, validate_workbook_plan,
)


def _claims() -> list[ClaimRef]:
    return [
        ClaimRef(id="C1", statement="Community solar capacity reached 7.9 GW in 2024", source_ids=[1], question_ids=["Q1"], value=7.9, unit="GW", as_of="2024"),
        ClaimRef(id="C2", statement="Capacity was 5.6 GW in 2022", source_ids=[2], question_ids=["Q1"], value=5.6, unit="GW", as_of="2022"),
        ClaimRef(id="C3", statement="Subscribers save 10 percent on bills on average", source_ids=[1], question_ids=["Q2"], value=10, unit="%"),
        ClaimRef(id="C4", statement="Interconnection queues delay many projects", source_ids=[3], question_ids=["Q3"]),
        ClaimRef(id="C5", statement="Low-income carve-outs exist in several states", source_ids=[3], question_ids=["Q4"]),
        ClaimRef(id="C6", statement="Solar will be free by 2030", status="unsupported", source_ids=[], question_ids=["Q2"]),
    ]


QUESTIONS = {"Q1": "How large is the market?", "Q2": "What do subscribers save?", "Q3": "What are the barriers?",
             "Q4": "What policy supports exist?"}


class Scripted:
    """Stand-in for a real LLM: answers each schema from a queue."""
    _mock_mode = False

    def __init__(self, **queues):
        self.queues = {k: list(v) for k, v in queues.items()}
        self.calls: list[str] = []

    def generate_json(self, prompt, schema, **kwargs):
        self.calls.append(schema.__name__)
        q = self.queues.get(schema.__name__) or []
        if not q:
            raise RuntimeError(f"no scripted answer for {schema.__name__}")
        return q.pop(0)


@pytest.mark.parametrize("n", [1, 2, 3, 5, 8, 12, 15])
def test_offline_deck_plan_meets_exact_length_and_rules(n):
    claims = _claims()
    plan = offline_deck_plan("Create a deck on community solar", claims, LengthTarget(unit="slides", value=n, explicit=True),
                             questions=QUESTIONS)
    report = validate_deck_plan(plan, claims)
    assert len(plan.slides) == n  # AC-7: explicit slide counts exactly
    assert plan.slides[0].archetype == "title"
    assert not report.blocking, report.issues
    assert report.claim_coverage_pct == 100.0  # AC-5


def test_archetypes_follow_the_evidence():
    plan = offline_deck_plan("deck on community solar", _claims(), LengthTarget(unit="slides", value=8, explicit=True),
                             questions=QUESTIONS)
    by_claims = {tuple(s.claim_ids): s.archetype for s in plan.slides}
    assert by_claims[("C1", "C2")] == "chart"  # two figures, same unit
    assert all(s.archetype != "chart" for s in plan.slides if not any(c in ("C1", "C2") for c in s.claim_ids))


def test_repairs_fix_duplicates_misfits_unknown_and_unsupported_claims():
    claims = _claims()
    bad = DeckPlan(title="Solar", length=LengthTarget(unit="slides", value=5, explicit=True), slides=[
        SlidePlan(purpose="p", archetype="bullets", headline="Community solar grows fast", claim_ids=["C1"]),
        SlidePlan(purpose="p", archetype="chart", headline="Savings are real", claim_ids=["C3", "C9"]),
        SlidePlan(purpose="p", archetype="bullets", headline="Community solar grows fast", claim_ids=["C2"]),
        SlidePlan(purpose="p", archetype="bullets", headline="Free power soon", claim_ids=["C6"]),
    ])
    before = validate_deck_plan(bad, claims)
    codes = {i.code for i in before.issues}
    assert {"length", "duplicate", "unknown_claim", "unsupported_claim", "archetype_fit"} <= codes
    from app.agents.planner import repair_deck_plan

    repair_deck_plan(bad, claims, brief="deck on community solar")
    after = validate_deck_plan(bad, claims)
    assert not after.blocking, after.issues
    assert len(bad.slides) == 5 and bad.slides[0].archetype == "title"
    assert all(c not in ("C6", "C9") for s in bad.slides for c in s.claim_ids)


def test_llm_plan_gets_one_revision_round_then_repairs():
    claims = _claims()
    length = LengthTarget(unit="slides", value=4, explicit=True)
    first = DeckPlan(title="Solar", slides=[SlidePlan(purpose="p", archetype="bullets", headline="Only one slide")])
    revised = DeckPlan(title="Solar", slides=[
        SlidePlan(purpose="open", archetype="title", headline="Community solar"),
        SlidePlan(purpose="p", archetype="chart", headline="Capacity grew from 5.6 to 7.9 GW", claim_ids=["C1", "C2"], blocks=["chart"]),
        SlidePlan(purpose="p", archetype="kpi", headline="Subscribers save a tenth on bills", claim_ids=["C3"], blocks=["kpi"]),
        SlidePlan(purpose="close", archetype="closing", headline="Questions and next steps"),
    ])
    llm = Scripted(DeckPlan=[first, revised])
    plan, report = plan_deck("deck on community solar", claims, length, llm_client=llm)
    assert llm.calls == ["DeckPlan", "DeckPlan"]
    assert report.revised and not report.blocking
    assert [s.archetype for s in plan.slides] == ["title", "chart", "kpi", "closing"]


def test_planner_failure_falls_back_to_offline_plan():
    plan, report = plan_deck("deck on community solar", _claims(), LengthTarget(unit="slides", value=6, explicit=True),
                             llm_client=Scripted())
    assert plan.origin == "fallback" and len(plan.slides) == 6 and not report.blocking


def test_deck_writer_is_conformed_to_its_plan():
    claims = _claims()
    plan = offline_deck_plan("deck on community solar", claims, LengthTarget(unit="slides", value=6, explicit=True),
                             questions=QUESTIONS)
    # the model writes too few slides and leaves one empty
    written = DeckModel(title="Solar", slides=[SlideModel(role="title", title="Solar"),
                                               SlideModel(role="title_content", title="Empty one")])
    deck = generate_deck_model("deck on community solar", TemplateProfile(), plan=plan, claims=claims,
                               llm_client=Scripted(DeckModel=[written]))
    assert len(deck.slides) == 6
    assert [s.archetype for s in deck.slides] == [s.archetype for s in plan.slides]
    report = validate_deck_content(deck, plan)
    assert not report.blocking, report.issues


def test_length_targets_follow_request_not_template():
    claims = _claims()
    assert deck_length({"unit": "slides", "value": 7, "explicit": True}).value == 7
    auto = deck_length(None, "standard", claims)
    assert not auto.explicit and 7 <= auto.value <= 12
    pages = document_length({"unit": "pages", "value": 3, "explicit": True})
    assert pages.value == 1200 and pages.tolerance_pct == 15.0


@pytest.mark.parametrize("length", [{"unit": "pages", "value": 1, "explicit": True},
                                    {"unit": "words", "value": 2500, "explicit": True}])
def test_document_plan_allocates_explicit_length(length):
    target = document_length(length)
    plan = offline_document_plan("report on community solar", _claims(), target, document_type="research_report",
                                 questions=QUESTIONS)
    total = sum(s.target_words for s in plan.sections) + (plan.summary_words if plan.executive_summary else 0)
    assert abs(total - target.value) <= target.value * 0.05  # AC-7 at plan level
    assert not validate_document_plan(plan, _claims()).blocking


def test_off_brief_sections_are_dropped_for_reports():
    from app.agents.planner import repair_document_plan
    from app.models.content_plan import DocumentPlan, SectionPlan

    plan = DocumentPlan(title="t", length=LengthTarget(unit="words", value=900), sections=[
        SectionPlan(heading="Background", purpose="p"), SectionPlan(heading="Pricing and fees", purpose="p"),
        SectionPlan(heading="Outlook", purpose="p")])
    repair_document_plan(plan, [], brief="report on community solar", document_type="research_report")
    assert [s.heading for s in plan.sections] == ["Background", "Outlook"]


def test_document_length_revision_round_hits_explicit_target():
    target = document_length({"unit": "words", "value": 600, "explicit": True})
    claims = _claims()
    plan = offline_document_plan("report on community solar", claims, target, questions=QUESTIONS)
    short = SectionBatch(sections=[Section(heading=s.heading, blocks=[ParagraphBlock(text="Too short.")]) for s in plan.sections],
                         executive_summary="Summary of the findings.")
    resized = SectionBatch(sections=[Section(heading="x", blocks=[ParagraphBlock(text=" ".join(["word"] * s.target_words))])
                                     for s in plan.sections])
    llm = Scripted(SectionBatch=[short, resized])
    doc = generate_document_model("report on community solar", TemplateProfile(), plan=plan, claims=claims, llm_client=llm)
    assert llm.calls == ["SectionBatch", "SectionBatch"]
    assert [s.heading for s in doc.sections] == [s.heading for s in plan.sections]
    assert within(document_words(doc), target), document_words(doc)


def test_workbook_plan_roles_checks_and_explicit_sheets():
    claims = _claims()
    plan, report = plan_workbook("model community solar capacity", claims, None)
    roles = [s.role for s in plan.sheets]
    assert roles[0] == "readme" and "data" in roles and "calculations" in roles and roles[-1] == "sources"
    assert not report.blocking, report.issues
    wb = write_workbook(plan, claims)
    calc = next(s for s in wb.sheets if s.role == "calculations")
    assert any(c.formula and c.formula.startswith("=SUMIF(Data!") for c in calc.cells)
    assert wb.checks and wb.checks[0].expected == pytest.approx(7.9 + 5.6)
    three, _ = plan_workbook("model", claims, LengthTarget(unit="sheets", value=3, explicit=True))
    assert len(three.sheets) == 3 and any(s.role == "data" for s in three.sheets)
    assert not validate_workbook_plan(three, claims).blocking


def test_typed_blocks_render_without_loss(tmp_path: Path):
    from app.services.pptx_renderer import render_pptx

    claims = _claims()
    plan = offline_deck_plan("deck on community solar", claims, LengthTarget(unit="slides", value=6, explicit=True),
                             questions=QUESTIONS)
    deck = generate_deck_model("deck on community solar", TemplateProfile(), plan=plan, claims=claims)
    chart_slide = next(s for s in legacy_view(deck).slides if s.archetype == "chart")
    assert any("7.9" in b.text for b in chart_slide.bullets)
    out = render_pptx(deck, Path("data/sample_templates/presentation_Template.pptx"), out_path=tmp_path / "d.pptx")
    import pptx

    text = " ".join(sh.text_frame.text for sl in pptx.Presentation(str(out)).slides for sh in sl.shapes if sh.has_text_frame)
    assert "7.9" in text


def test_plan_content_node_plans_every_output():
    from app.agents.graph import node_plan_content

    state = {"user_message": "Create a 5-slide deck and a report on community solar",
             "plan": {"outputs": ["pptx", "docx", "xlsx"], "topic": "community solar", "document_type": "research_report"},
             "claims": [c.model_dump() for c in _claims()], "questions": QUESTIONS,
             "interpretation": {"output_spec": {"length": {"unit": "slides", "value": 5, "explicit": True}}},
             "registry": {1: {}, 2: {}, 3: {}}, "template_profiles": {}}
    out = node_plan_content(state)
    assert len(out["content_plans"]["pptx"]["slides"]) == 5 and out["plan"]["slide_count"] == 5
    assert out["content_plans"]["docx"]["sections"] and out["content_plans"]["xlsx"]["sheets"]
    assert all(r["claim_coverage_pct"] == 100.0 for r in out["plan_reports"].values())  # AC-5
