"""Templates whose design lives on the sample slides keep that design (clone rendering, ported from DocForge)."""

import re
from pathlib import Path

import pptx
import pytest
from pptx.enum.shapes import MSO_SHAPE_TYPE

from app.models.deck_model import BulletItem, DeckModel, SlideModel
from app.models.template_profile import PptStyleProfile, TemplateProfile
from app.services.analyzers.slide_design import all_text, analyze_slide_design
from app.services.layout.backdrop import color_under
from app.services.layout.colors import contrast_ratio, darken_until, tint
from app.services.layout.text import Para, fit_paragraphs, measure
from app.services.pptx_renderer import render_pptx

TEMPLATES = sorted(Path("data/sample_templates").glob("*.pptx"))
GREEN = Path("data/sample_templates/Green Cream Simple Aesthetic Watercolor Presentation.pptx")


def _deck() -> DeckModel:
    return DeckModel(
        title="Electric Vehicle Charging Infrastructure in India",
        subtitle="Market outlook and strategy 2026",
        slides=[
            SlideModel(role="title", title="Electric Vehicle Charging Infrastructure in India", subtitle="Market outlook and strategy 2026"),
            SlideModel(role="title_content", title="Agenda", bullets=[BulletItem(text=t) for t in ["Market size and growth", "Policy landscape", "Key players"]]),
            SlideModel(role="section_header", title="Market Overview", subtitle="Where the opportunity lies"),
            SlideModel(role="title_content", title="Charging demand is growing faster than supply across metro corridors", bullets=[
                BulletItem(text="Public chargers grew 3x between 2023 and 2025, but EV sales grew 5x", source_ids=[1]),
                BulletItem(text="Highway corridors remain underserved outside the main metro routes", source_ids=[2]),
                BulletItem(text="Fleet operators drive most fast-charging utilisation today"),
                BulletItem(text="Two- and three-wheelers account for over 90% of EVs", level=1),
                BulletItem(text="Grid upgrades are the main lead-time constraint for new sites"),
            ]),
            SlideModel(role="two_content", title="Build vs Partner",
                       left=[BulletItem(text="Full control of site quality"), BulletItem(text="Higher capex")],
                       right=[BulletItem(text="Faster rollout via fuel stations"), BulletItem(text="Revenue share lowers margin")]),
            SlideModel(role="title_only", title="Thank You", subtitle="Questions and discussion"),
        ],
    )


SOURCES = {1: {"title": "NITI Aayog EV report", "url": "https://niti.gov.in/ev"}, 2: {"title": "BEE guidelines", "url": "https://beeindia.gov.in"}}


@pytest.mark.parametrize("template", TEMPLATES, ids=lambda p: p.stem[:12])
def test_sample_templates_keep_their_design(template: Path, tmp_path: Path):
    """Artwork kept on every slide, sample text gone, text fits its box and stays on the slide."""
    src = pptx.Presentation(template)
    design = analyze_slide_design(src)
    assert design.design_mode == "slides"
    assert any(b.kind == "cover" for b in design.slide_bases)

    deck = _deck()
    ours = " ".join([deck.title, deck.subtitle] + [x for s in deck.slides for x in [s.title, s.subtitle] + [b.text for b in s.bullets + s.left + s.right]]).lower()
    sample_words = {
        w.lower() for sl in src.slides for sh in sl.shapes for w in re.findall(r"[A-Za-z]{6,}", all_text(sh)) if w.lower() not in ours
    }

    out = tmp_path / "deck.pptx"
    render_pptx(deck, template, None, out, sources_map=SOURCES)
    prs = pptx.Presentation(out)
    assert len(prs.slides) == len(deck.slides) + 1  # + Sources slide

    for n, slide in enumerate(prs.slides, 1):
        template_shapes = [sh for sh in slide.shapes if not sh.name.startswith(("content:", "decor:"))]
        has_bg = slide._element.cSld.bg is not None or slide.slide_layout._element.cSld.bg is not None
        assert template_shapes or has_bg, f"slide {n} lost the template artwork"
        text = " ".join(all_text(sh) for sh in slide.shapes)
        residue = {w for w in re.findall(r"[A-Za-z]{6,}", text.lower()) if w in sample_words}
        assert not residue, f"slide {n} still shows sample text {residue}"
        assert "lorem" not in text.lower() and "(continued)" not in text.lower()
        for sh in slide.shapes:
            if not sh.name.startswith("content:") or not sh.has_text_frame:
                continue
            assert int(sh.left) >= 0 and int(sh.top) >= 0, f"slide {n} {sh.name} off-slide"
            assert int(sh.left) + int(sh.width) <= prs.slide_width and int(sh.top) + int(sh.height) <= prs.slide_height
            tf = sh.text_frame
            paras = [Para(p.text, p.runs[0].font.size.pt, font=p.runs[0].font.name) for p in tf.paragraphs if p.runs and p.runs[0].font.size]
            # an explicit zero inset is a real inset (templates often set lIns="0"), not a missing one
            width = int(sh.width) - int(tf.margin_left) - int(tf.margin_right)
            assert measure(paras, width).height_emu <= int(sh.height) * 1.02, f"slide {n} {sh.name} overflows"


def test_green_template_artwork_and_readable_titles(tmp_path: Path):
    out = tmp_path / "green.pptx"
    render_pptx(_deck(), GREEN, None, out, sources_map=SOURCES)
    prs = pptx.Presentation(out)
    for slide in prs.slides:
        pictures = [sh for sh in slide.shapes if sh.shape_type == MSO_SHAPE_TYPE.PICTURE or sh._element.xpath(".//a:blip")]
        assert pictures, "watercolour artwork must be kept on every slide"
        title = next(sh for sh in slide.shapes if sh.name in ("content:title", "content:headline"))
        run = title.text_frame.paragraphs[0].runs[0]
        if run.font.color.type is not None:
            under = color_under(slide, (int(title.left), int(title.top), int(title.width), int(title.height)), stop_at=title._element)
            assert contrast_ratio(str(run.font.color.rgb), under) >= 3


def test_old_cached_profile_without_slide_design_still_clones(tmp_path: Path):
    """Profiles cached before slide analysis existed have no slide_design: it is computed at render time."""
    profile = TemplateProfile(file_type="pptx", ppt_style=PptStyleProfile())
    out = tmp_path / "cached.pptx"
    render_pptx(_deck(), GREEN, profile, out)
    prs = pptx.Presentation(out)
    assert len(prs.slides) == len(_deck().slides) + 1  # + Sources slide
    assert all(sh.name != "AutoDecoration_AccentBar" for s in prs.slides for sh in s.shapes)


def test_template_without_sample_slides_uses_layouts(tmp_path: Path):
    tmpl = pptx.Presentation(GREEN)
    for sld_id in list(tmpl.slides._sldIdLst):
        tmpl.part.drop_rel(sld_id.rId)
        tmpl.slides._sldIdLst.remove(sld_id)
    path = tmp_path / "layouts_only.pptx"
    tmpl.save(path)
    assert analyze_slide_design(pptx.Presentation(path)).design_mode == "layouts"
    out = tmp_path / "out.pptx"
    render_pptx(_deck(), path, None, out)
    assert len(pptx.Presentation(out).slides) == len(_deck().slides) + 1


def test_colour_and_fit_helpers():
    assert tint("000000", 1.0) == "FFFFFF"
    assert contrast_ratio(darken_until("F2C94C", "FFFFFF", 4.5), "FFFFFF") >= 4.5
    paras = [Para("word " * 80, 28)]
    fitted, ok = fit_paragraphs(paras, 914400 * 4, 914400 * 2, min_pt=12)
    assert ok and fitted[0].size_pt < 28
