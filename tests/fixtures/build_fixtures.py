"""Programmatic, licence-clean template fixtures (TESTING §4). Built at test time so they are reproducible."""

from __future__ import annotations

import copy
import io
import shutil
import zipfile
from pathlib import Path

TEMPLATE_CONTENT_TYPES = {
    ".potx": ("presentationml.presentation.main+xml", "presentationml.template.main+xml"),
    ".dotx": ("wordprocessingml.document.main+xml", "wordprocessingml.template.main+xml"),
    ".xltx": ("spreadsheetml.sheet.main+xml", "spreadsheetml.template.main+xml"),
}


def build_xlsx_template(path: Path) -> Path:
    """A small financial-model template: readme, assumptions with names, data table, calculations, dashboard, sources."""
    from openpyxl import Workbook
    from openpyxl.chart import BarChart, Reference
    from openpyxl.formatting.rule import CellIsRule
    from openpyxl.styles import Font, PatternFill
    from openpyxl.workbook.defined_name import DefinedName
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table

    wb = Workbook()
    readme = wb.active
    readme.title = "README"
    readme["A1"] = "Revenue model template"
    readme["A2"] = "Fill the Assumptions sheet; everything else is calculated."

    a = wb.create_sheet("Assumptions")
    a.append(["Driver", "Value", "Unit"])
    a.append(["Growth rate", 0.08, "% per year"])
    a.append(["Churn", 0.02, "% per month"])
    for c in a[1]:
        c.font = Font(name="Arial", bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E79")
    a["B2"].number_format = "0.0%"
    a["B3"].number_format = "0.0%"
    wb.defined_names["growth"] = DefinedName("growth", attr_text="Assumptions!$B$2")
    wb.defined_names["churn"] = DefinedName("churn", attr_text="Assumptions!$B$3")
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="1")
    a.add_data_validation(dv)
    dv.add("B2:B3")

    d = wb.create_sheet("Data")
    d.append(["Month", "Customers", "Price"])
    for m in range(1, 7):
        d.append([f"2026-{m:02d}", 100 + m * 10, 49])
    d["C2"].number_format = '"$"#,##0'
    d.add_table(Table(displayName="Customers", ref="A1:C7"))
    d.freeze_panes = "A2"

    c = wb.create_sheet("Calculations")
    c.append(["Month", "Revenue", "Revenue next year"])
    for r in range(2, 8):
        c.append([f"=Data!A{r}", f"=Data!B{r}*Data!C{r}", f"=B{r}*(1+Assumptions!$B$2)"])
        c[f"B{r}"].number_format = '"$"#,##0'
    c.conditional_formatting.add("B2:B7", CellIsRule(operator="greaterThan", formula=["6000"], fill=PatternFill("solid", fgColor="C6EFCE")))

    dash = wb.create_sheet("Dashboard")
    dash["A1"] = "Total revenue"
    dash["B1"] = "=SUM(Calculations!B2:B7)"
    chart = BarChart()
    chart.title = "Revenue by month"
    chart.add_data(Reference(c, min_col=2, min_row=1, max_row=7), titles_from_data=True)
    dash.add_chart(chart, "D2")

    s = wb.create_sheet("Sources")
    s.append(["Figure", "Source"])
    s.append(["Price", "Internal price list 2026"])
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def as_template_package(src: Path, dest: Path) -> Path:
    """Copy a regular .pptx/.docx/.xlsx as its template variant (.potx/.dotx/.xltx) by switching the content type."""
    regular, template = TEMPLATE_CONTENT_TYPES[dest.suffix.lower()]
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(dest, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(regular.encode(), template.encode())
            zout.writestr(item, data)
    return dest


def build_pptx_with_master_logo(path: Path) -> Path:
    """A layouts-mode deck whose logo sits on the slide master (branding shared by every slide)."""
    from PIL import Image
    from pptx import Presentation
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT
    from pptx.oxml.ns import qn
    from pptx.util import Inches

    prs = Presentation()
    buf = io.BytesIO()
    Image.new("RGB", (120, 60), (200, 30, 30)).save(buf, format="PNG")
    buf.seek(0)
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Quarterly review"
    # python-pptx cannot add pictures to a master: add it to the slide, then move the element and its image part.
    pic = slide.shapes.add_picture(buf, Inches(8.8), Inches(0.2), Inches(0.9), Inches(0.45))
    master = prs.slide_master
    r_id = master.part.relate_to(slide.part.related_part(pic._element.blip_rId), RT.IMAGE)
    el = copy.deepcopy(pic._element)
    el.blipFill.find(qn("a:blip")).set(qn("r:embed"), r_id)
    master.shapes._spTree.append(el)
    slide.shapes._spTree.remove(pic._element)
    path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(path)
    return path


def copy_sample(name: str, dest_dir: Path) -> Path:
    src = Path("data/sample_templates") / name
    dest = dest_dir / name
    shutil.copyfile(src, dest)
    return dest
