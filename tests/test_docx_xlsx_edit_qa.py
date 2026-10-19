"""Stages 9–12: Word finishing, Excel rendering and recalculation, restyle and edit regression, quality reports and the
acceptance evaluation (FR-DOCX-1..5, FR-XLSX-1..7, FR-EDIT-3, FR-VAL-1..4, AC-2..AC-10)."""

from __future__ import annotations

from pathlib import Path

import docx
import openpyxl
import pytest

from app.models.document_model import DocumentModel, ParagraphBlock, Section
from app.models.style import StyleOverrides
from app.models.workbook_model import CellModel, ChartModel, ExpectedCheck, NamedRange, SheetModel, TableModel, WorkbookModel
from app.services.docx_renderer import render_docx
from app.services.xlsx_check import recalculate
from app.services.xlsx_renderer import render_xlsx

TEMPLATE = Path("data/sample_templates/proposal_Template.docx")
SOURCES = {1: {"title": "Solar report", "url": "https://a.org/r", "publisher": "Energy Agency", "published_at": "2024-03-01"},
           2: {"title": "Market data", "url": "https://b.org/m", "authors": ["Ann Lee", "Raj Patel"], "published_at": "2023"}}


def _doc(**kw) -> DocumentModel:
    return DocumentModel(title="Community solar", sections=[
        Section(heading="Background", blocks=[ParagraphBlock(text="Capacity reached 7.9 GW.", source_ids=[1, 2])]),
        Section(heading="Outlook", blocks=[ParagraphBlock(text="Growth continues.")])], **kw)


def _xml(path: Path) -> str:
    d = docx.Document(str(path))
    return d.element.xml + "".join(s.footer._element.xml for s in d.sections)


def test_docx_toc_page_numbers_and_apa_citations(tmp_path):
    out = render_docx(_doc(table_of_contents=True, citation_style="apa"), TEMPLATE, out_path=tmp_path / "d.docx", sources_map=SOURCES)
    xml = _xml(out)
    assert 'TOC \\o "1-3"' in xml and "PAGE" in xml  # FR-DOCX-1, FR-DOCX-3
    text = "\n".join(p.text for p in docx.Document(str(out)).paragraphs)
    assert "(Energy Agency, 2024; Lee and Patel, 2023)" in text  # FR-DOCX-2 in-text APA
    assert "Energy Agency. (2024). Solar report. https://a.org/r" in text
    assert "[1]" not in text


def test_harvard_reference_entry():
    from app.services.docx_finish import in_text, reference_entry

    assert in_text([1], "harvard", SOURCES) == " (Energy Agency 2024)"
    assert reference_entry(2, "harvard", SOURCES).startswith("Lee and Patel (2023) Market data. Available at: https://b.org/m")


def test_style_overrides_go_into_named_styles(tmp_path):
    style = StyleOverrides()
    style.apply([{"path": "heading_1.size_pt", "value": 20}, {"path": "accent_color", "value": "red"},
                 {"path": "body.font", "value": "Georgia"}])
    out = render_docx(_doc(style=style), TEMPLATE, out_path=tmp_path / "s.docx")
    d = docx.Document(str(out))
    assert d.styles["Heading 1"].font.size.pt == 20 and str(d.styles["Heading 1"].font.color.rgb) == "C0392B"
    assert d.styles["Normal"].font.name == "Georgia"
    heading = next(p for p in d.paragraphs if p.text == "Background")
    assert all(r.font.size is None for r in heading.runs)  # FR-DOCX-5: no direct formatting


def _workbook(**kw) -> WorkbookModel:
    data = SheetModel(name="Data", role="data", tables=[TableModel(name="Capacity", range="A1:C4", headers=["Year", "Capacity (GW)", "Share %"],
                                                                  rows=[[2022, 5.6, 0.2], [2023, 6.8, 0.25], [2024, 7.9, 0.3]])])
    calc = SheetModel(name="Calculations", role="calculations", cells=[
        CellModel(coordinate="A1", value="Total"), CellModel(coordinate="B1", formula="=SUM(Data!B2:B4)"),
        CellModel(coordinate="A2", value="Growth"), CellModel(coordinate="B2", formula="=(Data!B4-Data!B2)/Data!B2*growth_factor")])
    assum = SheetModel(name="Assumptions", role="assumptions", tables=[TableModel(name="Inputs", range="A1:B2", headers=["Driver", "Value"], rows=[["factor", 1]])])
    report = SheetModel(name="Report", role="report", charts=[ChartModel(title="Capacity", chart_type="column",
                                                                         data_range="Data!$B$1:$B$4", categories_range="Data!$A$2:$A$4")])
    return WorkbookModel(title="Solar", sheets=[assum, data, calc, report], named_ranges=[NamedRange(name="growth_factor", ref="Assumptions!B2")],
                         checks=[ExpectedCheck(ref="Calculations!B1", expected=20.3)], **kw)


def test_xlsx_tables_formats_names_charts_and_recalculation(tmp_path):
    out = render_xlsx(_workbook(), tmp_path / "w.xlsx")
    wb = openpyxl.load_workbook(str(out))
    data = wb["Data"]
    assert "Capacity" in data.tables and data.freeze_panes == "A2"  # FR-XLSX-2
    assert data["C2"].number_format == "0.0%" and data["A2"].number_format == "0"
    assert "growth_factor" in wb.defined_names  # FR-XLSX-3
    assert wb["Report"]._charts and data.conditional_formatting  # FR-XLSX-4, FR-XLSX-5
    report = recalculate(out, _workbook().checks)
    assert report.errors == [] and report.formulas == 2 and report.evaluated == 2  # FR-XLSX-6, AC-4
    assert report.checks[0].ok and report.checks[0].actual == pytest.approx(20.3)


@pytest.mark.parametrize("formula,code", [("=Missing!A1", "#REF!"), ("=NOPE(1)", "#NAME?"), ("=1/0", "#DIV/0!"),
                                          ("=undefined_name*2", "#NAME?"), ("=Calculations!C1", "circular")])
def test_recalculation_reports_formula_errors(tmp_path, formula, code):
    wb = WorkbookModel(sheets=[SheetModel(name="Calculations", cells=[CellModel(coordinate="C1", formula=formula)])])
    out = render_xlsx(wb, tmp_path / "e.xlsx")
    report = recalculate(out)
    assert report.errors and code in report.errors[0]


def test_failed_expected_value_check_is_an_error(tmp_path):
    from app.services.quality_report import validate_xlsx_file

    wb = _workbook()
    wb.checks = [ExpectedCheck(ref="Calculations!B1", expected=99)]
    issues, _ = validate_xlsx_file(render_xlsx(wb, tmp_path / "w.xlsx"), wb.checks)
    assert any("Expected 99" in i.message for i in issues)


def test_restyle_is_pure_and_changes_no_content_unit():
    from app.agents.interpreter import parse_style_changes
    from app.services.edit_guard import is_pure_restyle, regression

    text = "Make the headings 24 pt"
    changes = parse_style_changes(text)
    assert is_pure_restyle(text, changes)
    assert not is_pure_restyle("Add a section on costs and make headings 24 pt", changes)
    old = _doc()
    new = old.model_copy(deep=True)
    new.style = StyleOverrides()
    new.style.apply(changes)
    assert regression(old, new, []).changed_units == []


def test_regression_guard_flags_out_of_scope_changes(tmp_path):
    from app.services.artifact_nodes import ensure_node_ids
    from app.services.edit_guard import regression

    old = ensure_node_ids(_doc())
    new = old.model_copy(deep=True)
    new.sections[1].blocks[0].text = "Changed outlook."
    assert regression(old, new, [old.sections[1].node_id]).ok
    bad = regression(old, new, [old.sections[0].node_id])
    assert not bad.ok and bad.out_of_scope_changed == [old.sections[1].node_id]
    a = render_docx(old, TEMPLATE, out_path=tmp_path / "a.docx")
    b = render_docx(new, TEMPLATE, out_path=tmp_path / "b.docx")
    files = regression(old, new, None, old_file=a, new_file=b, kind="docx")
    assert files.file_units_compared >= 2 and len(files.file_changed_units) == 1


def test_quality_report_separates_findings_and_not_performed(tmp_path):
    from app.agents.validator import ValidationIssue
    from app.services.quality_report import build_quality_report

    out = render_docx(_doc(), TEMPLATE, out_path=tmp_path / "q.docx")
    state = {"plan_reports": {"docx": {"repairs": ["dropped duplicate 'Outlook'"], "revised": True, "claim_coverage_pct": 100.0}}}
    q = build_quality_report("docx", out, state, file_issues=[ValidationIssue(severity="warning", where="docx:file", message="Empty table.")],
                             offline=True)
    assert q.counts["corrected"] == 2 and q.counts["unresolved"] == 1
    ids = {n.check_id for n in q.not_performed}
    assert {"docx.visual", "content.llm_review"} <= ids


def test_acceptance_evaluation_passes(tmp_path):
    from app.eval.acceptance import evaluate

    report = evaluate(tmp_path)
    m = report["metrics"]
    assert report["passed"], m
    assert m["AC-2_reopen_pct"] == 100 and m["AC-3_pptx_unresolved_critical"] == 0 and m["AC-4_formula_errors"] == 0
    assert m["AC-7_length_ok_pct"] == 100 and m["AC-10_quality_report_pct"] == 100
