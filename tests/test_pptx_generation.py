"""Stage 8 — PowerPoint generation (FR-PPTX-1..8, AC-3, AC-8).

Every archetype is rendered on the three sample templates and on the generated (no-template) design, then the file
is reopened and checked: native charts and tables with the planned data, zero critical file findings, no template
sample text, accessible chart descriptions. Preview PNGs guard against blank or broken slides.
"""

from __future__ import annotations

from pathlib import Path

import pptx
import pytest

from app.agents.ppt_analyzer import analyze_presentation
from app.agents.validator import validate_pptx_file
from app.services.analyzers.summary import NoLLM
from app.services.pptx_renderer import render_pptx
from app.services.renderers.pptx_archetypes import Palette, palette_from_profile
from app.services.renderers.pptx_generated import preset_for
from tests.fixtures.archetype_deck import archetype_deck

TEMPLATES = [None, "Company_Template.pptx", "Green Cream Simple Aesthetic Watercolor Presentation.pptx",
             "presentation_Template.pptx"]
SOURCES = {i: {"title": f"Source {i}", "url": f"https://example.org/{i}"} for i in range(1, 5)}


def _render(tmp_path: Path, name: str | None, hint: str = "") -> tuple[Path, Path | None]:
    tpl = Path("data/sample_templates") / name if name else None
    profile = analyze_presentation(tpl, llm_client=NoLLM()) if tpl else None
    out = tmp_path / "deck.pptx"
    render_pptx(archetype_deck(), tpl, profile, out, sources_map=SOURCES, design_hint=hint)
    return out, tpl


@pytest.mark.parametrize("name", TEMPLATES)
def test_every_archetype_renders_natively_with_zero_critical_findings(tmp_path, name):
    out, tpl = _render(tmp_path, name)
    prs = pptx.Presentation(str(out))
    deck = archetype_deck()
    assert len(prs.slides) == len(deck.slides) + 1  # + Sources slide (FR-PPTX-6, FR-PPTX-7)
    charts = [sh for s in prs.slides for sh in s.shapes if sh.has_chart]
    tables = [sh for s in prs.slides for sh in s.shapes if sh.has_table]
    assert len(charts) == 1 and len(tables) == 1  # FR-PPTX-1/4: native, editable
    values = list(charts[0].chart.plots[0].series[0].values)
    assert values == [5.6, 6.8, 7.9]
    assert "column chart" in charts[0]._element.xpath("./*[1]/p:cNvPr")[0].get("descr")
    assert tables[0].table.cell(1, 0).text == "Utility-led"
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    for expected in ("10 %", "2018", "Partner", "Subscription", "People want to see"):
        assert expected in text, expected
    errors = [i.message for i in validate_pptx_file(out, tpl) if i.severity == "error"]
    assert errors == []  # AC-3 (overflow, off-slide, sample text, empty placeholders) and AC-8 sample text


def test_template_sample_charts_do_not_leak_into_visual_slides(tmp_path):
    out, _ = _render(tmp_path, "presentation_Template.pptx")
    prs = pptx.Presentation(str(out))
    text = " ".join(sh.text_frame.text for s in prs.slides for sh in s.shapes if sh.has_text_frame)
    assert "67%" not in text
    assert sum(1 for s in prs.slides for sh in s.shapes if sh.has_chart) == 1


def test_generated_design_follows_tone_and_is_16_by_9(tmp_path):
    out, _ = _render(tmp_path, None, hint="friendly tone for residents")
    prs = pptx.Presentation(str(out))
    assert round(prs.slide_width / prs.slide_height, 2) == round(16 / 9, 2)
    accent = preset_for("friendly")[0].accent
    xml = b"".join(s.part.blob for s in prs.slides)
    assert accent.encode() in xml
    assert preset_for("board update")[0] != preset_for("pitch to investors")[0]


def test_palette_from_template_tokens():
    profile = analyze_presentation(Path("data/sample_templates/Company_Template.pptx"), llm_client=NoLLM())
    pal = palette_from_profile(profile)
    assert pal.accent == profile.design_tokens.primary.upper()
    assert palette_from_profile(None) == Palette()


def test_no_template_phrases_select_generated_design():
    from app.agents.graph import _no_template
    from app.agents.supervisor import parse_plan

    plan = parse_plan("Create a 6-slide deck on heat pumps without a template")
    assert str(plan.template_mode).upper().endswith("NONE")
    assert _no_template({"plan": {"template_mode": plan.template_mode}})


@pytest.mark.parametrize("name", [None, "Company_Template.pptx"])
def test_preview_pngs_are_not_blank(tmp_path, name):
    from PIL import Image, ImageStat

    from app.services.renderers.pptx_preview import render_previews

    out, _ = _render(tmp_path, name)
    paths = render_previews(out, tmp_path / "png", width_px=320)
    assert len(paths) == len(archetype_deck().slides) + 1
    for p in paths:
        stat = ImageStat.Stat(Image.open(p).convert("L"))
        assert stat.stddev[0] > 3, f"{p.name} looks blank"
