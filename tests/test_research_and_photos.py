"""Scholarly sources, template photo frames and acronym clarification (Step 4 of the port). No network."""

import io
from pathlib import Path
from unittest.mock import patch

import httpx
import pptx
import pytest
from fastapi.testclient import TestClient
from PIL import Image

from app.models.deck_model import BulletItem, DeckModel, SlideModel
from app.services import academic_search as acad
from app.services.images import OpenverseProvider, fetch_deck_photos, resolve_photo
from app.services.pptx_renderer import render_pptx

COMPANY = Path("data/sample_templates/Company_Template.pptx")

OPENALEX = {"results": [{
    "id": "https://openalex.org/W1", "doi": "https://doi.org/10.1/abc", "display_name": "Cold climate heat pumps",
    "publication_year": 2023, "publication_date": "2023-05-01", "cited_by_count": 40, "type": "article",
    "primary_location": {"source": {"display_name": "Energy and Buildings"}},
    "authorships": [{"author": {"display_name": "A. Author"}}],
    "abstract_inverted_index": {"Heat": [0], "pumps": [1], "work": [2], "cold.": [3]},
}]}
S2 = {"data": [{"title": "Defrost losses", "abstract": "Defrost costs energy.", "year": 2022, "venue": "Applied Energy",
                "authors": [{"name": "B. Author"}], "citationCount": 5, "url": "https://s2.org/p", "externalIds": {}}]}
ARXIV = b"""<?xml version="1.0"?><feed xmlns="http://www.w3.org/2005/Atom" xmlns:arxiv="http://arxiv.org/schemas/atom">
<entry><id>http://arxiv.org/abs/2401.1</id><title>Frost models
 for heat pumps</title><summary>We model frost.</summary><published>2024-01-02T00:00:00Z</published>
<author><name>C. Author</name></author></entry></feed>"""


def _client(handler) -> httpx.Client:
    return httpx.Client(transport=httpx.MockTransport(handler))


def test_scholarly_providers_parse_their_apis():
    oa = acad.search_openalex("heat pumps", client=_client(lambda r: httpx.Response(200, json=OPENALEX)))
    assert oa[0]["url"] == "https://doi.org/10.1/abc" and oa[0]["snippet"] == "Heat pumps work cold."
    assert acad.scholarly_title(oa[0]) == "Cold climate heat pumps (Energy and Buildings, 2023)"
    s2 = acad.search_semantic_scholar("heat pumps", client=_client(lambda r: httpx.Response(200, json=S2)))
    assert s2[0]["venue"] == "Applied Energy" and s2[0]["url"] == "https://s2.org/p"
    ax = acad.search_arxiv("heat pumps", client=_client(lambda r: httpx.Response(200, content=ARXIV)))
    assert ax[0]["title"] == "Frost models for heat pumps" and ax[0]["kind"] == "preprint"
    with pytest.raises(acad.AcademicSearchError):
        acad.search_openalex("x", client=_client(lambda r: httpx.Response(429)))


def test_scholarly_search_survives_a_failing_index_and_dedupes():
    paper = {"title": "Same Paper", "url": "u", "snippet": "s"}
    with patch.object(acad, "search_openalex", return_value=[paper]), \
         patch.object(acad, "search_semantic_scholar", side_effect=acad.AcademicSearchError("rate limit")), \
         patch.object(acad, "search_arxiv", return_value=[{**paper, "title": "same paper"}]):
        assert len(acad.search_scholarly("q")) == 1


def test_research_registers_papers_as_citable_sources():
    from app.agents import web_research
    from app.services.source_registry import SourceRegistry

    paper = {"title": "Cold climate heat pumps", "url": "https://doi.org/10.1/abc", "snippet": "Heat pumps work in the cold.",
             "published_date": "2023-05-01", "venue": "Energy and Buildings", "kind": "paper"}
    registry = SourceRegistry()
    with patch.object(web_research, "_search_web_single_query", return_value=[]), \
         patch.object(web_research, "_search_scholarly_cached", return_value=[paper]):
        result = web_research.research("heat pumps in cold climates", registry=registry)
    assert len(result.source_ids) == 1
    src = registry.get(result.source_ids[0])
    assert src["title"] == "Cold climate heat pumps (Energy and Buildings, 2023)"
    assert src["url"] == "https://doi.org/10.1/abc"


def _png(width=1200, height=700) -> bytes:
    buf = io.BytesIO()
    Image.new("RGB", (width, height), (40, 120, 200)).save(buf, format="PNG")
    return buf.getvalue()


def _openverse(results_for: dict[str, int]):
    """Openverse stand-in: ``results_for`` maps a query to how many photos it finds."""
    def handler(request: httpx.Request) -> httpx.Response:
        if request.url.host == "api.openverse.org":
            q = request.url.params["q"]
            n = results_for.get(q, 0)
            return httpx.Response(200, json={"results": [
                {"url": f"https://img.example/{q.replace(' ', '_')}/{i}.png", "width": 1800, "height": 1000,
                 "creator": f"Creator {i}", "provider": "flickr", "license": "by", "license_version": "2.0",
                 "foreign_landing_url": f"https://flickr.example/{i}", "title": q}
                for i in range(n)
            ]})
        return httpx.Response(200, content=_png())
    return OpenverseProvider(client=_client(handler))


def test_photos_are_never_repeated_and_long_queries_fall_back(tmp_path):
    provider = _openverse({"Electric Vehicle Charging": 2})
    used, cache = set(), {}
    first = resolve_photo("Electric Vehicle Charging Infrastructure in India", provider, tmp_path, used=used, cache=cache)
    second = resolve_photo("Electric Vehicle Charging Infrastructure in India", provider, tmp_path, used=used, cache=cache)
    third = resolve_photo("Electric Vehicle Charging Infrastructure in India", provider, tmp_path, used=used, cache=cache)
    assert first and second and third is None
    assert first.path != second.path and len(used) == 2
    assert first.credit == "Photo: Creator 0 (flickr), CC BY 2.0"


def _deck() -> DeckModel:
    return DeckModel(title="Electric Vehicle Charging", slides=[
        SlideModel(role="title", title="Electric Vehicle Charging", subtitle="Outlook 2026"),
        SlideModel(role="title_content", title="Agenda", bullets=[BulletItem(text="Market"), BulletItem(text="Policy")]),
        SlideModel(role="title_content", title="Demand", bullets=[BulletItem(text="Chargers grew fast")]),
    ])


def test_template_photo_frames_get_credited_photos(tmp_path):
    deck = _deck()
    photos = fetch_deck_photos(deck, COMPANY, None, tmp_path / "media", provider=_openverse({"Electric Vehicle Charging": 5}))
    assert 0 in photos  # the cover's photo frame comes first
    out = tmp_path / "deck.pptx"
    render_pptx(deck, COMPANY, None, out, photos=photos)
    prs = pptx.Presentation(out)
    cover = prs.slides[0]
    assert any(sh.name == "content:image" for sh in cover.shapes)
    assert photos[0].credit in cover.notes_slide.notes_text_frame.text
    sources = " ".join(sh.text_frame.text for sh in prs.slides[-1].shapes if sh.has_text_frame)
    assert "Image: Photo: Creator" in sources


def test_photos_go_only_where_they_will_show_and_never_when_images_are_off(tmp_path):
    green = Path("data/sample_templates/Green Cream Simple Aesthetic Watercolor Presentation.pptx")
    photos = fetch_deck_photos(_deck(), green, None, tmp_path, provider=_openverse({"Electric Vehicle Charging": 5}))
    assert 0 not in photos  # the cover of this design has no photo slot: its artwork is kept as designed
    assert all(b - a >= 2 for a, b in zip(sorted(photos), sorted(photos)[1:]))  # never two photo slides in a row
    assert fetch_deck_photos(_deck(), COMPANY, None, tmp_path) == {}  # MOCK_LLM test run: provider disabled


# ── Ambiguous acronyms ──────────────────────────────────────────────────────


def test_acronym_only_topic_asks_a_question():
    from app.agents.supervisor import ambiguous_acronym_question, parse_plan

    assert ambiguous_acronym_question("Make a 10-slide deck on CJP")
    assert not ambiguous_acronym_question("Make a deck on EV charging in India")
    assert not ambiguous_acronym_question("Make a deck on CJP (Crime and Justice Programme)")
    assert not ambiguous_acronym_question("write a report on GDP")
    plan = parse_plan("Make a 10-slide deck on CJP")
    assert plan.error == "CLARIFICATION_NEEDED" and "CJP" in plan.clarification_needed


def test_clarification_reply_and_answer_completes_the_request():
    from app.main import app

    with TestClient(app) as api:
        token = api.post("/auth/login", json={"username": "demo", "password": "demo123"}).json()["access_token"]
        api.headers.update({"Authorization": f"Bearer {token}"})
        first = api.post("/chat", json={"message": "Make a 10-slide deck on CJP"}).json()
        msgs = api.get(f"/sessions/{first['session_id']}/messages").json()
        assert "CJP" in msgs[-1]["content"] and "?" in msgs[-1]["content"]

        seen = {}

        def fake_invoke(state):
            seen["message"] = state["user_message"]
            return {**state, "reply": "Done."}

        with patch("app.api.chat.graph_app.invoke", side_effect=fake_invoke):
            api.post("/chat", json={"message": "Crime and Justice Programme", "session_id": first["session_id"]})
        assert seen["message"] == "Make a 10-slide deck on CJP (meaning: Crime and Justice Programme)"
