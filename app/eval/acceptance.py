"""Acceptance evaluation (Stage 12): AC-2, AC-3, AC-4, AC-5, AC-6, AC-7, AC-8, AC-10 measured on ``eval/acceptance``.

Offline and deterministic: claims from the case file stand in for research, planners and writers run in offline
mode, every file is rendered, re-opened from disk and checked like a delivered artifact. ``python -m app.cli eval
acceptance`` prints the report; the live variant needs a working Gemini key.
"""

from __future__ import annotations

import json
import tempfile
from pathlib import Path
from typing import Any

CASES = Path("eval/acceptance/cases.json")
TEMPLATES = Path("data/sample_templates")
SOURCES = {i: {"title": f"Source {i}", "url": f"https://example.org/{i}", "publisher": f"Publisher {i}",
               "published_at": "2024-05-01"} for i in range(1, 5)}


class _Offline:
    """Deterministic planners and writers regardless of the configured key."""
    _mock_mode = True


OFFLINE = _Offline()


def _case(case: dict[str, Any], claims: list[Any], questions: dict[str, str], out: Path) -> dict[str, Any]:
    from app.agents.doc_generator import generate_document_model
    from app.agents.planner import plan_deck, plan_document, plan_workbook
    from app.agents.ppt_generator import generate_deck_model
    from app.agents.validator import validate_pptx_file
    from app.agents.xlsx_generator import generate_workbook_model
    from app.models.template_profile import TemplateProfile
    from app.services.analyzers.summary import NoLLM
    from app.services.docx_renderer import render_docx
    from app.services.length_rules import deck_length, document_length, workbook_length
    from app.services.plan_validation import document_words
    from app.services.pptx_renderer import render_pptx
    from app.services.quality_report import build_quality_report, validate_docx_file, validate_xlsx_file
    from app.services.xlsx_renderer import render_xlsx

    fmt, brief, length = case["format"], case["brief"], case.get("length")
    tpl = TEMPLATES / case["template"] if case.get("template") else None
    row: dict[str, Any] = {"id": case["id"], "scenario": case["scenario"], "format": fmt}
    path = out / f"{case['id']}.{fmt}"
    if fmt == "pptx":
        from app.agents.ppt_analyzer import analyze_presentation

        plan, report = plan_deck(brief, claims, deck_length(length, claims=claims), questions=questions, valid_sources=set(SOURCES), llm_client=OFFLINE)
        profile = analyze_presentation(tpl, llm_client=NoLLM()) if tpl else None
        deck = generate_deck_model(brief, profile or TemplateProfile(), plan=plan, claims=claims, llm_client=OFFLINE)
        render_pptx(deck, tpl, profile, path, sources_map=SOURCES, design_hint=brief)
        issues = validate_pptx_file(path, tpl)
        import pptx

        row["length_ok"] = (len(pptx.Presentation(str(path)).slides) - 1 == length["value"]) if length else True
        row["critical"] = [i.message for i in issues if i.severity == "error"]
        row["sample_text"] = [i.message for i in issues if "sample text" in i.message]
        recalc = None
    elif fmt == "docx":
        from app.agents.doc_analyzer import analyze_document

        target = document_length(length, claims=claims)
        plan, report = plan_document(brief, claims, target, questions=questions, valid_sources=set(SOURCES), llm_client=OFFLINE)
        profile = analyze_document(tpl, llm_client=NoLLM())
        doc = generate_document_model(brief, profile, plan=plan, claims=claims, llm_client=OFFLINE)
        doc.citation_style = case.get("citation_style", "numeric")
        render_docx(doc, tpl, profile, path, sources_map=SOURCES)
        issues = validate_docx_file(path)
        total = sum(s.target_words for s in plan.sections) + (plan.summary_words if plan.executive_summary else 0)
        row["length_ok"] = abs(total - target.value) <= target.value * 0.15  # plan allocation; written length needs a live model
        row["written_words"] = document_words(doc)
        row["critical"] = [i.message for i in issues if i.severity == "error"]
        recalc = None
    else:
        target = workbook_length(length, 0) if length else None
        plan, report = plan_workbook(brief, claims, target, valid_sources=set(SOURCES), llm_client=OFFLINE)
        wb = generate_workbook_model(brief, plan=plan, claims=claims, llm_client=OFFLINE)
        render_xlsx(wb, path)
        issues, recalc = validate_xlsx_file(path, wb.checks)
        row["length_ok"] = (len(wb.sheets) == length["value"]) if length else True
        row["critical"] = [i.message for i in issues if i.severity == "error"]
        row["formula_errors"] = recalc.errors if recalc else ["did not open"]
        row["checks"] = [c.ok for c in recalc.checks] if recalc else []
    row["reopens"] = not any("does not open" in m for m in row["critical"])
    row["claim_coverage_pct"] = report.claim_coverage_pct
    quality = build_quality_report(fmt, path, {"plan_reports": {fmt: report.model_dump()}}, file_issues=issues,
                                   recalc=recalc, offline=True)
    row["quality"] = {"verdict": quality.verdict, **quality.counts}
    return row


def _edit_regression(claims: list[Any], questions: dict[str, str]) -> dict[str, Any]:
    """AC-6: targeted and style-only edits change exactly their scope."""
    from app.agents.editor import _apply_ops_to_pptx
    from app.models.edit_ops import EditPlan
    from app.models.style import StyleOverrides
    from app.services.artifact_nodes import ensure_node_ids
    from app.services.edit_guard import regression
    from app.agents.offline_writer import write_deck
    from app.agents.planner import offline_deck_plan
    from app.models.content_plan import LengthTarget

    deck = ensure_node_ids(write_deck(offline_deck_plan("deck on community solar", claims,
                                                        LengthTarget(unit="slides", value=6, explicit=True),
                                                        questions=questions), claims))
    results = []
    plan = EditPlan.model_validate({"target": "pptx", "summary": "s", "ops": [
        {"type": "update_slide", "index": 2, "slide": {"title": "New title for slide two", "role": "title_content"}}]})
    try:
        edited = ensure_node_ids(_apply_ops_to_pptx(deck, plan.ops))
        g = regression(deck, edited, [deck.slides[1].node_id])
        results.append({"edit": "update slide 2", "ok": g.ok and g.changed_units == [deck.slides[1].node_id]})
    except Exception as exc:  # the op schema differs: record, do not hide
        results.append({"edit": "update slide 2", "ok": False, "error": str(exc)[:120]})
    restyled = deck.model_copy(deep=True)
    restyled.style = StyleOverrides(accent_color="C0392B")
    g = regression(deck, restyled, [])
    results.append({"edit": "restyle accent", "ok": g.changed_units == []})
    return {"cases": results, "pass_pct": round(100 * sum(r["ok"] for r in results) / len(results), 1)}


def evaluate(out_dir: Path | None = None) -> dict[str, Any]:
    from app.services.plan_inputs import ClaimRef

    data = json.loads(CASES.read_text(encoding="utf-8"))
    claims = [ClaimRef.model_validate(c) for c in data["claims"]]
    out = out_dir or Path(tempfile.mkdtemp(prefix="acceptance_"))
    out.mkdir(parents=True, exist_ok=True)
    rows = [_case(c, claims, data["questions"], out) for c in data["cases"]]
    n = len(rows)
    xlsx = [r for r in rows if r["format"] == "xlsx"]
    pptx_rows = [r for r in rows if r["format"] == "pptx"]
    checks = [ok for r in xlsx for ok in r["checks"]]
    metrics = {
        "AC-2_reopen_pct": round(100 * sum(r["reopens"] for r in rows) / n, 1),
        "AC-3_pptx_unresolved_critical": sum(len(r["critical"]) for r in pptx_rows),
        "AC-4_formula_errors": sum(len(r["formula_errors"]) for r in xlsx),
        "AC-4_checks_pass_pct": round(100 * sum(checks) / len(checks), 1) if checks else None,
        "AC-5_claim_coverage_min_pct": min(r["claim_coverage_pct"] for r in rows),
        "AC-6_edit_regression": _edit_regression(claims, data["questions"]),
        "AC-7_length_ok_pct": round(100 * sum(r["length_ok"] for r in rows) / n, 1),
        "AC-8_sample_text_findings": sum(len(r.get("sample_text", [])) for r in pptx_rows),
        "AC-10_quality_report_pct": round(100 * sum(1 for r in rows if r.get("quality")) / n, 1),
    }
    passed = (metrics["AC-2_reopen_pct"] == 100 and metrics["AC-3_pptx_unresolved_critical"] == 0
              and metrics["AC-4_formula_errors"] == 0 and metrics["AC-4_checks_pass_pct"] in (None, 100.0)
              and metrics["AC-5_claim_coverage_min_pct"] >= 95 and metrics["AC-7_length_ok_pct"] == 100
              and metrics["AC-8_sample_text_findings"] == 0 and metrics["AC-6_edit_regression"]["pass_pct"] == 100)
    return {"cases": rows, "metrics": metrics, "passed": passed, "output_dir": str(out)}
