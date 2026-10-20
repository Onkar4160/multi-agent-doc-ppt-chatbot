"""Deterministic Quality Validator for generated DOCX documents and PPTX presentation decks."""

from __future__ import annotations

import re
from typing import Any, Literal

from pydantic import BaseModel, Field

from app.models.deck_model import DeckModel
from app.models.document_model import DocumentModel
from app.models.workbook_model import WorkbookModel
from app.services.ir_compat import legacy_view
from app.services.plan_validation import block_texts, document_words


class ValidationIssue(BaseModel):
    """An issue identified by deterministic quality validation."""

    severity: Literal["error", "warning"]
    where: str  # e.g., "pptx:slide_3" or "docx:section_2" or "xlsx:sheet_1"
    message: str


class ValidationReport(BaseModel):
    """Unified validation report detailing pass/fail status and issues."""

    passed: bool
    score: float = Field(ge=0.0, le=100.0)
    issues: list[ValidationIssue] = Field(default_factory=list)


PLACEHOLDER_REGEX = re.compile(
    r"(click to add|lorem ipsum|insert text|\[placeholder\]|<placeholder>|sample text)",
    re.IGNORECASE,
)
NUMBER_OR_PERCENT_REGEX = re.compile(
    r"(\d+%|\$\d+|\d+\s*percent|\b\d{2,}\b|\bv?\d+\.\d+\b|\b(?:19|20)\d{2}\b|\b(?:January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{4}\b|\bversion\s+\d+\b|\brelease\s+\d+\b)",
    re.IGNORECASE,
)


def validate_scope(
    target_nodes: list[str] | None,
    actually_changed_nodes: list[str] | None,
) -> list[ValidationIssue]:
    """Validate that modifications stayed strictly within planned target nodes."""
    issues: list[ValidationIssue] = []
    if not target_nodes:
        return issues

    target_set = {str(n).lower().strip() for n in target_nodes}
    changed_set = {str(n).lower().strip() for n in (actually_changed_nodes or [])}

    # If the planned target is appending at the end, adding new nodes at the end is expected
    if "end" in target_set or "append" in target_set or any("append" in t for t in target_set):
        return issues

    def _normalize_node(n: str) -> set[str]:
        cleaned = re.sub(r"[_\-\s]+", " ", n).strip()
        num_match = re.search(r"\b(\d+)\b", cleaned)
        res = {n, cleaned}
        if num_match:
            res.add(num_match.group(1))
            res.add(f"slide {num_match.group(1)}")
            res.add(f"page {num_match.group(1)}")
        return res

    normalized_targets: set[str] = set()
    for t in target_set:
        normalized_targets.update(_normalize_node(t))

    unauthorized: list[str] = []
    for c in changed_set:
        c_norms = _normalize_node(c)
        if not c_norms.intersection(normalized_targets):
            unauthorized.append(c)

    if unauthorized:
        issues.append(
            ValidationIssue(
                severity="error",
                where="scope:targeted_edit",
                message=(
                    f"Scope violation: Planned target nodes were {sorted(target_set)}, "
                    f"but unapproved nodes were modified: {sorted(unauthorized)}."
                ),
            )
        )
    return issues


def _uncited_figure(text: str, request_text: str = "") -> bool:
    """A number, date or metric that needs a citation; figures the user wrote in the request (e.g. '26/11',
    'in 2026') are part of the subject, not claims."""
    allowed = {m.group(0).lower() for m in NUMBER_OR_PERCENT_REGEX.finditer(request_text or "")}
    return any(m.group(0).lower() not in allowed for m in NUMBER_OR_PERCENT_REGEX.finditer(text or ""))


def validate_outputs(
    doc_model: DocumentModel | None = None,
    deck_model: DeckModel | None = None,
    workbook_model: WorkbookModel | None = None,
    expected_slide_count: int = 12,
    valid_source_ids: set[int] | None = None,
    expected_doc_outline: list[str] | None = None,
    follow_template_outline: bool = False,
    min_target_words: int | None = 1200,
    target_nodes: list[str] | None = None,
    actually_changed_nodes: list[str] | None = None,
    request_text: str = "",
) -> ValidationReport:
    """Perform deterministic quality validation on DocumentModel, DeckModel, and WorkbookModel.

    Args:
        doc_model: Generated DocumentModel instance.
        deck_model: Generated DeckModel instance.
        workbook_model: Generated WorkbookModel instance.
        expected_slide_count: Target slide count from Plan.
        valid_source_ids: Set of valid integer source IDs in registry.
        expected_doc_outline: List of expected section headings.
        follow_template_outline: If True, enforce fuzzy match against template headings.
        target_nodes: Planned target node IDs for scoped edits.
        actually_changed_nodes: Actually modified node IDs.

    Returns:
        ValidationReport containing pass status, score (0-100), and issue list.
    """
    issues: list[ValidationIssue] = []
    valid_sids = valid_source_ids or set()

    # Scope validation
    if target_nodes:
        issues.extend(validate_scope(target_nodes, actually_changed_nodes))

    # ── 1. Validate PPTX Deck Model ──────────────────────────────────────────
    if deck_model:
        slides = deck_model.slides
        
        # Check slide count
        if len(slides) != expected_slide_count:
            issues.append(ValidationIssue(
                severity="error",
                where="pptx:deck",
                message=f"Slide count mismatch: generated {len(slides)} slides, expected {expected_slide_count}.",
            ))

        for idx, slide in enumerate(slides, start=1):
            slide_ref = f"pptx:slide_{idx}"
            
            # Combine all slide text (typed blocks count like bullets)
            slide = _legacy_slide(slide)
            bullets_text = []
            for b in (slide.bullets or []):
                bullets_text.append(b.text)
            for b in (slide.left or []):
                bullets_text.append(b.text)
            for b in (slide.right or []):
                bullets_text.append(b.text)

            all_slide_text = " ".join([slide.title, slide.subtitle or ""] + bullets_text)

            # Check placeholders
            if PLACEHOLDER_REGEX.search(all_slide_text):
                issues.append(ValidationIssue(
                    severity="error",
                    where=slide_ref,
                    message=f"Placeholder text detected on slide {idx}.",
                ))

            # Content slide checks (title_content, two_content)
            if slide.role in ("title_content", "two_content") and slide.archetype in (None, "bullets", "two_column", "comparison"):
                bullet_count = len(slide.bullets) if slide.role == "title_content" else (len(slide.left) + len(slide.right))
                if bullet_count < 3:
                    issues.append(ValidationIssue(
                        severity="warning",
                        where=slide_ref,
                        message=f"Content slide {idx} has fewer than 3 bullets ({bullet_count} found).",
                    ))

                word_count = len(all_slide_text.split())
                if word_count > 110:
                    issues.append(ValidationIssue(
                        severity="warning",
                        where=slide_ref,
                        message=f"Content slide {idx} word count ({word_count}) exceeds maximum 110 words.",
                    ))

            # Source citation checks for slides
            for b in (slide.bullets or []) + (slide.left or []) + (slide.right or []):
                for sid in getattr(b, "source_ids", []):
                    if valid_sids and sid not in valid_sids:
                        issues.append(ValidationIssue(
                            severity="error",
                            where=slide_ref,
                            message=f"Slide {idx} cites invalid source_id [{sid}].",
                        ))
                # Check numbers/percentages citation
                if _uncited_figure(b.text, request_text) and not getattr(b, "source_ids", []):
                    b_snip = b.text[:70] + "..." if len(b.text) > 70 else b.text
                    issues.append(ValidationIssue(
                        severity="error",
                        where=slide_ref,
                        message=f"Bullet on slide {idx} contains numbers/dates/metrics ('{b_snip}') but lacks citation source_ids.",
                    ))

    # ── 2. Validate DOCX Document Model ──────────────────────────────────────
    if doc_model:
        sections = doc_model.sections
        
        # Check section count (proportional / soft check)
        if len(sections) == 0:
            issues.append(ValidationIssue(
                severity="error",
                where="docx:document",
                message="DOCX has zero sections; minimum 1 required.",
            ))
        elif len(sections) < 3:
            issues.append(ValidationIssue(
                severity="warning",
                where="docx:document",
                message=f"DOCX has only {len(sections)} sections; document may be unusually short given the brief.",
            ))

        # Check total word count against minimum target
        total_word_count = document_words(doc_model)
        if min_target_words:
            half_target = min_target_words // 2
            if total_word_count < half_target:
                issues.append(ValidationIssue(
                    severity="error",
                    where="docx:document",
                    message=f"DOCX word count ({total_word_count}) is under 50% of minimum target ({half_target} words); document must be expanded.",
                ))
            elif total_word_count < min_target_words:
                issues.append(ValidationIssue(
                    severity="warning",
                    where="docx:document",
                    message=f"DOCX word count ({total_word_count}) is below target {min_target_words} words.",
                ))

        # Check outline fuzzy match only when the brief indicates the template's structure should be followed
        if expected_doc_outline and follow_template_outline:
            generated_headings = [s.heading.lower() for s in sections]
            for exp_h in expected_doc_outline[:6]:
                kw = exp_h.split()[-1].lower() if exp_h.split() else exp_h.lower()
                if not any(kw in gh for gh in generated_headings):
                    issues.append(ValidationIssue(
                        severity="warning",
                        where="docx:document",
                        message=f"Expected section heading '{exp_h}' not found in document.",
                    ))

        has_sources_section = False
        for idx, sec in enumerate(sections, start=1):
            sec_ref = f"docx:section_{idx}"
            if "source" in sec.heading.lower() or "reference" in sec.heading.lower():
                has_sources_section = True

            for block in sec.blocks:
                b_text = " ".join(block_texts(block))

                # Check placeholders
                if PLACEHOLDER_REGEX.search(b_text):
                    issues.append(ValidationIssue(
                        severity="error",
                        where=sec_ref,
                        message=f"Placeholder text detected in section {idx}.",
                    ))

                # Source citation checks
                for sid in getattr(block, "source_ids", []):
                    if valid_sids and sid not in valid_sids:
                        issues.append(ValidationIssue(
                            severity="error",
                            where=sec_ref,
                            message=f"Section {idx} block cites invalid source_id [{sid}].",
                        ))

                # Check numbers/percentages citation
                if _uncited_figure(b_text, request_text) and not getattr(block, "source_ids", []):
                    b_snip = b_text[:70] + "..." if len(b_text) > 70 else b_text
                    issues.append(ValidationIssue(
                        severity="error",
                        where=sec_ref,
                        message=f"Section {idx} text contains numbers/dates/metrics ('{b_snip}') but lacks citation source_ids.",
                    ))

    # ── 3. Validate XLSX Workbook Model ──────────────────────────────────────
    if workbook_model:
        sheets = workbook_model.sheets
        if not sheets:
            issues.append(ValidationIssue(
                severity="error",
                where="xlsx:workbook",
                message="Workbook has zero sheets; minimum 1 required.",
            ))
        for s_idx, sheet in enumerate(sheets, start=1):
            sheet_ref = f"xlsx:sheet_{s_idx}"
            if not sheet.name or not sheet.name.strip():
                issues.append(ValidationIssue(
                    severity="error",
                    where=sheet_ref,
                    message=f"Sheet {s_idx} is missing a valid name.",
                ))
            if PLACEHOLDER_REGEX.search(sheet.name):
                issues.append(ValidationIssue(
                    severity="error",
                    where=sheet_ref,
                    message=f"Placeholder text detected in sheet name '{sheet.name}'.",
                ))
            # Check tables
            for t_idx, table in enumerate(sheet.tables, start=1):
                if not table.headers:
                    issues.append(ValidationIssue(
                        severity="warning",
                        where=f"{sheet_ref}:table_{t_idx}",
                        message=f"Table {t_idx} in sheet '{sheet.name}' has no column headers.",
                    ))
                for row in table.rows:
                    for cell in row:
                        if isinstance(cell, str) and PLACEHOLDER_REGEX.search(cell):
                            issues.append(ValidationIssue(
                                severity="error",
                                where=f"{sheet_ref}:table_{t_idx}",
                                message=f"Placeholder cell value detected in table {t_idx}.",
                            ))
            # Check individual cells and formulas
            for cell in sheet.cells:
                if cell.formula and not cell.formula.startswith("="):
                    issues.append(ValidationIssue(
                        severity="warning",
                        where=f"{sheet_ref}:cell_{cell.coordinate}",
                        message=f"Formula in cell {cell.coordinate} does not start with '=': '{cell.formula}'",
                    ))

    # Compute overall score and pass status
    error_count = sum(1 for i in issues if i.severity == "error")
    warning_count = sum(1 for i in issues if i.severity == "warning")
    
    score = max(0.0, 100.0 - (error_count * 20.0) - (warning_count * 5.0))
    passed = error_count == 0

    return ValidationReport(passed=passed, score=score, issues=issues)


# ── Rendered PPTX checks (ported from DocForge ``validate/pptx_checks.py``) ──
# They read the file, not the model, so they judge what ships.

_EMU_PER_INCH = 914400
_TOLERANCE = int(0.02 * _EMU_PER_INCH)
# Minimum sizes for a 7.5-inch-tall slide, scaled to the real slide height.
_FLOORS = (
    ("content:headline", 20),
    ("content:title", 24),
    ("content:references", 10),
    ("content:", 14),
)
_FIT_MARKER = "overflows its box"


def _floor(name: str) -> float | None:
    return next((size for prefix, size in _FLOORS if name.startswith(prefix)), None)


def _shape_box(shape: Any) -> tuple[int, int, int, int]:
    return int(shape.left or 0), int(shape.top or 0), int(shape.width or 0), int(shape.height or 0)


def _paras(shape: Any) -> tuple[list[Any], list[tuple[float, bool, str | None]]]:
    """Measurement paragraphs plus per-paragraph (size, bold, colour) for typography checks."""
    from app.services.layout.text import Para

    paras, styles = [], []
    for p in shape.text_frame.paragraphs:
        runs = p.runs
        first = runs[0].font if runs else None
        size = first.size.pt if first is not None and first.size is not None else 18.0
        bold = bool(first.bold) if first is not None else False
        color = None
        if first is not None and first.color is not None and first.color.type is not None:
            try:
                color = str(first.color.rgb)
            except AttributeError:
                color = None
        p_pr = p._p.pPr
        indent = int(p_pr.get("marL", "0")) if p_pr is not None else 0
        paras.append(
            Para(
                "".join(r.text for r in runs),
                size,
                bold=bold,
                font=first.name if first is not None else None,
                indent_emu=indent,
                space_before_pt=p.space_before.pt if p.space_before is not None else 0.0,
                space_after_pt=p.space_after.pt if p.space_after is not None else 0.0,
                line_spacing=p.line_spacing if isinstance(p.line_spacing, float) else 1.0,
            )
        )
        styles.append((size, bold, color))
    return paras, styles


def _template_samples(template_path: Any) -> set[str]:
    """Sample sentences on the template's slides (text that must never ship)."""
    import pptx

    from app.services.analyzers.slide_design import all_text

    try:
        prs = pptx.Presentation(str(template_path))
    except Exception:
        return set()
    return {t for slide in prs.slides for s in slide.shapes for t in [all_text(s)] if len(t) > 12}


def validate_pptx_file(path: Any, template_path: Any = None) -> list[ValidationIssue]:
    """Measured text overflow, font floors, contrast on the real backdrop, sample residue and bounds of a deck."""
    import pptx

    from app.services.layout.backdrop import color_under
    from app.services.layout.colors import contrast_ratio
    from app.services.layout.text import measure

    try:
        prs = pptx.Presentation(str(path))
    except Exception as exc:
        return [ValidationIssue(severity="error", where="pptx:file", message=f"Deck does not open: {exc}")]

    issues: list[ValidationIssue] = []
    slide_w, slide_h = int(prs.slide_width), int(prs.slide_height)
    scale = slide_h / (7.5 * _EMU_PER_INCH)
    samples = _template_samples(template_path) if template_path else set()

    for i, slide in enumerate(prs.slides, start=1):
        where = f"pptx:slide_{i}"
        for shape in slide.shapes:
            name = shape.name or ""
            # Only our own shapes: template artwork often bleeds off the slide on purpose.
            if name.startswith(("content:", "decor:", "footer:")):
                x, y, w, h = _shape_box(shape)
                if x < -_TOLERANCE or y < -_TOLERANCE or x + w > slide_w + _TOLERANCE or y + h > slide_h + _TOLERANCE:
                    issues.append(ValidationIssue(
                        severity="error", where=where, message=f"Slide {i}: '{name}' extends beyond the slide."
                    ))
            if not getattr(shape, "has_text_frame", False):
                continue
            tf = shape.text_frame
            text = tf.text.strip()
            if not text:
                if shape.is_placeholder or name.startswith("content:"):
                    issues.append(ValidationIssue(
                        severity="error", where=where, message=f"Slide {i}: empty text box '{name}'."
                    ))
                continue
            if PLACEHOLDER_REGEX.search(text):
                issues.append(ValidationIssue(
                    severity="error", where=where, message=f"Slide {i}: placeholder text in '{name}'."
                ))
            if any(sample in text for sample in samples):
                issues.append(ValidationIssue(
                    severity="error", where=where, message=f"Slide {i}: template sample text remains in '{name}'."
                ))
            paras, styles = _paras(shape)
            if tf.word_wrap is not False:
                width = int(shape.width) - int(tf.margin_left) - int(tf.margin_right)
                height = int(shape.height) - int(tf.margin_top) - int(tf.margin_bottom)
                m = measure(paras, width)
                if m.height_emu > height * 1.02 or m.widest_word_overflow:
                    over = max((m.height_emu - height) / _EMU_PER_INCH, 0)
                    issues.append(ValidationIssue(
                        severity="error",
                        where=where,
                        message=f"Slide {i}: text in '{name}' {_FIT_MARKER} by {over:.2f} in"
                        + (" (a word is wider than the box)." if m.widest_word_overflow else "."),
                    ))
            floor = _floor(name)
            if floor is not None:
                floor = round(floor * scale, 1)
                small = [s for s, _, _ in styles if s < floor - 0.05]
                if small:
                    issues.append(ValidationIssue(
                        severity="warning",
                        where=where,
                        message=f"Slide {i}: '{name}' uses {min(small):.0f}pt, below the {floor}pt minimum.",
                    ))
            if name.startswith(("content:", "footer:")):
                under = color_under(slide, _shape_box(shape), stop_at=shape._element)
                for size, bold, color in styles:
                    if not color or not under:
                        continue
                    needed = 3.0 if size >= 18 * scale or (bold and size >= 14 * scale) else 4.5
                    ratio = contrast_ratio(color, under)
                    if ratio < needed:
                        issues.append(ValidationIssue(
                            severity="warning",
                            where=where,
                            message=f"Slide {i}: text in '{name}' has contrast {ratio:.1f}:1 (needs {needed}:1).",
                        ))
                        break
    return issues


def fit_problems(issues: list[Any]) -> list[str]:
    """Text-fit findings as instructions the deck generator can act on (one line per slide)."""
    by_slide: dict[int, str] = {}
    for issue in issues:
        data = issue if isinstance(issue, dict) else issue.model_dump()
        where = data.get("where", "")
        msg = data.get("message", "")
        # The Sources slide, chart notes and footers are generated by the renderer: rewriting slides cannot fix them.
        if any(f"'{p}" in msg for p in ("content:references", "footer:", "decor:")):
            continue
        if _FIT_MARKER in msg and where.startswith("pptx:slide_"):
            by_slide.setdefault(int(where.rsplit("_", 1)[-1]), data["message"])
    return [
        f"Slide {n}: text does not fit ({msg}). Shorten it: headline at most 8 words, at most 4 bullets of "
        "at most 12 words; keep the facts and source_ids."
        for n, msg in sorted(by_slide.items())
    ]


def merge_report(report: dict[str, Any] | None, extra: list[ValidationIssue]) -> dict[str, Any]:
    """Add rendered-file issues to a model validation report, recomputing score and pass status."""
    issues = [ValidationIssue.model_validate(i) for i in (report or {}).get("issues", [])] + list(extra)
    errors = sum(1 for i in issues if i.severity == "error")
    warnings = sum(1 for i in issues if i.severity == "warning")
    return ValidationReport(
        passed=errors == 0, score=max(0.0, 100.0 - errors * 20.0 - warnings * 5.0), issues=issues
    ).model_dump()


def _legacy_slide(slide: Any) -> Any:
    """The slide as the renderer shows it: typed blocks flattened into bullets or columns."""
    if not getattr(slide, "blocks", None):
        return slide
    return legacy_view(DeckModel(title="", slides=[slide])).slides[0]

