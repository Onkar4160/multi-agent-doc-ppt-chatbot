"""Document parsing service for DOCX, PPTX, PDF, and image files into structured ParsedContent."""

from __future__ import annotations

import logging
from pathlib import Path
from typing import Any

import docx
from pydantic import BaseModel, Field
import pymupdf
import pptx

from app.models.template_profile import FontInfo
from app.services.ocr import ocr_image, ocr_scanned_pdf

logger = logging.getLogger(__name__)


class ParsedBlock(BaseModel):
    """A typed block of document content (heading, paragraph, table, list item, image)."""
    type: str  # "heading" | "paragraph" | "list" | "table" | "image" | "slide_title" | "slide_body" | "ocr_text"
    text: str = ""
    level: int = 0
    style_name: str | None = None
    font_info: FontInfo | None = None
    table_data: list[list[str]] | None = None
    metadata: dict[str, Any] = Field(default_factory=dict)


class ParsedContent(BaseModel):
    """Unified structured parse output from any document type."""
    file_type: str
    blocks: list[ParsedBlock] = Field(default_factory=list)
    text_preview: str = ""
    page_or_slide_count: int = 1
    is_scanned: bool = False
    raw_text: str = ""


# ── 1. Parse DOCX ──────────────────────────────────────────────────────────

def parse_docx(file_path: str | Path) -> ParsedContent:
    """Parse a Microsoft Word (.docx) document."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"DOCX file not found: {file_path}")

    try:
        doc = docx.Document(path)
    except Exception as exc:
        raise ValueError(f"Corrupt or invalid DOCX file '{path.name}': {exc}") from exc

    blocks: list[ParsedBlock] = []
    text_parts: list[str] = []

    # Process Body Paragraphs & Tables in document order
    for elem in doc.element.body:
        if elem.tag.endswith('p'):
            p = docx.text.paragraph.Paragraph(elem, doc)
            text = p.text.strip()
            if not text:
                continue

            style_name = p.style.name if p.style else "Normal"
            font_info = _extract_paragraph_font(p)

            # Detect headings by style name or font size + bold
            is_heading_style = style_name.lower().startswith("heading") or "title" in style_name.lower()
            is_heading_font = bool(font_info.bold and font_info.size_pt and font_info.size_pt >= 14)

            if is_heading_style or is_heading_font:
                try:
                    level = int("".join(c for c in style_name if c.isdigit()) or ("1" if font_info.size_pt and font_info.size_pt >= 16 else "2"))
                except ValueError:
                    level = 1
                b_type = "heading"
            elif style_name.lower().startswith("list"):
                level = 1
                b_type = "list"
            else:
                level = 0
                b_type = "paragraph"

            blocks.append(ParsedBlock(
                type=b_type,
                text=text,
                level=level,
                style_name=style_name,
                font_info=font_info,
            ))
            text_parts.append(text)

        elif elem.tag.endswith('tbl'):
            table = docx.table.Table(elem, doc)
            grid: list[list[str]] = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                grid.append(row_data)

            tbl_text = "\n".join(" | ".join(r) for r in grid if any(r))
            blocks.append(ParsedBlock(
                type="table",
                text=tbl_text,
                table_data=grid,
                style_name=table.style.name if table.style else "Table Grid",
            ))
            if tbl_text:
                text_parts.append(tbl_text)

    # Process Inline Shapes (Images)
    image_count = len(doc.inline_shapes)
    if image_count > 0:
        blocks.append(ParsedBlock(
            type="image",
            text=f"[{image_count} embedded image(s)]",
            metadata={"image_count": image_count},
        ))

    full_text = "\n".join(text_parts)
    text_preview = full_text[:1000] + ("..." if len(full_text) > 1000 else "")

    return ParsedContent(
        file_type="docx",
        blocks=blocks,
        text_preview=text_preview,
        page_or_slide_count=max(1, len(doc.paragraphs) // 25),
        is_scanned=False,
        raw_text=full_text,
    )


# ── 2. Parse PPTX ──────────────────────────────────────────────────────────

def parse_pptx(file_path: str | Path) -> ParsedContent:
    """Parse a Microsoft PowerPoint (.pptx) presentation."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"PPTX file not found: {file_path}")

    try:
        prs = pptx.Presentation(path)
    except Exception as exc:
        raise ValueError(f"Corrupt or invalid PPTX file '{path.name}': {exc}") from exc

    blocks: list[ParsedBlock] = []
    text_parts: list[str] = []
    slide_count = len(prs.slides)

    for idx, slide in enumerate(prs.slides, start=1):
        slide_title = ""
        if slide.shapes.title and slide.shapes.title.text:
            slide_title = slide.shapes.title.text.strip()

        if slide_title:
            blocks.append(ParsedBlock(
                type="slide_title",
                text=slide_title,
                level=1,
                metadata={"slide_number": idx, "layout_name": slide.slide_layout.name},
            ))
            text_parts.append(f"Slide {idx}: {slide_title}")

        body_texts: list[str] = []
        for shape in slide.shapes:
            if shape.has_text_frame and shape != slide.shapes.title:
                txt = shape.text_frame.text.strip()
                if txt:
                    body_texts.append(txt)
            elif shape.has_table:
                grid = [[cell.text.strip() for cell in row.cells] for row in shape.table.rows]
                tbl_txt = "\n".join(" | ".join(r) for r in grid if any(r))
                if tbl_txt:
                    blocks.append(ParsedBlock(
                        type="table",
                        text=tbl_txt,
                        table_data=grid,
                        metadata={"slide_number": idx},
                    ))

        if body_texts:
            full_body = "\n".join(body_texts)
            blocks.append(ParsedBlock(
                type="slide_body",
                text=full_body,
                metadata={"slide_number": idx},
            ))
            text_parts.append(full_body)

    full_text = "\n\n".join(text_parts)
    text_preview = full_text[:1000] + ("..." if len(full_text) > 1000 else "")

    return ParsedContent(
        file_type="pptx",
        blocks=blocks,
        text_preview=text_preview,
        page_or_slide_count=slide_count,
        is_scanned=False,
        raw_text=full_text,
    )


# ── 3. Parse PDF ───────────────────────────────────────────────────────────

def parse_pdf(file_path: str | Path, ocr_scanned: bool = True) -> ParsedContent:
    """Parse a PDF document using PyMuPDF; automatically detects scanned PDFs."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"PDF file not found: {file_path}")

    try:
        doc = pymupdf.open(path)
    except Exception as exc:
        raise ValueError(f"Corrupt or invalid PDF file '{path.name}': {exc}") from exc

    blocks: list[ParsedBlock] = []
    text_parts: list[str] = []
    page_count = len(doc)

    for idx, page in enumerate(doc, start=1):
        page_blocks = page.get_text("blocks")
        for b in page_blocks:
            if len(b) >= 5 and isinstance(b[4], str):
                txt = b[4].strip()
                if txt:
                    blocks.append(ParsedBlock(
                        type="paragraph",
                        text=txt,
                        metadata={"page_number": idx, "bbox": b[:4]},
                    ))
                    text_parts.append(txt)

    doc.close()
    full_text = "\n".join(text_parts)
    is_scanned = len(full_text.strip()) < 50

    if is_scanned and ocr_scanned:
        logger.info("PDF '%s' appears to be scanned. Running OCR...", path.name)
        ocr_text = ocr_scanned_pdf(path)
        if ocr_text:
            full_text = ocr_text
            blocks = [ParsedBlock(type="ocr_text", text=ocr_text)]

    text_preview = full_text[:1000] + ("..." if len(full_text) > 1000 else "")

    return ParsedContent(
        file_type="pdf",
        blocks=blocks,
        text_preview=text_preview,
        page_or_slide_count=page_count,
        is_scanned=is_scanned,
        raw_text=full_text,
    )


# ── 4. Parse Image ─────────────────────────────────────────────────────────

def parse_image(file_path: str | Path) -> ParsedContent:
    """Parse an image file (PNG/JPG) using Gemini Vision OCR."""
    path = Path(file_path)
    if not path.exists():
        raise FileNotFoundError(f"Image file not found: {file_path}")

    image_bytes = path.read_bytes()
    ext = path.suffix.lower().lstrip(".")
    mime = f"image/{'jpeg' if ext in ('jpg', 'jpeg') else 'png'}"

    raw_text = ocr_image(image_bytes, mime_type=mime)
    blocks = [ParsedBlock(type="ocr_text", text=raw_text)] if raw_text else []
    text_preview = raw_text[:1000] + ("..." if len(raw_text) > 1000 else "")

    return ParsedContent(
        file_type="image",
        blocks=blocks,
        text_preview=text_preview,
        page_or_slide_count=1,
        is_scanned=True,
        raw_text=raw_text,
    )


# ── Dispatcher ─────────────────────────────────────────────────────────────

def parse_file(file_path: str | Path, ocr_scanned: bool = True) -> ParsedContent:
    """Route file to appropriate parser by extension."""
    path = Path(file_path)
    ext = path.suffix.lower()

    if ext == ".docx":
        return parse_docx(path)
    elif ext == ".pptx":
        return parse_pptx(path)
    elif ext == ".pdf":
        return parse_pdf(path, ocr_scanned=ocr_scanned)
    elif ext in (".png", ".jpg", ".jpeg"):
        return parse_image(path)
    else:
        raise ValueError(f"Unsupported file format '{ext}' for parsing")


# ── Internal Helpers ───────────────────────────────────────────────────────

def _extract_paragraph_font(p: docx.text.paragraph.Paragraph) -> FontInfo:
    """Extract primary font formatting from paragraph runs."""
    font_name = None
    size_pt = None
    bold = None
    italic = None
    color_hex = None

    for run in p.runs:
        if run.font.name:
            font_name = run.font.name
        if run.font.size:
            size_pt = run.font.size.pt
        if run.bold is not None:
            bold = run.bold
        if run.italic is not None:
            italic = run.italic
        if run.font.color and run.font.color.rgb:
            color_hex = str(run.font.color.rgb)
        if font_name and size_pt:
            break

    return FontInfo(
        name=font_name or "Calibri",
        size_pt=size_pt or 11.0,
        bold=bold or False,
        italic=italic or False,
        color_hex=color_hex,
    )
