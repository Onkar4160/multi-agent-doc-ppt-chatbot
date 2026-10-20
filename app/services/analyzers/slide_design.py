"""Templates whose design lives on the sample slides (Canva, SlideModel and most marketplace templates).

Their layouts are often blank; backgrounds, watercolour art, panels and photo frames sit on the sample slides.
For those, each sample slide becomes a reusable *base*: decorations are kept, the sample text is removed, the
title box keeps its styling for our headline, and the area the sample body text used becomes the content
region. This module only analyses; ``app.services.renderers.pptx_clone`` performs the cloning. Ported from DocForge.
"""

from __future__ import annotations

import colorsys
import io
import re
from collections import Counter

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from app.services.layout.backdrop import color_under, pixels_of
from app.models.template_profile import Box, SlideBase, SlideDesign

_PROMO = re.compile(
    r"slidemodel|coupon|subscribe now|download (free )?powerpoint|freebundle|envato|slidesgo|"
    r"freepik|this presentation template",
    re.IGNORECASE,
)
_CLOSING = re.compile(r"\b(thank|questions|q\s*&\s*a|contact us|the end|get in touch)\b", re.IGNORECASE)
_SCHEME = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}
_META = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.HEADER}
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"


def _box(shape) -> Box | None:
    if shape.left is None or shape.width is None:
        return None
    return Box(x=int(shape.left), y=int(shape.top), w=int(shape.width), h=int(shape.height))


def _ph_type(shape):
    if not shape.is_placeholder:
        return None
    try:
        return shape.placeholder_format.type
    except (ValueError, AttributeError):
        return None


def all_text(shape) -> str:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return " ".join(all_text(s) for s in shape.shapes).strip()
    if getattr(shape, "has_text_frame", False):
        return shape.text_frame.text.strip()
    if getattr(shape, "has_table", False):
        return " ".join(c.text for r in shape.table.rows for c in r.cells).strip()
    return ""


def _runs(shape):
    if getattr(shape, "has_text_frame", False):
        for p in shape.text_frame.paragraphs:
            yield from p.runs


def max_size(shape) -> float:
    sizes = [r.font.size.pt for r in _runs(shape) if r.font.size is not None]
    if sizes:
        return max(sizes)
    from app.services.renderers.pptx_canvas import inherited_size

    if shape.is_placeholder:
        return inherited_size(shape) or 18.0
    return 18.0


def run_style(shape) -> tuple[str | None, str | None]:
    """(font name, hex colour) of the first styled run."""
    for r in _runs(shape):
        color = None
        try:
            if r.font.color is not None and r.font.color.type is not None:
                color = str(r.font.color.rgb)
        except AttributeError:
            color = None
        if r.font.name or color:
            return r.font.name, color
    return None, None


def _is_visual(shape) -> bool:
    """Picture, picture-filled or filled shape, or a text-free group: the template's artwork."""
    if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
        return True
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return not all_text(shape)
    el = shape._element
    if el.findall(f".//{_A}blipFill"):
        return True
    return bool(el.findall(f".//{_A}solidFill")) and not all_text(shape)


def _union(boxes: list[Box]) -> Box | None:
    boxes = [b for b in boxes if b is not None]
    if not boxes:
        return None
    x, y = min(b.x for b in boxes), min(b.y for b in boxes)
    return Box(x=x, y=y, w=max(b.right for b in boxes) - x, h=max(b.bottom for b in boxes) - y)


def _contains(outer: Box, inner: Box) -> bool:
    return (
        outer.x <= inner.x
        and outer.right >= inner.right
        and outer.y <= inner.y
        and outer.bottom >= inner.bottom
    )


def _extend_down(region: Box, blockers: list[Box], slide_w: int, slide_h: int) -> Box:
    """Sample body boxes are sized to their sample text; the real space usually continues downwards."""
    limit = int(slide_h * 0.9)
    for b in blockers:
        overlaps_x = b.x < region.right and b.right > region.x
        contains = _contains(b, region)
        if overlaps_x and not contains and region.bottom <= b.y < limit:
            limit = b.y
    # Stay inside a panel that contains the region (e.g. a card behind the sample text).
    for b in blockers:
        if (
            b.x <= region.x
            and b.right >= region.right
            and b.y <= region.y
            and region.bottom <= b.bottom < limit
        ):
            limit = b.bottom - int(slide_h * 0.03)
    return Box(x=region.x, y=region.y, w=region.w, h=max(region.h, limit - region.y))


def _avoid_sides(region: Box, blockers: list[Box], slide_w: int) -> Box:
    """Artwork reaching into the side of the area (shapes beside the sample text) narrows it."""
    gap = int(slide_w * 0.015)
    left, right = region.x, region.right
    for b in blockers:
        overlap = b.intersection_area(region)
        if _contains(b, region) or overlap == 0 or overlap >= 0.5 * b.area():
            continue  # backdrop panel, unrelated, or sample structure that is removed with the text
        if b.x > region.x + region.w // 2:
            right = min(right, b.x - gap)
        elif b.right < region.right - region.w // 2:
            left = max(left, b.right + gap)
    if right - left < 0.5 * region.w:
        return region  # artwork across the middle: keep the sample's area
    return Box(x=left, y=region.y, w=right - left, h=region.h)


def analyse_slide(index: int, slide, slide_w: int, slide_h: int) -> SlideBase | None:
    shapes = list(slide.shapes)
    texts = {i: all_text(s) for i, s in enumerate(shapes)}
    joined = " ".join(texts.values())
    if not joined.strip() or _PROMO.search(joined):
        return None  # image-only or promotional slide: never a base

    remove: list[int] = []
    photos: list[int] = []
    text_slots: list[int] = []
    for i, s in enumerate(shapes):
        ptype = _ph_type(s)
        if ptype == PP_PLACEHOLDER.SLIDE_NUMBER:
            continue
        if ptype in _META:
            remove.append(i)
        elif ptype == PP_PLACEHOLDER.PICTURE and not texts[i]:
            photos.append(i)
        elif getattr(s, "has_chart", False) or getattr(s, "has_table", False):
            remove.append(i)
        elif texts[i]:
            text_slots.append(i)
    if not text_slots:
        return None

    def title_key(i: int):
        s = shapes[i]
        wordy = any(ch.isalpha() for ch in texts[i]) and len(texts[i]) <= 90
        return (wordy and s.shape_type != MSO_SHAPE_TYPE.GROUP, max_size(s), -int(s.top or 0))

    title = max(text_slots, key=title_key)
    others = [i for i in text_slots if i != title]
    title_text = texts[title]
    if index == 0:
        kind = "cover"
    elif _CLOSING.search(title_text):
        kind = "closing"
    else:
        kind = "content"

    subtitle = None
    if kind in ("cover", "closing") and others:
        short = [i for i in others if len(texts[i]) <= 160 and shapes[i].shape_type != MSO_SHAPE_TYPE.GROUP]
        title_mid = int(shapes[title].top or 0) + int(shapes[title].height or 0) // 2
        below = [i for i in short if int(shapes[i].top or 0) >= title_mid]
        short = below or short  # a subtitle sits under the title; text above it is a kicker
        if short:
            subtitle = max(short, key=lambda i: (max_size(shapes[i]), -int(shapes[i].top or 0)))
    body = [i for i in others if i != subtitle]
    remove += body
    # Repeats of the title word are typography (e.g. a stacked "DISCUSSION"), not a content area.
    body = [i for i in body if texts[i].strip().lower() != title_text.strip().lower()]

    title_box = _box(shapes[title])
    region = _union([_box(shapes[i]) for i in body])
    kept_visuals = [
        _box(s)
        for i, s in enumerate(shapes)
        if i not in remove and i not in text_slots and _is_visual(s) and _box(s) is not None
    ]
    # Artwork on the slide's layout shows through the clone too.
    kept_visuals += [
        _box(s) for s in slide.slide_layout.shapes if not s.is_placeholder and _is_visual(s) and _box(s) is not None
    ]
    full = slide_w * slide_h
    blockers = [b for b in kept_visuals if b.area() < 0.9 * full]  # full-slide backgrounds never block
    if region is not None:
        if title_box is not None and region.y < title_box.bottom:
            cut = title_box.bottom + int(slide_h * 0.02)
            region = Box(x=region.x, y=cut, w=region.w, h=max(region.bottom - cut, int(slide_h * 0.1)))
        region = _avoid_sides(_extend_down(region, blockers, slide_w, slide_h), blockers, slide_w)
        margin = int(slide_w * 0.03)
        x0, y0 = max(region.x, margin), max(region.y, margin)
        x1, y1 = min(region.right, slide_w - margin), min(region.bottom, slide_h - margin)
        region = Box(x=x0, y=y0, w=max(x1 - x0, 0), h=max(y1 - y0, 0)) if x1 > x0 and y1 > y0 else None
    frames: list[int] = []
    if region is not None:
        # Sample structure inside the content area (column panels, bullet icons, diagram parts) goes;
        # panels that contain the whole area and artwork that only touches it stay.
        for i, s in enumerate(shapes):
            b = _box(s)
            if i in remove or i in text_slots or i in photos or b is None or _contains(b, region):
                continue
            thick = Box(x=b.x, y=b.y, w=max(b.w, 12700), h=max(b.h, 12700))  # lines have no area
            if thick.intersection_area(region) >= 0.5 * thick.area():
                remove.append(i)
        for i, s in enumerate(shapes):
            b = _box(s)
            if i in remove or i in text_slots or i in photos or b is None:
                continue
            fill = _own_fill(s)
            share = b.area() / full
            framed = fill and _light(fill) and 0.04 <= share <= 0.3 and 0.5 <= b.w / max(b.h, 1) <= 2.2
            if framed and b.intersection_area(region) <= 0.2 * b.area():
                frames.append(i)
    sample_photos = [
        i for i, s in enumerate(shapes)
        if i not in remove and i not in text_slots and i not in photos and _box(s) is not None
        and _box(s).area() >= 0.025 * full and is_sample_photo(s, slide)
    ]
    area = region.area() / full if region else 0.0
    region_color = None
    if region is not None:
        first_removed = shapes[min(remove)]._element if remove else None
        region_color = color_under(slide, (region.x, region.y, region.w, region.h), stop_at=first_removed)
    if kind == "content" and area < 0.08:
        # Few text items and little body: a divider. Otherwise a cramped diagram slide we avoid.
        kind = "section" if len(text_slots) <= 2 else "diagram"
    # Bases with many small sample items are diagrams; prefer simple ones.
    score = area - 0.04 * max(0, len(body) - 3) - (0.1 if len(photos) > 1 else 0)
    # Artwork beside a short sample title covers our longer headline once the title box is widened.
    title_blocked = False
    if title_box is not None and kind == "content":
        wide = Box(x=title_box.x, y=title_box.y, w=max(title_box.w, region.w if region else 0), h=title_box.h)
        if any(b.intersection_area(wide) > 0.1 * wide.area() for b in blockers):
            score -= 0.2
            title_blocked = True
    return SlideBase(
        index=index,
        kind=kind,
        title_shape=title,
        subtitle_shape=subtitle,
        remove_shapes=sorted(set(remove)),
        photo_shapes=photos,
        frame_shapes=frames,
        sample_photos=sample_photos,
        title_blocked=title_blocked,
        region=region,
        score=round(score, 3),
        title_text=title_text[:60],
        region_color=region_color,
    )


def photo_blips(shape) -> list:
    """``a:blip`` elements of every picture or picture-filled shape inside ``shape`` (groups at any depth)."""
    el = shape._element if hasattr(shape, "_element") else shape
    return list(el.iter(f"{_A}blip"))


def _photographic(blob: bytes) -> bool:
    """Opaque with many distinct colours: a photograph (not a logo, icon or transparent watercolour leaf)."""
    from PIL import Image

    try:
        img = Image.open(io.BytesIO(blob)).convert("RGBA")
    except Exception:  # noqa: BLE001 - unreadable media is not a photo slot
        return False
    img.thumbnail((48, 48))
    px = list(pixels_of(img))
    if not px:
        return False
    opaque = [p for p in px if p[3] > 240]
    if len(opaque) < 0.95 * len(px):
        return False
    colours = {(r // 16, g // 16, b // 16) for r, g, b, _a in opaque}
    return len(colours) >= 60


def is_sample_photo(shape, slide) -> bool:
    """A picture (or group holding one) whose image is a photograph: the template's stand-in for a topic photo."""
    for blip in photo_blips(shape):
        rid = blip.get(_R_EMBED)
        if not rid:
            continue
        try:
            if _photographic(slide.part.related_part(rid).blob):
                return True
        except KeyError:
            continue
    return False


def _own_fill(shape) -> str | None:
    """Solid colour of a shape, or of the first filled child of a group."""
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return next((c for c in (_own_fill(ch) for ch in shape.shapes) if c), None)
    node = shape._element.find(f".//{_A}spPr/{_A}solidFill/{_A}srgbClr")
    return node.get("val").upper() if node is not None and node.get("val") else None


def _light(hex_color: str) -> bool:
    _h, lum, _s = colorsys.rgb_to_hls(*(int(hex_color[i : i + 2], 16) / 255 for i in (0, 2, 4)))
    return lum > 0.88


def design_mode(prs, bases: list[SlideBase]) -> str:
    """'slides' when the artwork sits on the sample slides rather than on the masters and layouts."""
    if len(bases) < 2 or not any(b.kind == "content" for b in bases):
        return "layouts"
    slides = list(prs.slides)
    on_slides = [sum(1 for s in slides[b.index].shapes if _is_visual(s)) for b in bases]
    used_layouts = {id(slides[b.index].slide_layout): slides[b.index].slide_layout for b in bases}
    on_layouts = [sum(1 for s in lay.shapes if not s.is_placeholder) for lay in used_layouts.values()]
    avg_slides = sum(on_slides) / len(on_slides)
    avg_layouts = sum(on_layouts) / max(1, len(on_layouts))
    return "slides" if avg_slides >= 1.5 and avg_slides > avg_layouts else "layouts"


def _hex(rgb) -> str:
    return "".join(f"{c:02X}" for c in rgb[:3])


def slide_style(
    prs, bases: list[SlideBase], theme_colors: dict[str, str] | None = None
) -> tuple[dict[str, str], list[str], str | None, str | None]:
    """Fonts (heading/body), accent palette, title text colour and dominant light background colour.

    Accents come from shape fills first (theme references resolved), then from the artwork's pixels.
    """
    theme_colors = theme_colors or {}
    from PIL import Image

    slides = list(prs.slides)
    heading_fonts: Counter = Counter()
    body_fonts: Counter = Counter()
    title_colors: Counter = Counter()
    pixels: Counter = Counter()
    for b in bases:
        shapes = list(slides[b.index].shapes)
        if b.title_shape is not None:
            name, color = run_style(shapes[b.title_shape])
            if name:
                heading_fonts[name] += 1
            if color:
                title_colors[color] += 1
        for i in [*b.remove_shapes, *([b.subtitle_shape] if b.subtitle_shape is not None else [])]:
            name, _ = run_style(shapes[i])
            if name:
                body_fonts[name] += 1
    fills: Counter = Counter()
    parts = [slides[b.index] for b in bases] + list(
        {id(slides[b.index].slide_layout): slides[b.index].slide_layout for b in bases}.values()
    )
    for part in parts:
        for fill in part._element.iter(f"{_A}solidFill"):
            if not len(fill):
                continue
            node = fill[0]
            val = node.get("val")
            if node.tag == f"{_A}schemeClr":
                val = theme_colors.get(_SCHEME.get(val, val))
            if val and len(val) == 6:
                fills[val.upper()] += 1
    seen: set[str] = set()
    for b in bases[:8]:
        slide = slides[b.index]
        for blip in slide._element.iter(f"{_A}blip"):
            rid = blip.get(_R_EMBED)
            if not rid or rid in seen:
                continue
            seen.add(rid)
            try:
                img = Image.open(io.BytesIO(slide.part.related_part(rid).blob)).convert("RGBA")
            except Exception:  # noqa: BLE001 - unreadable media is just skipped
                continue
            img.thumbnail((64, 64))
            for r, g, bl, a in pixels_of(img):
                if a > 200:
                    pixels[(r // 16 * 16, g // 16 * 16, bl // 16 * 16)] += 1
    fonts = {}
    if heading_fonts:
        fonts["heading"] = heading_fonts.most_common(1)[0][0]
    if body_fonts:
        fonts["body"] = body_fonts.most_common(1)[0][0]
    accents, background = [], None
    for hexed, _count in fills.most_common(30):
        _h, lum, sat = colorsys.rgb_to_hls(*(int(hexed[i : i + 2], 16) / 255 for i in (0, 2, 4)))
        if 0.15 < lum < 0.8 and sat > 0.2 and all(_distance(hexed, a) > 60 for a in accents):
            accents.append(hexed)
    for rgb, _count in pixels.most_common(60):
        h, lum, sat = colorsys.rgb_to_hls(*(c / 255 for c in rgb))
        if lum > 0.85 and background is None:
            background = _hex(rgb)
        if 0.2 < lum < 0.75 and sat > 0.2:
            hexed = _hex(rgb)
            if all(_distance(hexed, a) > 60 for a in accents):
                accents.append(hexed)
        if len(accents) >= 4:
            break
    text_color = title_colors.most_common(1)[0][0] if title_colors else theme_colors.get("dk1")
    if not accents:
        # A monochrome design (greys, black and white): its accents are the dark neutrals it fills shapes with,
        # not the theme's default accent colour, which the slides never show.
        for hexed, _count in fills.most_common(30):
            _h, lum, _s = colorsys.rgb_to_hls(*(int(hexed[i : i + 2], 16) / 255 for i in (0, 2, 4)))
            if 0.08 < lum < 0.6 and all(_distance(hexed, a) > 40 for a in accents):
                accents.append(hexed)
            if len(accents) >= 3:
                break
        if not accents and text_color:
            accents.append(text_color)
    return fonts, accents, text_color, background


def _distance(a: str, b: str) -> float:
    return sum(abs(int(a[i : i + 2], 16) - int(b[i : i + 2], 16)) for i in (0, 2, 4))


_SLOTS = ("dk1", "lt1", "dk2", "lt2", "accent1", "accent2", "accent3", "accent4", "accent5", "accent6", "hlink")


def theme_info(prs) -> tuple[dict[str, str], dict[str, str]]:
    """Theme colour scheme slots (hex, no '#') and major/minor fonts of the first master."""
    from lxml import etree
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    try:
        root = etree.fromstring(prs.slide_master.part.part_related_by(RT.THEME).blob)
    except (KeyError, etree.XMLSyntaxError):
        return {}, {}
    colors: dict[str, str] = {}
    scheme = root.find(f".//{_A}clrScheme")
    if scheme is not None:
        for slot in _SLOTS:
            node = scheme.find(f"{_A}{slot}")
            if node is None or not len(node):
                continue
            value = node[0].get("val") if node[0].tag == f"{_A}srgbClr" else node[0].get("lastClr")
            if value:
                colors[slot] = value.upper()
    fonts: dict[str, str] = {}
    for key, tag in (("major", "majorFont"), ("minor", "minorFont")):
        latin = root.find(f".//{_A}{tag}/{_A}latin")
        if latin is not None and latin.get("typeface"):
            fonts[key] = latin.get("typeface")
    return colors, fonts


def analyze_slide_design(prs) -> SlideDesign:
    """Analyse the sample slides of an open presentation into reusable bases and decide the design mode."""
    slide_w, slide_h = int(prs.slide_width), int(prs.slide_height)
    theme_colors, theme_fonts = theme_info(prs)
    bases: list[SlideBase] = []
    for i, slide in enumerate(prs.slides):
        try:
            base = analyse_slide(i, slide, slide_w, slide_h)
        except Exception:  # noqa: BLE001 - one odd slide must not break template analysis
            continue
        if base is not None:
            bases.append(base)
    mode = design_mode(prs, bases)
    fonts, palette, text_color, background = ({}, [], None, None)
    if mode == "slides":
        fonts, palette, text_color, background = slide_style(prs, bases, theme_colors)
    return SlideDesign(
        slide_width=slide_w,
        slide_height=slide_h,
        design_mode=mode,
        slide_bases=bases,
        slide_fonts=fonts,
        slide_palette=palette,
        slide_text_color=text_color,
        slide_background=background,
        theme_colors=theme_colors,
        theme_fonts=theme_fonts,
    )
