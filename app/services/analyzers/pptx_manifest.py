"""PPTX template manifest: element inventory with roles, layout archetypes, branding, design tokens (Stage 6).

Every shape on every sample slide is described (kind, box, stacking order, text style, fill) and given one role:

- ``structural``     footer, date, slide-number and header placeholders — the slide frame;
- ``placeholder``    empty content placeholders waiting for content;
- ``sample_content`` sample text, tables, charts and photos that must not survive into output;
- ``branding``       logos and brand text: small pictures at the slide edge, or elements repeated at the same
                     place on several slides, on a layout or on the master;
- ``decorative``     the remaining artwork: backgrounds, shapes, lines, illustrations.
"""

from __future__ import annotations

from collections import Counter
from typing import Any

from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from app.models.template_profile import Box, DesignTokens, ElementInfo, FontInfo, SlideInventory
from app.services.analyzers.slide_design import all_text, run_style

_META = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.HEADER}
_PROMPT_TEXT = ("click to add", "click icon to add", "add text", "enter text here")
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _ph_type(shape: Any):
    if not shape.is_placeholder:
        return None
    try:
        return shape.placeholder_format.type
    except (ValueError, AttributeError):
        return None


def _kind(shape: Any) -> str:
    st = shape.shape_type
    if getattr(shape, "has_chart", False):
        return "chart"
    if getattr(shape, "has_table", False):
        return "table"
    if shape.is_placeholder:
        return "placeholder"
    if st == MSO_SHAPE_TYPE.PICTURE:
        return "picture"
    if st == MSO_SHAPE_TYPE.GROUP:
        return "group"
    if st == MSO_SHAPE_TYPE.FREEFORM:
        return "freeform"
    if st == MSO_SHAPE_TYPE.LINE or getattr(shape, "begin_x", None) is not None:
        return "line"
    if st == MSO_SHAPE_TYPE.TEXT_BOX:
        return "text"
    if st == MSO_SHAPE_TYPE.AUTO_SHAPE:
        return "shape"
    return "other"


def _box(shape: Any) -> Box | None:
    if shape.left is None or shape.width is None:
        return None
    return Box(x=int(shape.left), y=int(shape.top), w=int(shape.width), h=int(shape.height))


def _has_image(shape: Any) -> bool:
    return shape.shape_type == MSO_SHAPE_TYPE.PICTURE or bool(shape._element.findall(f".//{_A}blip"))


def _fill(shape: Any) -> str | None:
    node = shape._element.find(f".//{_A}spPr/{_A}solidFill/{_A}srgbClr")
    return node.get("val").upper() if node is not None and node.get("val") else None


def _position_key(shape: Any, sw: int, sh: int) -> tuple:
    b = _box(shape)
    if b is None:
        return ()
    return (_kind(shape), round(b.x / sw, 2), round(b.y / sh, 2), round(b.w / sw, 2), round(b.h / sh, 2))


def _is_logo_like(shape: Any, sw: int, sh: int) -> bool:
    """A small picture in the outer band of the slide (where logos sit)."""
    b = _box(shape)
    if b is None or not _has_image(shape):
        return False
    area = b.area() / (sw * sh)
    edge = b.x < 0.15 * sw or b.right > 0.85 * sw or b.y < 0.15 * sh or b.bottom > 0.85 * sh
    return area < 0.02 and edge


def classify(shape: Any, sw: int, sh: int, repeated: set[tuple]) -> tuple[str, str]:
    """(role, reason) of one sample-slide shape."""
    ptype = _ph_type(shape)
    text = all_text(shape)
    kind = _kind(shape)
    if ptype in _META:
        return "structural", "footer, date, header or slide-number placeholder"
    if kind in ("chart", "table"):
        return "sample_content", f"sample {kind}"
    if shape.is_placeholder:
        if not text or any(p in text.lower() for p in _PROMPT_TEXT):
            if ptype == PP_PLACEHOLDER.PICTURE and _has_image(shape) and shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                return "sample_content", "sample photo in a picture placeholder"
            return "placeholder", "empty placeholder waiting for content"
        return "sample_content", "placeholder holding sample text"
    key = _position_key(shape, sw, sh)
    if text:
        # The same words at the same place on several slides is brand text (a company name); a title at the
        # same place with different words is just the layout's title position.
        if (key, text.strip().lower()) in repeated and len(text) <= 40:
            return "branding", "the same short text at the same place on several slides"
        return "sample_content", "sample text"
    is_picture = shape.shape_type == MSO_SHAPE_TYPE.PICTURE  # logos are pictures; image-filled groups are artwork
    if is_picture and (_is_logo_like(shape, sw, sh) or (key in repeated and _box(shape).area() / (sw * sh) < 0.05)):
        return "branding", "small picture at the slide edge" if _is_logo_like(shape, sw, sh) else "small picture repeated on several slides"
    return "decorative", "artwork without text"


def inventory(prs: Any) -> list[SlideInventory]:
    """Element inventory of every sample slide."""
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    slides = list(prs.slides)
    keys = Counter()
    for slide in slides:
        per_slide = set()
        for s in slide.shapes:
            key = _position_key(s, sw, sh)
            if key:
                per_slide.add(key)
                text = all_text(s).strip().lower()
                if text:
                    per_slide.add((key, text))
        keys.update(per_slide)
    repeated = {k for k, n in keys.items() if n >= max(2, len(slides) // 2)} if len(slides) >= 3 else set()
    out = []
    for i, slide in enumerate(slides):
        elements = []
        for z, shape in enumerate(slide.shapes):
            role, reason = classify(shape, sw, sh, repeated)
            name, color = run_style(shape)
            size = None
            if getattr(shape, "has_text_frame", False):
                sizes = [r.font.size.pt for p in shape.text_frame.paragraphs for r in p.runs if r.font.size]
                size = max(sizes) if sizes else None
            ptype = _ph_type(shape)
            elements.append(ElementInfo(
                shape_index=z, name=shape.name, kind=_kind(shape), role=role, box=_box(shape), z=z,
                placeholder_type=str(ptype).split(".")[-1].split(" ")[0].lower() if ptype is not None else None,
                text_preview=all_text(shape)[:80],
                font=FontInfo(name=name, size_pt=size, color_hex=color) if (name or color or size) else None,
                fill_color=_fill(shape), has_image=_has_image(shape), reason=reason,
            ))
        out.append(SlideInventory(index=i, layout_name=slide.slide_layout.name, elements=elements))
    return out


def layout_branding(prs: Any) -> list[ElementInfo]:
    """Logo-like pictures and brand text on the masters and layouts (they appear on every slide that uses them)."""
    sw, sh = int(prs.slide_width), int(prs.slide_height)
    found: dict[tuple, ElementInfo] = {}
    for master in prs.slide_masters:
        for part in [master, *master.slide_layouts]:
            for z, shape in enumerate(part.shapes):
                if shape.is_placeholder:
                    continue
                text = all_text(shape)
                if _is_logo_like(shape, sw, sh) or (text and len(text) <= 40 and _kind(shape) == "text"):
                    key = _position_key(shape, sw, sh)
                    found.setdefault(key, ElementInfo(
                        shape_index=z, name=shape.name, kind=_kind(shape), role="branding", box=_box(shape), z=z,
                        text_preview=text[:80], has_image=_has_image(shape),
                        reason=f"on {'the master' if part is master else 'layout ' + part.name}",
                    ))
    return list(found.values())


def layout_archetype(name: str, placeholder_types: list[str]) -> str:
    """What kind of slide a layout is for, from its placeholders and name."""
    n = name.lower()
    normal = [t.lower().split(" ")[0] for t in placeholder_types]  # "center_title (3)" → "center_title"
    types = [t for t in normal if t not in ("footer", "date", "slide_number", "header")]
    bodies = sum(1 for t in types if t in ("body", "object", "content", "obj"))
    has_title = any(t in ("title", "center_title", "vertical_title") for t in types)
    has_picture = any(t in ("picture", "pic") for t in types)
    if any(w in n for w in ("thank", "closing", "the end", "contact")):
        return "closing"
    if "quote" in n:
        return "quote"
    if "timeline" in n or "roadmap" in n:
        return "timeline"
    if not types or n.strip() == "blank" or n.endswith(" blank"):
        return "blank"
    if "center_title" in types or "subtitle" in types or "title slide" in n:
        return "title"
    if "section" in n or "header" in n:
        return "section"
    if "comparison" in n or bodies >= 4:
        return "comparison"
    if has_picture or "caption" in n or "picture" in n:
        return "image_text"
    if "two" in n or bodies >= 2:
        return "two_column"
    if has_title and bodies == 1:
        return "bullets"
    if has_title and bodies == 0:
        return "title_only"
    return "other"


def design_tokens(prs: Any, slide_design: Any, theme_colors: dict[str, str], theme_fonts: dict[str, str]) -> DesignTokens:
    """Design tokens: from the artwork when the design lives on the slides, else from the theme."""
    if slide_design is not None and slide_design.design_mode == "slides":
        from app.services.layout.colors import tokens_from_template

        t = tokens_from_template(slide_design)
        return DesignTokens(heading_font=t.heading_font, body_font=t.body_font, primary=t.primary, secondary=t.secondary,
                            accents=list(slide_design.slide_palette), text=t.text, background=t.background, source="slides")
    accents = [theme_colors[k] for k in ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6") if k in theme_colors]
    return DesignTokens(
        heading_font=theme_fonts.get("major"), body_font=theme_fonts.get("minor"),
        primary=accents[0] if accents else None, secondary=accents[1] if len(accents) > 1 else None,
        accents=accents, text=theme_colors.get("dk1"), background=theme_colors.get("lt1"), source="theme",
    )
