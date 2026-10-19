"""DOCX template manifest: sections, page setup, headers/footers, styles by role, numbering, cover, tables (Stage 6).

Text alone does not describe a Word template; this captures the structure a generated document must reuse.
"""

from __future__ import annotations

import re
from typing import Any

from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

from app.models.template_profile import DesignTokens, DocxManifest, DocxSection, DocxStyleInfo

_ROLE_STYLES = {
    "title": ("Title",),
    "subtitle": ("Subtitle",),
    "heading_1": ("Heading 1",),
    "heading_2": ("Heading 2",),
    "heading_3": ("Heading 3",),
    "body": ("Normal", "Body Text"),
    "caption": ("Caption",),
    "list_bullet": ("List Bullet", "List Paragraph"),
    "list_number": ("List Number",),
    "quote": ("Quote", "Intense Quote"),
    "toc_heading": ("TOC Heading",),
}
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _inches(length: Any) -> float | None:
    return round(length.inches, 2) if length is not None else None


def _text_of(container: Any) -> str:
    return " ".join(p.text.strip() for p in container.paragraphs if p.text.strip())[:200]


def _has_page_field(part_element: Any) -> bool:
    xml = part_element.xml if hasattr(part_element, "xml") else ""
    return bool(re.search(r"(instrText[^>]*>\s*PAGE\b|fldSimple[^>]+instr=\"\s*PAGE\b)", xml))


def _style_info(doc: Any, names: tuple[str, ...], theme_fonts: dict[str, str]) -> DocxStyleInfo | None:
    for name in names:
        try:
            style = doc.styles[name]
        except KeyError:
            continue
        font = style.font
        base = style.base_style
        family = font.name or (base.font.name if base is not None else None)
        if family is None:  # theme font reference (majorHAnsi / minorHAnsi)
            rfonts = style.element.find(f"{qn('w:rPr')}/{qn('w:rFonts')}")
            ref = rfonts.get(qn("w:asciiTheme")) if rfonts is not None else None
            if ref:
                family = theme_fonts.get("major" if "major" in ref else "minor")
        size = font.size.pt if font.size else (base.font.size.pt if base is not None and base.font.size else None)
        default_family, default_size = _doc_defaults(doc, theme_fonts)
        if family is None:  # inherited from the document defaults (usually the theme's minor/major font)
            family = theme_fonts.get("major") if name in ("Title", "Heading 1", "Heading 2", "Heading 3") and theme_fonts.get("major") else default_family
        if size is None:
            size = default_size
        color = None
        try:
            if font.color is not None and font.color.type is not None and font.color.rgb is not None:
                color = str(font.color.rgb)
        except (AttributeError, ValueError):
            color = None
        return DocxStyleInfo(style_name=name, font=family, size_pt=size, bold=font.bold, color=color)
    return None


def _doc_defaults(doc: Any, theme_fonts: dict[str, str]) -> tuple[str | None, float | None]:
    """Font family and size every style inherits when it sets none (w:docDefaults)."""
    rpr = doc.styles.element.find(f"{qn('w:docDefaults')}/{qn('w:rPrDefault')}/{qn('w:rPr')}")
    if rpr is None:
        return theme_fonts.get("minor"), None
    family = None
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is not None:
        family = rfonts.get(qn("w:ascii"))
        ref = rfonts.get(qn("w:asciiTheme"))
        if family is None and ref:
            family = theme_fonts.get("major" if "major" in ref else "minor")
    sz = rpr.find(qn("w:sz"))
    size = int(sz.get(qn("w:val"))) / 2 if sz is not None and sz.get(qn("w:val")) else None
    return family or theme_fonts.get("minor"), size


def _theme(doc: Any) -> tuple[dict[str, str], dict[str, str]]:
    from lxml import etree

    for rel in doc.part.rels.values():
        if rel.reltype.endswith("/theme"):
            root = etree.fromstring(rel.target_part.blob)
            colors: dict[str, str] = {}
            scheme = root.find(f".//{_A}clrScheme")
            for node in scheme if scheme is not None else []:
                if len(node):
                    val = node[0].get("val") if node[0].tag == f"{_A}srgbClr" else node[0].get("lastClr")
                    if val:
                        colors[node.tag.split("}")[-1]] = val.upper()
            fonts = {}
            for key, tag in (("major", "majorFont"), ("minor", "minorFont")):
                latin = root.find(f".//{_A}{tag}/{_A}latin")
                if latin is not None and latin.get("typeface"):
                    fonts[key] = latin.get("typeface")
            return colors, fonts
    return {}, {}


def docx_manifest(doc: Any) -> DocxManifest:
    theme_colors, theme_fonts = _theme(doc)
    sections = []
    for sec in doc.sections:
        sections.append(DocxSection(
            orientation="landscape" if sec.orientation == WD_ORIENT.LANDSCAPE else "portrait",
            page_width_in=_inches(sec.page_width), page_height_in=_inches(sec.page_height),
            margins_in={k: v for k, v in {
                "top": _inches(sec.top_margin), "bottom": _inches(sec.bottom_margin),
                "left": _inches(sec.left_margin), "right": _inches(sec.right_margin)}.items() if v is not None},
            header_text=_text_of(sec.header) if not sec.header.is_linked_to_previous or not sections else "",
            footer_text=_text_of(sec.footer) if not sec.footer.is_linked_to_previous or not sections else "",
            has_page_number=_has_page_field(sec.footer._element) or _has_page_field(sec.header._element),
            different_first_page=bool(sec.different_first_page_header_footer),
        ))
    styles = {role: info for role, names in _ROLE_STYLES.items() if (info := _style_info(doc, names, theme_fonts))}

    body = doc.element.body
    cover_lines: list[str] = []
    cover_images = 0
    seen_heading = False
    list_styles: set[str] = set()
    images = 0
    paragraphs = 0
    for p in doc.paragraphs:
        style = p.style.name if p.style is not None else ""
        has_image = bool(p._p.findall(f".//{_A}blip"))
        images += int(has_image)
        if p.text.strip():
            paragraphs += 1
        if style.startswith("Heading"):
            seen_heading = True  # the cover is everything before the first heading
        if not seen_heading:
            if p.text.strip() and len(cover_lines) < 8:
                cover_lines.append(p.text.strip()[:120])
            cover_images += int(has_image)
        if (p._p.pPr is not None and p._p.pPr.numPr is not None) or "List" in style:
            list_styles.add(style or "numbered")
    tables = []
    for t in doc.tables:
        first = t.rows[0] if t.rows else None
        header_bold = bool(first) and all(
            any(r.bold for p in c.paragraphs for r in p.runs) for c in first.cells if c.text.strip()
        ) if first else False
        tables.append({"rows": len(t.rows), "columns": len(t.columns), "style": t.style.name if t.style is not None else None,
                       "header_row": header_bold})
    numbering = 0
    try:
        numbering = len(doc.part.numbering_part.element.findall(qn("w:abstractNum")))
    except Exception:  # no numbering part
        numbering = 0
    has_toc = bool(re.search(r"TOC\s+\\\\[ohzu]", body.xml)) or "TOC Heading" in {p.style.name for p in doc.paragraphs if p.style is not None}
    return DocxManifest(
        sections=sections, styles_by_role=styles, list_styles_used=sorted(list_styles), numbering_definitions=numbering,
        cover_lines=cover_lines, cover_images=cover_images, tables=tables, has_toc=has_toc, body_paragraphs=paragraphs,
        images=images, theme_colors=theme_colors, theme_fonts=theme_fonts,
    )


def docx_tokens(manifest: DocxManifest) -> DesignTokens:
    styles = manifest.styles_by_role
    head = styles.get("heading_1") or styles.get("title")
    body = styles.get("body")
    tc = manifest.theme_colors
    accents = [tc[k] for k in ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6") if k in tc]
    primary = (head.color if head and head.color and head.color != "000000" else None) or (accents[0] if accents else None)
    return DesignTokens(
        heading_font=(head.font if head else None) or manifest.theme_fonts.get("major"),
        body_font=(body.font if body else None) or manifest.theme_fonts.get("minor"),
        primary=primary, secondary=accents[1] if len(accents) > 1 else None, accents=accents,
        text=(body.color if body and body.color else None) or tc.get("dk1"), background=tc.get("lt1"), source="styles",
    )
