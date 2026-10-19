"""XLSX template manifest: sheets and their roles, tables, defined names, formulas and dependencies, formats,
conditional formatting, validations, charts (Stage 6, FR-TPL-6)."""

from __future__ import annotations

import re
from typing import Any

from openpyxl.utils import get_column_letter

from app.models.template_profile import DesignTokens, XlsxManifest, XlsxSheetInfo

_ROLE_NAMES = (
    ("readme", r"read\s*me|about|instructions?|notes|cover|guide"),
    ("assumptions", r"assumptions?|inputs?|drivers?|parameters?|settings"),
    ("sources", r"sources?|references?|citations?"),
    ("report", r"dashboard|summary|report|overview|kpis?|charts?|output"),
    ("calculations", r"calc|calculations?|model|forecast|projections?|workings?|p&l|cash\s*flow"),
    ("data", r"data|raw|input\s*data|records|transactions|list"),
)
_SHEET_REF = re.compile(r"(?:'([^']+)'|\b([A-Za-z_][\w.]*))!\$?[A-Z]{1,3}\$?\d+")
_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"


def _role(name: str, formula_cells: int, value_cells: int, charts: int) -> str:
    for role, pattern in _ROLE_NAMES:
        if re.search(pattern, name, re.I):
            return role
    if charts:
        return "report"
    if formula_cells and formula_cells >= value_cells * 0.3:
        return "calculations"
    return "data" if value_cells else "other"


def _chart_info(chart: Any) -> dict[str, Any]:
    title = None
    try:
        tx = chart.title.tx.rich if chart.title and chart.title.tx else None
        title = "".join(r.t for p in tx.p for r in (p.r or [])) if tx else None
    except AttributeError:
        title = None
    refs = []
    for s in getattr(chart, "series", []) or []:
        try:
            refs.append(s.val.numRef.f)
        except AttributeError:
            continue
    return {"type": type(chart).__name__.replace("Chart", "").lower(), "title": title, "series": refs}


def xlsx_manifest(wb: Any) -> XlsxManifest:
    sheets = []
    names = set(wb.sheetnames)
    for ws in wb.worksheets:
        formulas, values, number_formats = [], 0, set()
        refs: set[str] = set()
        for row in ws.iter_rows():
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                if isinstance(v, str) and v.startswith("="):
                    formulas.append(f"{cell.coordinate}: {v[:80]}")
                    for quoted, bare in _SHEET_REF.findall(v):
                        sheet = quoted or bare
                        if sheet in names and sheet != ws.title:
                            refs.add(sheet)
                else:
                    values += 1
                if cell.number_format and cell.number_format != "General":
                    number_formats.add(cell.number_format)
        header: list[str] = []
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10), values_only=True):
            texts = [str(v) for v in row if v is not None]
            if len(texts) >= 2 and all(not str(t).startswith("=") for t in texts):
                header = texts[:20]
                break
        charts = [_chart_info(c) for c in getattr(ws, "_charts", [])]
        sheets.append(XlsxSheetInfo(
            name=ws.title,
            role=_role(ws.title, len(formulas), values, len(charts)),
            visible=ws.sheet_state == "visible",
            used_range=ws.dimensions,
            max_row=ws.max_row,
            max_col=ws.max_column,
            header=header,
            freeze_panes=ws.freeze_panes if isinstance(ws.freeze_panes, str) else None,
            tables=[{"name": t.name, "ref": t.ref, "columns": [c.name for c in t.tableColumns]} for t in ws.tables.values()],
            formula_count=len(formulas),
            formula_samples=formulas[:5],
            references_sheets=sorted(refs),
            number_formats=sorted(number_formats)[:15],
            conditional_formats=sum(len(rule.rules) for rule in ws.conditional_formatting),
            data_validations=len(ws.data_validations.dataValidation) if ws.data_validations else 0,
            merged_ranges=len(ws.merged_cells.ranges),
            charts=charts,
            images=len(getattr(ws, "_images", [])),
        ))
    defined = {}
    try:
        for name, dn in wb.defined_names.items():
            defined[name] = dn.attr_text
    except AttributeError:  # older openpyxl: a list of DefinedName
        for dn in getattr(wb.defined_names, "definedName", []):
            defined[dn.name] = dn.attr_text
    return XlsxManifest(sheets=sheets, defined_names=defined, theme_colors=_theme_colors(wb))


def _theme_colors(wb: Any) -> dict[str, str]:
    from lxml import etree

    blob = getattr(wb, "loaded_theme", None)
    if not blob:
        return {}
    root = etree.fromstring(blob)
    out: dict[str, str] = {}
    scheme = root.find(f".//{_A}clrScheme")
    for node in scheme if scheme is not None else []:
        if len(node):
            val = node[0].get("val") if node[0].tag == f"{_A}srgbClr" else node[0].get("lastClr")
            if val:
                out[node.tag.split("}")[-1]] = val.upper()
    return out


def xlsx_tokens(wb: Any, manifest: XlsxManifest) -> DesignTokens:
    """Header styling of the first sheet with a header row, else the theme."""
    fill = font = None
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=1, max_row=min(ws.max_row, 10)):
            cells = [c for c in row if c.value is not None]
            if len(cells) >= 2:
                c = cells[0]
                if c.fill is not None and c.fill.fgColor is not None and c.fill.fill_type == "solid":
                    rgb = c.fill.fgColor.rgb
                    fill = rgb[-6:].upper() if isinstance(rgb, str) else None
                font = c.font.name if c.font is not None else None
                break
        if font:
            break
    tc = manifest.theme_colors
    accents = [tc[k] for k in ("accent1", "accent2", "accent3", "accent4", "accent5", "accent6") if k in tc]
    return DesignTokens(heading_font=font, body_font=font, primary=fill or (accents[0] if accents else None),
                        secondary=accents[1] if len(accents) > 1 else None, accents=accents,
                        text=tc.get("dk1"), background=tc.get("lt1"), source="sheet")


def column_letter(n: int) -> str:
    return get_column_letter(n)
