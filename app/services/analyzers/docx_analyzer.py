"""Analyse a DOCX template to extract a TemplateProfileSchema."""

from __future__ import annotations

import logging
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor

from app.llm.schemas import TemplateProfileSchema

logger = logging.getLogger(__name__)


def analyze_docx(path: str | Path) -> TemplateProfileSchema:
    """Extract fonts, sizes, colors, margins, heading styles from a DOCX file."""
    doc = Document(str(path))

    fonts: set[str] = set()
    font_sizes: dict[str, float] = {}
    colors: dict[str, str] = {}
    heading_styles: list[str] = []
    section_patterns: list[str] = []

    # ── Margins from first section ───────────────────────
    margins: dict[str, float] = {}
    if doc.sections:
        sec = doc.sections[0]
        margins = {
            "top": _emu_to_inches(sec.top_margin),
            "bottom": _emu_to_inches(sec.bottom_margin),
            "left": _emu_to_inches(sec.left_margin),
            "right": _emu_to_inches(sec.right_margin),
        }

    # ── Walk paragraphs ──────────────────────────────────
    for para in doc.paragraphs:
        style_name = para.style.name if para.style else ""

        # Track heading styles
        if style_name.startswith("Heading"):
            if style_name not in heading_styles:
                heading_styles.append(style_name)
            section_patterns.append(style_name)

        # Collect fonts and sizes from runs
        for run in para.runs:
            if run.font.name:
                fonts.add(run.font.name)
            if run.font.size:
                pt = run.font.size.pt
                font_sizes[style_name or "body"] = pt
            if run.font.color and run.font.color.rgb:
                colors[style_name or "body"] = f"#{run.font.color.rgb}"

    # ── Detect bullet style ──────────────────────────────
    bullet_style = "•"
    for para in doc.paragraphs:
        text = para.text.strip()
        if text and text[0] in ("•", "–", "-", "▪", "►", "○"):
            bullet_style = text[0]
            break

    # ── Detect line spacing ──────────────────────────────
    line_spacing = 1.15
    for para in doc.paragraphs:
        if para.paragraph_format.line_spacing:
            ls = para.paragraph_format.line_spacing
            if isinstance(ls, (int, float)):
                line_spacing = float(ls)
            else:
                line_spacing = ls.pt / 12.0 if hasattr(ls, "pt") else 1.15
            break

    # ── Detect tone from content ─────────────────────────
    all_text = " ".join(p.text for p in doc.paragraphs[:10])
    tone = _detect_tone(all_text)

    return TemplateProfileSchema(
        fonts=sorted(fonts),
        font_sizes=font_sizes,
        colors=colors,
        heading_styles=heading_styles,
        margins=margins,
        tone=tone,
        section_patterns=section_patterns,
        bullet_style=bullet_style,
        line_spacing=line_spacing,
    )


def _emu_to_inches(emu) -> float:
    """Convert EMU (English Metric Units) to inches."""
    if emu is None:
        return 1.0
    return round(emu / 914400, 2)


def _detect_tone(text: str) -> str:
    """Simple heuristic tone detection from sample text."""
    text_lower = text.lower()
    if any(w in text_lower for w in ("hereby", "pursuant", "whereas", "thereof")):
        return "formal/legal"
    if any(w in text_lower for w in ("we're excited", "awesome", "check out", "!")):
        return "casual/marketing"
    if any(w in text_lower for w in ("methodology", "hypothesis", "findings", "abstract")):
        return "academic"
    return "professional"
