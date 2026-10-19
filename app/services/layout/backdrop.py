"""What colour is behind a box on a slide: solid fills and picture pixels, groups included.

Used by the template analyzer (is a content area light or dark?) and by the contrast check, which would
otherwise treat text over artwork as text over the white slide background.
"""

from __future__ import annotations

import contextlib
import io
from functools import lru_cache

from pptx.enum.dml import MSO_FILL
from pptx.enum.shapes import MSO_SHAPE_TYPE

_A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
_R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"

Box = tuple[float, float, float, float]  # x, y, w, h in EMU


def _leaves(shapes, transform=(0.0, 0.0, 1.0, 1.0)):
    """(shape, absolute box) for every non-group shape, in z-order; group child coordinates resolved."""
    ox, oy, sx, sy = transform
    for shape in shapes:
        if shape.left is None or shape.width is None:
            continue
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            xfrm = shape._element.grpSpPr.find(f"{_A}xfrm")
            if xfrm is None:
                continue
            off, ext = xfrm.find(f"{_A}off"), xfrm.find(f"{_A}ext")
            ch_off, ch_ext = xfrm.find(f"{_A}chOff"), xfrm.find(f"{_A}chExt")
            if None in (off, ext, ch_off, ch_ext) or not int(ch_ext.get("cx")) or not int(ch_ext.get("cy")):
                continue
            gx = int(ext.get("cx")) / int(ch_ext.get("cx"))
            gy = int(ext.get("cy")) / int(ch_ext.get("cy"))
            # child -> group space: off + (c - chOff) * g ; then group -> parent via the outer transform
            inner = (
                ox + sx * (int(off.get("x")) - int(ch_off.get("x")) * gx),
                oy + sy * (int(off.get("y")) - int(ch_off.get("y")) * gy),
                sx * gx,
                sy * gy,
            )
            yield from _leaves(shape.shapes, inner)
        else:
            yield shape, (ox + sx * shape.left, oy + sy * shape.top, sx * shape.width, sy * shape.height)


def pixels_of(img):
    """Pixel tuples; Pillow 12 renamed `getdata` to `get_flattened_data`."""
    getter = getattr(img, "get_flattened_data", None)
    return getter() if getter is not None else img.getdata()


@lru_cache(maxsize=64)
def _thumb(blob: bytes):
    from PIL import Image

    img = Image.open(io.BytesIO(blob)).convert("RGBA")
    img.thumbnail((256, 256))
    return img


def _image_color(blob: bytes, shape_box: Box, area: Box) -> str | None:
    """Average opaque colour of the part of a stretched picture that lies under ``area``."""
    try:
        img = _thumb(blob)
    except Exception:  # noqa: BLE001 - unreadable media: unknown colour
        return None
    x, y, w, h = shape_box
    if w <= 0 or h <= 0:
        return None
    ax0, ay0 = max(area[0], x), max(area[1], y)
    ax1, ay1 = min(area[0] + area[2], x + w), min(area[1] + area[3], y + h)
    if ax1 <= ax0 or ay1 <= ay0:
        return None
    iw, ih = img.size
    crop = img.crop(
        (
            int((ax0 - x) / w * iw),
            int((ay0 - y) / h * ih),
            max(int((ax1 - x) / w * iw), int((ax0 - x) / w * iw) + 1),
            max(int((ay1 - y) / h * ih), int((ay0 - y) / h * ih) + 1),
        )
    )
    pixels = [p for p in pixels_of(crop) if p[3] > 128]
    if len(pixels) < 0.5 * crop.size[0] * crop.size[1]:
        return None  # mostly transparent here: whatever is below shows through
    n = len(pixels)
    return "".join(f"{round(sum(p[i] for p in pixels) / n):02X}" for i in range(3))


def _solid(shape) -> str | None:
    try:
        if shape.fill.type == MSO_FILL.SOLID:
            return str(shape.fill.fore_color.rgb)
    except (AttributeError, TypeError, ValueError):
        return None
    return None


def _blob(shape, slide) -> bytes | None:
    blip = next(shape._element.iter(f"{_A}blip"), None)
    rid = blip.get(_R_EMBED) if blip is not None else None
    if not rid:
        return None
    try:
        return slide.part.related_part(rid).blob
    except KeyError:
        return None


_SCHEME = {"tx1": "dk1", "bg1": "lt1", "tx2": "dk2", "bg2": "lt2"}


@lru_cache(maxsize=16)
def _theme_colors(blob: bytes) -> dict[str, str]:
    from lxml import etree

    scheme = etree.fromstring(blob).find(f".//{_A}clrScheme")
    out: dict[str, str] = {}
    for node in scheme if scheme is not None else []:
        if len(node):
            val = node[0].get("val") if node[0].tag == f"{_A}srgbClr" else node[0].get("lastClr")
            if val:
                out[node.tag.split("}")[-1]] = val.upper()
    return out


def _scheme_color(part, name: str | None) -> str | None:
    """Resolve a theme colour reference (accent1, bg1, ...) through the master's theme."""
    from pptx.opc.constants import RELATIONSHIP_TYPE as RT

    if not name:
        return None
    try:
        pres = part.part.package.presentation_part.presentation
        theme = pres.slide_master.part.part_related_by(RT.THEME)
        return _theme_colors(theme.blob).get(_SCHEME.get(name, name))
    except Exception:  # noqa: BLE001 - unreadable theme: unknown colour
        return None


def _background(slide, area: Box) -> str | None:
    """Slide background, inherited from the layout or master when the slide has none (solid or picture)."""
    parts = [slide]
    with contextlib.suppress(AttributeError):  # layouts and masters have no further parents
        parts += [slide.slide_layout, slide.slide_layout.slide_master]
    for part in parts:
        bg = part._element.cSld.bg
        if bg is None:
            continue
        solid = bg.find(f".//{_A}solidFill/{_A}srgbClr")
        if solid is not None:
            return solid.get("val").upper()
        scheme = bg.find(f".//{_A}solidFill/{_A}schemeClr")
        if scheme is not None:
            return _scheme_color(part, scheme.get("val"))
        blip = next(bg.iter(f"{_A}blip"), None)
        rid = blip.get(_R_EMBED) if blip is not None else None
        if not rid:
            return None
        try:
            blob = part.part.related_part(rid).blob
        except KeyError:
            return None
        pres = part.part.package.presentation_part.presentation
        return _image_color(blob, (0, 0, int(pres.slide_width), int(pres.slide_height)), area)
    return None


def color_under(slide, area: Box, *, stop_at=None, default: str | None = "FFFFFF") -> str | None:
    """Colour behind ``area`` (EMU box): topmost fill or picture covering its centre, below ``stop_at``."""
    result = _background(slide, area) or default
    cx, cy = area[0] + area[2] / 2, area[1] + area[3] / 2
    for shape, (x, y, w, h) in _leaves(slide.shapes):
        if stop_at is not None and shape._element is stop_at:
            break
        if not (x <= cx <= x + w and y <= cy <= y + h):
            continue
        color = _solid(shape)
        if color is None:
            blob = _blob(shape, slide)
            if blob is not None:
                color = _image_color(blob, (x, y, w, h), area)
        if color:
            result = color
    return result


def is_dark(hex_color: str | None) -> bool:
    if not hex_color:
        return False
    r, g, b = (int(hex_color[i : i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b < 0.45
