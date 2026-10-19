"""File checks for DOCX/XLSX and the per-version quality report (FR-VAL-1..5, AC-2, AC-4, AC-10; ADR-032).

Every delivered file is re-opened from disk. The report separates findings that were detected, corrected
(deterministically or by the LLM) and left unresolved, and lists checks that were not performed with the reason.
"""

from __future__ import annotations

import shutil
import subprocess
import tempfile
from pathlib import Path
from typing import Any

from app.agents.validator import PLACEHOLDER_REGEX, ValidationIssue
from app.models.quality import EvidenceSummary, Finding, NotPerformed, QualityReport, Repair

SEVERITY = {"error": "critical", "warning": "minor"}


def validate_docx_file(path: str | Path) -> list[ValidationIssue]:
    import docx

    try:
        doc = docx.Document(str(path))
    except Exception as exc:
        return [ValidationIssue(severity="error", where="docx:file", message=f"Document does not open: {exc}")]
    issues = []
    paras = doc.paragraphs
    for i, p in enumerate(paras):
        if p.style.name.startswith("Heading") and not p.text.strip():
            issues.append(ValidationIssue(severity="warning", where="docx:file", message=f"Empty heading at paragraph {i + 1}."))
        if PLACEHOLDER_REGEX.search(p.text):
            issues.append(ValidationIssue(severity="error", where="docx:file", message=f"Placeholder text: '{p.text[:60]}'."))
    for t in doc.tables:
        if not any(c.text.strip() for row in t.rows for c in row.cells):
            issues.append(ValidationIssue(severity="warning", where="docx:file", message="Empty table."))
    return issues


def validate_xlsx_file(path: str | Path, checks: list[Any] | None = None) -> tuple[list[ValidationIssue], Any]:
    from app.services.xlsx_check import recalculate

    try:
        report = recalculate(path, checks)
    except Exception as exc:
        return [ValidationIssue(severity="error", where="xlsx:file", message=f"Workbook does not open: {exc}")], None
    issues = [ValidationIssue(severity="error", where="xlsx:formula", message=f"Formula error {e}") for e in report.errors]
    issues += [ValidationIssue(severity="error", where="xlsx:check",
                               message=f"Expected {c.expected:g} at {c.ref}, got {c.actual}") for c in report.checks if not c.ok]
    return issues, report


def pdf_conversion(path: str | Path) -> tuple[bool | None, str]:
    """LibreOffice round trip when available: (converted?, detail); None when LibreOffice is not installed."""
    soffice = shutil.which("soffice") or shutil.which("libreoffice")
    if not soffice:
        return None, "LibreOffice is not installed"
    with tempfile.TemporaryDirectory() as out:
        try:
            subprocess.run([soffice, "--headless", "--convert-to", "pdf", "--outdir", out, str(path)],
                           capture_output=True, timeout=120, check=False)
        except Exception as exc:
            return False, str(exc)
        return any(Path(out).glob("*.pdf")), "converted" if any(Path(out).glob("*.pdf")) else "conversion failed"


def build_quality_report(fmt: str, path: str | Path, state: dict[str, Any], *, file_issues: list[ValidationIssue] | None = None,
                         recalc: Any = None, offline: bool = False) -> QualityReport:
    findings: list[Finding] = []
    for issue in file_issues or []:
        findings.append(Finding(check_id=f"{issue.where.split(':')[0]}.file", tier=2, severity=SEVERITY.get(issue.severity, "minor"),
                                status="unresolved" if issue.severity == "error" else "detected", unit=issue.where,
                                message=issue.message))
    for issue in (state.get("validation") or {}).get("issues", []):
        if str(issue.get("where", "")).startswith(fmt) and not any(f.message == issue.get("message") for f in findings):
            findings.append(Finding(check_id=f"{fmt}.content", tier=3, severity=SEVERITY.get(issue.get("severity"), "minor"),
                                    status="unresolved" if issue.get("severity") == "error" else "detected",
                                    unit=issue.get("where"), message=issue.get("message", "")))
    plan_report = (state.get("plan_reports") or {}).get(fmt) or {}
    for fix in plan_report.get("repairs", []):
        findings.append(Finding(check_id=f"{fmt}.plan", tier=1, severity="minor", status="corrected", message=fix,
                                repair=Repair(action=fix, by="deterministic")))
    if plan_report.get("revised"):
        findings.append(Finding(check_id=f"{fmt}.plan", tier=1, severity="major", status="corrected",
                                message="Plan revised once by the model after validation", repair=Repair(action="revise plan", by="llm")))
    not_performed = []
    converted, detail = pdf_conversion(path)
    if converted is None:
        not_performed.append(NotPerformed(check_id="file.pdf_conversion", tier=2, reason=detail))
    elif not converted:
        findings.append(Finding(check_id="file.pdf_conversion", tier=2, severity="major", status="unresolved", message=detail))
    if fmt in ("docx", "xlsx"):
        not_performed.append(NotPerformed(check_id=f"{fmt}.visual", tier=4, reason="no page renderer for this format; structural checks only"))
    if offline:
        not_performed.append(NotPerformed(check_id="content.llm_review", tier=5, reason="LLM is mocked or unavailable"))
    if recalc is not None and recalc.not_checked:
        not_performed.append(NotPerformed(check_id="xlsx.recalc", tier=2,
                                          reason=f"{len(recalc.not_checked)} formulas use functions the checker does not evaluate"))
    dossier = state.get("dossier") or {}
    evidence = EvidenceSummary(sources_used=len(state.get("registry") or {}), coverage_pct=plan_report.get("claim_coverage_pct"),
                               gaps=list(dossier.get("gaps", []))[:5])
    return QualityReport(findings=findings, not_performed=not_performed, evidence_summary=evidence, iterations=1)
