"""Word finishing: citation styles, table of contents, page numbers, style overrides (FR-DOCX-1..5, ADR-030)."""

from __future__ import annotations

import re
from typing import Any

from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


def _info(sources_map: dict | None, sid: int) -> dict[str, Any]:
    return (sources_map or {}).get(sid) or (sources_map or {}).get(str(sid)) or {}


def _year(info: dict[str, Any]) -> str:
    m = re.search(r"(19|20)\d{2}", str(info.get("published_at") or info.get("date") or ""))
    return m.group(0) if m else "n.d."


def _author(info: dict[str, Any]) -> str:
    authors = info.get("authors") or []
    if authors:
        first = str(authors[0]).split()[-1]
        return first + (" et al." if len(authors) > 2 else (f" and {str(authors[1]).split()[-1]}" if len(authors) == 2 else ""))
    publisher = info.get("publisher") or ""
    if publisher:
        return publisher
    url = info.get("url") or info.get("url_or_filename") or ""
    host = re.sub(r"^https?://(www\.)?", "", url).split("/")[0]
    return host or (info.get("title") or "Anon.")[:40]


def in_text(source_ids: list[int], style: str, sources_map: dict | None) -> str:
    """Citation marker: ' [1][2]' (numeric) or ' (Author, 2024; Other, n.d.)' (APA / Harvard)."""
    ids = sorted(set(source_ids or []))
    if not ids or style == "none":
        return ""
    if style in ("apa", "harvard"):
        sep = ", " if style == "apa" else " "
        return " (" + "; ".join(f"{_author(_info(sources_map, i))}{sep}{_year(_info(sources_map, i))}" for i in ids) + ")"
    return " " + "".join(f"[{i}]" for i in ids)


def reference_entry(sid: int, style: str, sources_map: dict | None) -> str:
    info = _info(sources_map, sid)
    title = info.get("title") or f"Source {sid}"
    url = info.get("url_or_filename") or info.get("url") or ""
    accessed = info.get("accessed_at") or info.get("retrieved_at") or ""
    if style == "apa":
        return f"{_author(info)}. ({_year(info)}). {title}." + (f" {url}" if url else "")
    if style == "harvard":
        tail = f" Available at: {url}" if url else ""
        tail += f" (Accessed: {accessed})" if accessed and url else ""
        return f"{_author(info)} ({_year(info)}) {title}.{tail}"
    acc = f" (accessed {accessed})" if accessed else ""
    return f"[{sid}] {title}" + (f" - {url}" if url else "") + acc


def _field(paragraph: Any, instr: str, placeholder: str = "") -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    text = OxmlElement("w:instrText")
    text.set(qn("xml:space"), "preserve")
    text.text = instr
    sep = OxmlElement("w:fldChar")
    sep.set(qn("w:fldCharType"), "separate")
    run._r.extend([begin, text, sep])
    paragraph.add_run(placeholder)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    paragraph.add_run()._r.append(end)


def add_toc(doc: Any, heading_style: str) -> None:
    """'Contents' heading and a TOC field (Word fills it on open / F9)."""
    h = doc.add_paragraph(style=heading_style)
    h.add_run("Contents")
    p = doc.add_paragraph()
    _field(p, 'TOC \\o "1-3" \\h \\z \\u', "Right-click and choose Update Field to build the table of contents.")
    settings = doc.settings.element
    if settings.find(qn("w:updateFields")) is None:
        upd = OxmlElement("w:updateFields")
        upd.set(qn("w:val"), "true")
        settings.append(upd)


def has_page_field(part_element: Any) -> bool:
    return any("PAGE" in (t.text or "") for t in part_element.iter(qn("w:instrText"))) or bool(
        part_element.xpath(".//w:fldSimple[contains(@w:instr,'PAGE')]"))


def ensure_page_numbers(doc: Any, title: str = "") -> list[str]:
    """Add 'Page N' to footers (and the title to empty headers) of sections that lack them."""
    done = []
    for i, section in enumerate(doc.sections, start=1):
        footer = section.footer
        if not has_page_field(footer._element):
            p = footer.add_paragraph() if footer.paragraphs and footer.paragraphs[0].text else footer.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.add_run("Page ")
            _field(p, "PAGE", "1")
            done.append(f"section {i}: page numbers")
        header = section.header
        if title and not any(p.text.strip() for p in header.paragraphs):
            header.paragraphs[0].text = title
            header.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            for r in header.paragraphs[0].runs:
                r.font.size = Pt(9)
            done.append(f"section {i}: running header")
    return done


def apply_style_overrides(doc: Any, style: Any) -> list[str]:
    """Write requested formatting into named styles (Heading 1/2, Normal), never as direct formatting."""
    if style is None:
        return []
    names = {s.name for s in doc.styles}
    done = []

    def set_(name: str, size: float | None = None, font: str | None = None, color: str | None = None) -> None:
        if name not in names:
            return
        f = doc.styles[name].font
        if size:
            f.size = Pt(size)
        if font:
            f.name = font
            rpr = doc.styles[name].element.get_or_add_rPr()
            fonts = rpr.find(qn("w:rFonts"))
            if fonts is not None:
                for k in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
                    fonts.attrib.pop(qn(k), None)
        if color:
            f.color.rgb = RGBColor.from_string(color)
        done.append(name)

    set_("Heading 1", style.heading_1_size_pt, style.heading_font, style.accent_color)
    set_("Heading 2", style.heading_2_size_pt, style.heading_font, style.accent_color)
    set_("Title", None, style.heading_font, style.accent_color)
    set_("Normal", style.body_size_pt, style.body_font)
    if style.line_spacing and "Normal" in names:
        doc.styles["Normal"].paragraph_format.line_spacing = style.line_spacing
    return done
