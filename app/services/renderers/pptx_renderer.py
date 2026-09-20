"""Render a DeckModel into an editable PPTX file, reusing the template's layouts."""

from __future__ import annotations

import io
import logging
from pathlib import Path

from pptx import Presentation
from pptx.util import Pt

from app.llm.schemas import DeckModel, TemplateProfileSchema

logger = logging.getLogger(__name__)


def render_pptx(
    model: DeckModel,
    profile: TemplateProfileSchema,
    template_path: str | None = None,
) -> bytes:
    """Render a DeckModel to PPTX bytes, reusing template slide layouts."""
    if template_path and Path(template_path).exists():
        prs = Presentation(template_path)
        # Remove existing slides but keep layouts
        while len(prs.slides) > 0:
            rId = prs.slides._sldIdLst[0].get("r:id")  # type: ignore[attr-defined]
            prs.part.drop_rel(rId)
            prs.slides._sldIdLst.remove(prs.slides._sldIdLst[0])  # type: ignore[attr-defined]
    else:
        prs = Presentation()

    layout_map = {layout.name: layout for layout in prs.slide_layouts}

    # ── Add slides ───────────────────────────────────────
    for slide_content in model.slides:
        layout = _find_layout(layout_map, slide_content.layout_name, prs)
        slide = prs.slides.add_slide(layout)

        # Populate placeholders
        for ph in slide.placeholders:
            ph_name = ph.name
            ph_idx = str(ph.placeholder_format.idx)

            # Try matching by name first, then by idx
            content = (
                slide_content.placeholders.get(ph_name)
                or slide_content.placeholders.get(ph_idx)
                or slide_content.placeholders.get(f"placeholder_{ph_idx}")
            )
            if content:
                ph.text = content

        # Speaker notes
        if slide_content.speaker_notes:
            if slide.has_notes_slide:
                notes_slide = slide.notes_slide
            else:
                notes_slide = slide.notes_slide  # creates it
            notes_slide.notes_text_frame.text = slide_content.speaker_notes

    # ── Sources slide ────────────────────────────────────
    if model.sources_slide:
        layout = _find_layout(layout_map, "Blank", prs)
        slide = prs.slides.add_slide(layout)
        # Add a text box for sources
        from pptx.util import Inches
        left, top, width, height = Inches(0.5), Inches(0.5), Inches(9), Inches(6)
        txBox = slide.shapes.add_textbox(left, top, width, height)
        tf = txBox.text_frame
        tf.word_wrap = True

        title_para = tf.paragraphs[0]
        title_para.text = "Sources"
        title_para.font.size = Pt(24)
        title_para.font.bold = True

        for src in model.sources_slide:
            para = tf.add_paragraph()
            sid = src.get("id", "")
            label = src.get("label", "")
            url = src.get("url", "")
            para.text = f"[{sid}] {label} – {url}"
            para.font.size = Pt(10)

    # ── Serialize ────────────────────────────────────────
    buf = io.BytesIO()
    prs.save(buf)
    return buf.getvalue()


def _find_layout(layout_map: dict, name: str, prs: Presentation):
    """Find a layout by name with fuzzy fallback."""
    if name in layout_map:
        return layout_map[name]

    # Fuzzy match
    name_lower = name.lower()
    for lname, layout in layout_map.items():
        if name_lower in lname.lower():
            return layout

    # Fallback to first layout
    logger.warning("Layout '%s' not found, using first available", name)
    return prs.slide_layouts[0]
