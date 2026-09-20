"""Render a DocumentModel into an editable DOCX file, respecting the template's styles."""

from __future__ import annotations

import io
import logging
from pathlib import Path

from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from app.llm.schemas import DocumentModel, TemplateProfileSchema

logger = logging.getLogger(__name__)

# Heading level → python-docx style name mapping
_HEADING_STYLE = {1: "Heading 1", 2: "Heading 2", 3: "Heading 3", 4: "Heading 4"}


def render_docx(
    model: DocumentModel,
    profile: TemplateProfileSchema,
    template_path: str | None = None,
) -> bytes:
    """Render a DocumentModel to DOCX bytes, reusing template styles when available."""
    if template_path and Path(template_path).exists():
        doc = Document(template_path)
        # Clear existing content but keep styles
        for _ in range(len(doc.paragraphs)):
            p = doc.paragraphs[0]
            p._element.getparent().remove(p._element)
    else:
        doc = Document()

    # ── Apply margins ────────────────────────────────────
    if profile.margins:
        for section in doc.sections:
            if "top" in profile.margins:
                section.top_margin = Inches(profile.margins["top"])
            if "bottom" in profile.margins:
                section.bottom_margin = Inches(profile.margins["bottom"])
            if "left" in profile.margins:
                section.left_margin = Inches(profile.margins["left"])
            if "right" in profile.margins:
                section.right_margin = Inches(profile.margins["right"])

    # ── Title ────────────────────────────────────────────
    title_para = doc.add_heading(model.title, level=0)
    _apply_font(title_para, profile, "Heading 0")

    # ── Sections ─────────────────────────────────────────
    for section in model.sections:
        level = min(section.heading_level, 4)
        heading = doc.add_heading(section.heading, level=level)
        _apply_font(heading, profile, _HEADING_STYLE.get(level, "Heading 1"))

        # Body text – split by newlines for paragraph breaks
        for line in section.body.split("\n"):
            line = line.strip()
            if not line:
                continue
            para = doc.add_paragraph(line)
            _apply_font(para, profile, "body")
            # Apply line spacing
            para.paragraph_format.line_spacing = profile.line_spacing

    # ── Sources section ──────────────────────────────────
    if model.sources_section:
        doc.add_heading("Sources", level=1)
        for src in model.sources_section:
            label = src.get("label", "")
            url = src.get("url", "")
            sid = src.get("id", "")
            para = doc.add_paragraph(f"[{sid}] {label}")
            if url:
                run = para.add_run(f" – {url}")
                run.font.size = Pt(9)
                run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)

    # ── Serialize ────────────────────────────────────────
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _apply_font(para, profile: TemplateProfileSchema, style_key: str) -> None:
    """Apply profile fonts and colors to a paragraph's runs."""
    primary_font = profile.fonts[0] if profile.fonts else None
    color_hex = profile.colors.get(style_key) or profile.colors.get("body")
    size_pt = profile.font_sizes.get(style_key)

    for run in para.runs:
        if primary_font:
            run.font.name = primary_font
        if size_pt:
            run.font.size = Pt(size_pt)
        if color_hex and color_hex.startswith("#"):
            hex_clean = color_hex.lstrip("#")
            if len(hex_clean) == 6:
                r, g, b = int(hex_clean[:2], 16), int(hex_clean[2:4], 16), int(hex_clean[4:6], 16)
                run.font.color.rgb = RGBColor(r, g, b)
