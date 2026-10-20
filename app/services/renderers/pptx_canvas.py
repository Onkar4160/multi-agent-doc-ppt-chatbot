"""Low-level, deterministic drawing helpers over python-pptx (ported from DocForge ``render/pptx/canvas.py``).

Naming contract (read by the validator): shapes carrying content are named ``content:<role>``,
purely decorative shapes ``decor:<role>``. Overlap checks only consider content shapes.
"""

from __future__ import annotations

from dataclasses import dataclass, field

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Pt

from app.services.layout.text import EMU_PER_INCH, Para, fit_paragraphs, measure

INSET = int(0.05 * EMU_PER_INCH)


def inch(v: float) -> int:
    return int(round(v * EMU_PER_INCH))


def rgb(hex_color: str) -> RGBColor:
    return RGBColor.from_string(hex_color.upper())


@dataclass(frozen=True)
class Rect:
    x: float
    y: float
    w: float
    h: float

    @property
    def right_x(self) -> float:
        return self.x + self.w

    @property
    def bottom_y(self) -> float:
        return self.y + self.h

    def inset(self, dx: float, dy: float | None = None) -> Rect:
        dy = dx if dy is None else dy
        return Rect(self.x + dx, self.y + dy, self.w - 2 * dx, self.h - 2 * dy)

    def split_h(self, n: int, gap: float) -> list[Rect]:
        w = (self.w - gap * (n - 1)) / n
        return [Rect(self.x + i * (w + gap), self.y, w, self.h) for i in range(n)]

    def split_v(self, n: int, gap: float) -> list[Rect]:
        h = (self.h - gap * (n - 1)) / n
        return [Rect(self.x, self.y + i * (h + gap), self.w, h) for i in range(n)]

    def take_top(self, h: float, gap: float = 0.0) -> tuple[Rect, Rect]:
        return Rect(self.x, self.y, self.w, h), Rect(self.x, self.y + h + gap, self.w, self.h - h - gap)

    def take_left(self, w: float, gap: float = 0.0) -> tuple[Rect, Rect]:
        return Rect(self.x, self.y, w, self.h), Rect(self.x + w + gap, self.y, self.w - w - gap, self.h)


@dataclass(frozen=True)
class TextStyle:
    color: str
    align: str = "left"  # left | center | right
    anchor: str = "top"  # top | middle | bottom
    italic: bool = False
    bullet_color: str | None = None


@dataclass
class BulletPara:
    """A paragraph spec with an optional bullet level (None = no bullet)."""

    para: Para
    level: int | None = None
    color: str | None = None
    italic: bool = False


@dataclass
class FitIssue:
    shape_name: str
    detail: str


@dataclass
class CanvasLog:
    fit_issues: list[FitIssue] = field(default_factory=list)
    estimated_measurement: bool = False


def content_height(paras: list[BulletPara], width_in: float) -> float:
    """Height in inches (including insets) that ``paras`` need at their nominal size in a box this wide."""
    indent = inch(BULLET_INDENT_IN)
    specs = [
        Para(**{**bp.para.__dict__, "indent_emu": indent * (bp.level + 1) if bp.level is not None else 0})
        for bp in paras
    ]
    return (measure(specs, inch(width_in) - 2 * INSET).height_emu + 2 * INSET) / EMU_PER_INCH


_ALIGN = {"left": PP_ALIGN.LEFT, "center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}
_ANCHOR = {"top": MSO_ANCHOR.TOP, "middle": MSO_ANCHOR.MIDDLE, "bottom": MSO_ANCHOR.BOTTOM}
BULLET_INDENT_IN = 0.28


class Canvas:
    def __init__(self, slide, log: CanvasLog, scale: float = 1.0) -> None:
        self.slide = slide
        self.log = log
        self.scale = scale  # minimum sizes are defined for a 7.5-inch-tall slide

    # ---------------------------------------------------------------- shapes

    def background(self, color: str) -> None:
        fill = self.slide.background.fill
        fill.solid()
        fill.fore_color.rgb = rgb(color)

    def shape(self, kind, r: Rect, *, fill: str | None, name: str, line: str | None = None):
        shp = self.slide.shapes.add_shape(kind, inch(r.x), inch(r.y), inch(r.w), inch(r.h))
        shp.name = name
        if fill is None:
            shp.fill.background()
        else:
            shp.fill.solid()
            shp.fill.fore_color.rgb = rgb(fill)
        if line is None:
            shp.line.fill.background()
        else:
            shp.line.color.rgb = rgb(line)
            shp.line.width = Pt(1)
        shp.shadow.inherit = False
        if shp.has_text_frame:
            shp.text_frame.text = ""
        return shp

    def rect(self, r: Rect, fill: str, name: str, *, rounded: bool = False):
        shp = self.shape(
            MSO_SHAPE.ROUNDED_RECTANGLE if rounded else MSO_SHAPE.RECTANGLE, r, fill=fill, name=name
        )
        if rounded:
            shp.adjustments[0] = 0.08
        return shp

    def circle(self, cx: float, cy: float, d: float, fill: str, name: str):
        return self.shape(MSO_SHAPE.OVAL, Rect(cx - d / 2, cy - d / 2, d, d), fill=fill, name=name)

    def line(self, x1: float, y1: float, x2: float, y2: float, color: str, name: str, width_pt: float = 2.0):
        conn = self.slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, inch(x1), inch(y1), inch(x2), inch(y2))
        conn.name = name
        conn.line.color.rgb = rgb(color)
        conn.line.width = Pt(width_pt)
        return conn

    def badge(
        self, cx: float, cy: float, d: float, label: str, fill: str, text_color: str, font: str, name: str
    ):
        """Filled circle with a centred label (numbers in circles are the deck's visual motif)."""
        shp = self.circle(cx, cy, d, fill, name)
        tf = shp.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = False
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        run = p.add_run()
        run.text = label
        run.font.size = Pt(max(10, round(d * 72 * 0.42)))
        run.font.bold = True
        run.font.name = font
        run.font.color.rgb = rgb(text_color)
        return shp

    # ---------------------------------------------------------------- text

    def text(
        self,
        r: Rect,
        paras: list[BulletPara],
        style: TextStyle,
        *,
        name: str,
        min_pt: float,
    ):
        """Text box whose paragraphs are shrunk (proportionally) to fit; misses are logged, not hidden."""
        box = self.slide.shapes.add_textbox(inch(r.x), inch(r.y), inch(r.w), inch(r.h))
        box.name = name
        tf = box.text_frame
        tf.word_wrap = True
        tf.auto_size = MSO_AUTO_SIZE.NONE
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(INSET)
        tf.vertical_anchor = _ANCHOR[style.anchor]

        indent = inch(BULLET_INDENT_IN)
        specs = [
            Para(**{**bp.para.__dict__, "indent_emu": indent * (bp.level + 1) if bp.level is not None else 0})
            for bp in paras
        ]
        width = inch(r.w) - 2 * INSET
        height = inch(r.h) - 2 * INSET
        fitted, ok = fit_paragraphs(specs, width, height, min_pt=min_pt * self.scale)
        if not ok:
            self.log.fit_issues.append(FitIssue(name, "text does not fit at minimum size"))

        for i, (bp, spec) in enumerate(zip(paras, fitted, strict=True)):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = _ALIGN[style.align]
            p.line_spacing = spec.line_spacing
            if spec.space_after_pt and i < len(fitted) - 1:
                p.space_after = Pt(spec.space_after_pt)
            if spec.space_before_pt and i > 0:
                p.space_before = Pt(spec.space_before_pt)
            run = p.add_run()
            run.text = spec.text
            font = run.font
            font.size = Pt(spec.size_pt)
            font.bold = spec.bold
            font.italic = bp.italic or style.italic
            if spec.font:
                font.name = spec.font
            font.color.rgb = rgb(bp.color or style.color)
            if bp.level is not None:
                _set_bullet(p, bp.level, style.bullet_color or bp.color or style.color, indent)
        return box

    def picture(self, r: Rect, path, *, name: str, alt: str):
        """Picture that fills ``r`` exactly, cropping the excess (no distortion)."""
        from PIL import Image

        with Image.open(path) as img:
            iw, ih = img.size
        pic = self.slide.shapes.add_picture(str(path), inch(r.x), inch(r.y), inch(r.w), inch(r.h))
        pic.name = name
        box_ratio, img_ratio = r.w / r.h, iw / ih
        if img_ratio > box_ratio:  # too wide: crop left/right
            excess = 1 - box_ratio / img_ratio
            pic.crop_left = pic.crop_right = excess / 2
        elif img_ratio < box_ratio:  # too tall: crop top/bottom (bias towards the top third)
            excess = 1 - img_ratio / box_ratio
            pic.crop_top, pic.crop_bottom = excess * 0.35, excess * 0.65
        pic._element.nvPicPr.cNvPr.set("descr", alt[:500])
        return pic

    def placeholder_text(self, ph, paras: list[BulletPara], *, name: str, min_pt: float):
        """Fill a template placeholder or styled text box, keeping the template's font, colour and alignment.

        Sizes start from the size the template uses for that box (explicit run size, else the inherited
        placeholder size) and shrink only when the text would overflow; the chosen size is written
        explicitly so the validator can measure it. Fixed line spacing is dropped because it would clip
        smaller text.
        """
        from copy import deepcopy

        from pptx.oxml.ns import qn

        ph.name = name
        tf = ph.text_frame
        first_run = next((r for p in tf.paragraphs for r in p.runs), None)
        rpr = first_run._r.find(qn("a:rPr")) if first_run is not None else None
        rpr = deepcopy(rpr) if rpr is not None else None
        ppr = tf.paragraphs[0]._p.find(qn("a:pPr"))
        ppr = deepcopy(ppr) if ppr is not None else None
        if ppr is not None:
            for node in ppr.findall(qn("a:lnSpc")):
                ppr.remove(node)
        base = first_run.font.size.pt if first_run is not None and first_run.font.size else None
        if base is None and ph.is_placeholder:
            base = inherited_size(ph)
        slot_font = first_run.font.name if first_run is not None else None
        slot_bold = bool(first_run is not None and first_run.font.bold)
        tf.word_wrap = True
        if paras:
            factor = base / paras[0].para.size_pt if base else 1.0
            paras = [
                BulletPara(
                    Para(
                        **{
                            **bp.para.__dict__,
                            "size_pt": round(bp.para.size_pt * factor, 1),
                            "font": slot_font or bp.para.font,
                            "bold": bp.para.bold or slot_bold,
                        }
                    ),
                    level=bp.level,
                    color=bp.color,
                    italic=bp.italic,
                )
                for bp in paras
            ]
        indent = inch(BULLET_INDENT_IN)
        specs = [
            Para(**{**bp.para.__dict__, "indent_emu": indent * (bp.level + 1) if bp.level else 0})
            for bp in paras
        ]
        width = int(ph.width) - int(tf.margin_left) - int(tf.margin_right)
        height = int(ph.height) - int(tf.margin_top) - int(tf.margin_bottom)
        from app.services.layout.fonts import has_exact_metrics

        if specs and not has_exact_metrics(specs[0].font):
            width = int(width * 0.94)  # PowerPoint will substitute this font: leave room for wider glyphs
        fitted, ok = fit_paragraphs(specs, width, height, min_pt=min_pt * self.scale)
        # Template boxes are sized to their sample words; grow (downwards) rather than overflow.
        # The 10% margin covers inherited line spacing the measurement cannot see.
        needed = int(measure(fitted, width).height_emu * 1.1) + int(tf.margin_top) + int(tf.margin_bottom)
        if needed > int(ph.height):
            set_geometry(ph, h=needed)
        tf.clear()
        for i, (bp, spec) in enumerate(zip(paras, fitted, strict=True)):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            if ppr is not None:
                old = p._p.find(qn("a:pPr"))
                if old is not None:
                    p._p.remove(old)
                p._p.insert(0, deepcopy(ppr))
            p.level = bp.level or 0
            run = p.add_run()
            if rpr is not None:
                run._r.insert(0, deepcopy(rpr))
            run.text = spec.text
            run.font.size = Pt(spec.size_pt)
            if not slot_font and spec.font:
                run.font.name = spec.font  # what we measured is what PowerPoint renders
            if spec.bold:
                run.font.bold = True
            if bp.italic:
                run.font.italic = True
            if bp.color:
                run.font.color.rgb = rgb(bp.color)
        return ph

    def tighten(self, shape) -> None:
        """Shrink a text box to the height its text needs, keeping its vertical anchor."""
        from pptx.enum.text import MSO_ANCHOR

        tf = shape.text_frame
        paras = [
            Para(
                p.text,
                (p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else 18),
                bold=bool(p.runs and p.runs[0].font.bold),
                font=p.runs[0].font.name if p.runs else None,
            )
            for p in tf.paragraphs
        ]
        width = int(shape.width) - int(tf.margin_left) - int(tf.margin_right)
        needed = int(measure(paras, width).height_emu * 1.1) + int(tf.margin_top) + int(tf.margin_bottom)
        if needed >= int(shape.height):
            return
        anchor = tf.vertical_anchor
        top, height = int(shape.top), int(shape.height)
        if anchor == MSO_ANCHOR.BOTTOM:
            top = top + height - needed
        elif anchor == MSO_ANCHOR.MIDDLE:
            top = top + (height - needed) // 2
        set_geometry(shape, y=top, h=needed)

    def separate(self, upper, lower, attached=None) -> None:
        """Move ``lower`` (and a frame ``attached`` to it) down when ``upper`` now overlaps it (titles can grow)."""
        top = min(int(lower.top), int(attached.top)) if attached is not None else int(lower.top)
        bottom = int(upper.top) + int(upper.height)
        overlaps = top < bottom and int(lower.top) + int(lower.height) > int(upper.top)
        if overlaps and int(lower.top) >= int(upper.top):
            shift = bottom + inch(0.08) - top
            set_geometry(lower, y=int(lower.top) + shift)
            if attached is not None:
                set_geometry(attached, y=int(attached.top) + shift)

    def notes(self, text: str | None) -> None:
        if text:
            self.slide.notes_slide.notes_text_frame.text = text


def set_geometry(shape, *, x=None, y=None, w=None, h=None) -> None:
    """Change position/size; all four are written because setting one on a placeholder that inherits its
    geometry would otherwise zero the others."""
    cur = (int(shape.left), int(shape.top), int(shape.width), int(shape.height))
    shape.left, shape.top = cur[0] if x is None else int(x), cur[1] if y is None else int(y)
    shape.width, shape.height = cur[2] if w is None else int(w), cur[3] if h is None else int(h)


_SZ_PATHS = (
    "./p:txBody/a:lstStyle/a:lvl1pPr/a:defRPr/@sz",
    "./p:txBody/a:p/a:pPr/a:defRPr/@sz",
    "./p:txBody/a:p/a:r/a:rPr/@sz",
    "./p:txBody/a:p/a:endParaRPr/@sz",
)


def inherited_size(ph) -> float | None:
    """Font size (pt) a placeholder inherits from its layout, master placeholder or master text styles."""
    chain = []
    layout_ph = getattr(ph, "_base_placeholder", None)
    if layout_ph is not None:
        chain.append(layout_ph._element)
        master_ph = getattr(layout_ph, "_base_placeholder", None)
        if master_ph is not None:
            chain.append(master_ph._element)
    for el in chain:
        for path in _SZ_PATHS:
            found = el.xpath(path)
            if found:
                return int(found[0]) / 100
    try:
        master = ph.part.slide_layout.slide_master._element
        style = "p:titleStyle" if ph.placeholder_format.idx == 0 else "p:bodyStyle"
        found = master.xpath(f"./p:txStyles/{style}/a:lvl1pPr/a:defRPr/@sz")
        return int(found[0]) / 100 if found else None
    except (AttributeError, ValueError):
        return None


def _set_bullet(paragraph, level: int, color: str, indent_emu: int) -> None:
    """Real paragraph bullets (never a literal bullet character in the text)."""
    p_pr = paragraph._p.get_or_add_pPr()
    p_pr.set("marL", str(indent_emu * (level + 1)))
    p_pr.set("indent", str(-indent_emu))
    p_pr.set("lvl", str(level))
    for tag in ("a:buClr", "a:buFont", "a:buChar", "a:buNone", "a:buAutoNum"):
        for old in p_pr.findall(qn(tag)):
            p_pr.remove(old)
    bu_clr = etree.SubElement(p_pr, qn("a:buClr"))
    srgb = etree.SubElement(bu_clr, qn("a:srgbClr"))
    srgb.set("val", color.upper())
    bu_font = etree.SubElement(p_pr, qn("a:buFont"))
    bu_font.set("typeface", "Arial")
    bu_char = etree.SubElement(p_pr, qn("a:buChar"))
    bu_char.set("char", "•" if level == 0 else "–")
