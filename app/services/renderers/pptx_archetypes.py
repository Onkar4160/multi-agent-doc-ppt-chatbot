"""Native, editable visual components for slide archetypes (FR-PPTX-4, ADR-029).

Charts are real PowerPoint charts (data editable in PowerPoint), tables are real tables, and timelines, process
flows, KPI tiles, comparisons, quotes and callouts are shapes with live text. Components draw into a content
region given in inches and take their colours and fonts from a ``Palette`` derived from the template (or the
generated design), so they sit inside any template's design.
"""

from __future__ import annotations

from dataclasses import dataclass
from typing import Any

from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Pt

from app.services.renderers.pptx_canvas import BulletPara, Canvas, Rect, TextStyle, content_height, inch, rgb
from app.services.layout.text import Para

VISUAL_TYPES = {"chart", "table", "kpi", "timeline", "process", "comparison", "quote", "callout"}
MIN_PT = 14.0

CHART_TYPES = {
    "bar": XL_CHART_TYPE.BAR_CLUSTERED, "column": XL_CHART_TYPE.COLUMN_CLUSTERED, "line": XL_CHART_TYPE.LINE_MARKERS,
    "pie": XL_CHART_TYPE.PIE, "doughnut": XL_CHART_TYPE.DOUGHNUT, "area": XL_CHART_TYPE.AREA,
    "scatter": XL_CHART_TYPE.LINE_MARKERS, "stacked_bar": XL_CHART_TYPE.BAR_STACKED,
    "stacked_column": XL_CHART_TYPE.COLUMN_STACKED,
}


@dataclass(frozen=True)
class Palette:
    accent: str = "1F4E79"
    accent2: str = "2E86AB"
    text: str = "1F2933"
    muted: str = "5B6770"
    surface: str = "EEF2F6"
    background: str = "FFFFFF"
    on_accent: str = "FFFFFF"
    heading_font: str | None = None
    body_font: str | None = None


def _hex(v: str | None) -> str | None:
    v = (v or "").lstrip("#").upper()
    return v if len(v) == 6 and all(c in "0123456789ABCDEF" for c in v) else None


def _lum(h: str) -> float:
    r, g, b = (int(h[i:i + 2], 16) / 255 for i in (0, 2, 4))
    return 0.2126 * r + 0.7152 * g + 0.0722 * b


def palette_from_profile(profile: Any) -> Palette:
    """Template colours and fonts (design tokens, Stage 6) as a component palette; defaults where missing."""
    t = getattr(profile, "design_tokens", None) if profile is not None else None
    if t is None:
        return Palette()
    accent = _hex(t.primary) or Palette.accent
    accents = [a for a in (_hex(x) for x in (t.accents or [])) if a and a != accent]
    second = _hex(t.secondary) or (accents[0] if accents else Palette.accent2)
    text = _hex(t.text) or Palette.text
    bg = _hex(t.background) or "FFFFFF"
    if _lum(bg) < 0.4:  # dark template: components keep light text on the slide
        text, surface, muted = "FFFFFF", "33414D", "C9D1D9"
    else:
        surface, muted = Palette.surface, Palette.muted
    return Palette(accent=accent, accent2=second, text=text, muted=muted, surface=surface, background=bg,
                   on_accent="FFFFFF" if _lum(accent) < 0.6 else "1F2933",
                   heading_font=t.heading_font, body_font=t.body_font)


def with_overrides(p: Palette, style: Any) -> Palette:
    """Apply a user's restyle (accent colour, fonts) to a palette."""
    from dataclasses import replace

    if style is None:
        return p
    return replace(p, accent=style.accent_color or p.accent, heading_font=style.heading_font or p.heading_font,
                   body_font=style.body_font or p.body_font)


def has_visual(blocks: list[Any]) -> bool:
    return any(getattr(b, "type", "") in VISUAL_TYPES for b in blocks or [])


_SCALE = [1.0]  # set per draw_blocks call: sizes are defined for a 7.5-inch-tall slide


def _para(text: str, size: float, p: Palette, *, bold: bool = False, heading: bool = False) -> Para:
    return Para(text=text, size_pt=size * _SCALE[0], bold=bold, font=(p.heading_font if heading else p.body_font), space_after_pt=4)


def _cite(block: Any) -> str:
    ids = sorted(set(getattr(block, "source_ids", []) or []))
    return " " + "".join(f"[{i}]" for i in ids) if ids else ""


def _alt(shape: Any, text: str) -> None:
    shape._element.xpath("./*[1]/p:cNvPr")[0].set("descr", text[:500])


def draw_blocks(canvas: Canvas, region: Rect, blocks: list[Any], palette: Palette) -> None:
    """Draw a slide's blocks into ``region``: the main visual first, a callout or bullets below it."""
    _SCALE[0] = canvas.scale
    visual = [b for b in blocks if getattr(b, "type", "") in VISUAL_TYPES - {"callout"}]
    extra = [b for b in blocks if getattr(b, "type", "") in ("callout", "bullets")]
    area = region
    if visual and extra:
        area, below = region.take_top(region.h * 0.78, 0.12)
        _side_text(canvas, below, extra, palette)
    elif extra and not visual:
        _side_text(canvas, region, extra, palette)
        return
    if visual:
        DRAW[visual[0].type](canvas, area, visual[0], palette)


def _side_text(c: Canvas, r: Rect, blocks: list[Any], p: Palette) -> None:
    from app.services.copyfit import strip_citations

    paras = []
    for b in blocks:
        if b.type == "callout":
            paras.append(BulletPara(_para(b.text + _cite(b), 16, p, bold=True), color=p.accent))
        else:
            paras += [BulletPara(_para(strip_citations(i.text) + _cite(i), 15, p), level=0) for i in b.items]
    c.text(r, paras, TextStyle(p.text, anchor="middle", bullet_color=p.accent), name="content:note", min_pt=MIN_PT)


def _chart(c: Canvas, r: Rect, b: Any, p: Palette) -> None:
    data = CategoryChartData()
    data.categories = [str(x) for x in b.categories]
    for s in b.series:
        data.add_series(s.name, [float(v) for v in s.values[: len(b.categories)]])
    kind = CHART_TYPES.get(b.chart_type, XL_CHART_TYPE.COLUMN_CLUSTERED)
    frame = c.slide.shapes.add_chart(kind, inch(r.x), inch(r.y), inch(r.w), inch(r.h), data)
    frame.name = "content:chart"
    chart = frame.chart
    unit = f" ({b.unit})" if b.unit else ""
    chart.has_title = bool(b.title)
    if b.title:
        chart.chart_title.text_frame.text = b.title + unit
        tp = chart.chart_title.text_frame.paragraphs[0]
        tp.runs[0].font.size, tp.runs[0].font.bold = Pt(round(16 * _SCALE[0])), True
        tp.runs[0].font.color.rgb = rgb(p.text)
    pie = kind in (XL_CHART_TYPE.PIE, XL_CHART_TYPE.DOUGHNUT)
    chart.has_legend = pie or len(b.series) > 1
    if chart.has_legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(round(14 * _SCALE[0]))
        chart.legend.font.color.rgb = rgb(p.text)
    plot = chart.plots[0]
    plot.has_data_labels = True
    labels = plot.data_labels
    labels.font.size = Pt(round(14 * _SCALE[0]))
    labels.font.color.rgb = rgb(p.text)
    labels.number_format = '0%' if pie else 'General'
    labels.number_format_is_linked = not pie
    if pie:
        labels.show_percentage, labels.show_value = True, False
    colors = [p.accent, p.accent2, p.muted]
    if pie:
        for i, point in enumerate(plot.series[0].points):
            point.format.fill.solid()
            point.format.fill.fore_color.rgb = rgb(colors[i % len(colors)] if i < 3 else _shade(p.accent, i))
    else:
        for i, series in enumerate(plot.series):
            fmt = series.format
            if kind == XL_CHART_TYPE.LINE_MARKERS:
                fmt.line.color.rgb = rgb(colors[i % len(colors)])
                fmt.line.width = Pt(2.5)
            else:
                fmt.fill.solid()
                fmt.fill.fore_color.rgb = rgb(colors[i % len(colors)])
        for axis in (chart.category_axis, chart.value_axis):
            axis.tick_labels.font.size = Pt(round(14 * _SCALE[0]))
            axis.tick_labels.font.color.rgb = rgb(p.text)
            axis.format.line.color.rgb = rgb(p.muted)
        chart.value_axis.has_major_gridlines = True
        chart.value_axis.major_gridlines.format.line.color.rgb = rgb(p.surface)
    desc = ", ".join(f"{cat}: {v:g}" for cat, v in zip(b.categories, b.series[0].values)) if b.series else ""
    _alt(frame, f"{b.chart_type} chart: {b.title or ''}{unit}. {desc}")
    if b.source_ids:
        c.text(Rect(r.x, r.bottom_y - 0.4 * _SCALE[0], r.w, 0.4 * _SCALE[0]), [BulletPara(_para("Source:" + _cite(b), 11, p))],
               TextStyle(p.muted, align="right"), name="content:references-chart", min_pt=10)


def _shade(h: str, i: int) -> str:
    f = 0.55 + 0.1 * (i % 4)
    return "".join(f"{min(255, int(int(h[k:k + 2], 16) * f + 255 * (1 - f))):02X}" for k in (0, 2, 4))


def _table(c: Canvas, r: Rect, b: Any, p: Palette) -> None:
    headers = list(b.headers)
    rows = [list(map(str, row)) for row in b.rows][:12]
    cols = max(len(headers), max((len(x) for x in rows), default=0), 1)
    n_rows = len(rows) + (1 if headers else 0)
    row_h = min(0.6, r.h / max(n_rows, 1))
    frame = c.slide.shapes.add_table(n_rows, cols, inch(r.x), inch(r.y), inch(r.w), inch(row_h * n_rows))
    frame.name = "content:table"
    table = frame.table
    size = 14 if n_rows <= 6 else 12
    for ri in range(n_rows):
        values = headers if (headers and ri == 0) else rows[ri - (1 if headers else 0)]
        for ci in range(cols):
            cell = table.cell(ri, ci)
            cell.text = values[ci] if ci < len(values) else ""
            para = cell.text_frame.paragraphs[0]
            for run in para.runs:
                run.font.size = Pt(size)
                run.font.bold = bool(headers and ri == 0)
                if p.body_font:
                    run.font.name = p.body_font
                run.font.color.rgb = rgb(p.on_accent if (headers and ri == 0) else p.text)
            cell.fill.solid()
            cell.fill.fore_color.rgb = rgb(p.accent if (headers and ri == 0) else (p.surface if ri % 2 == 0 else p.background))
    _alt(frame, "Table: " + ", ".join(headers))


def _kpi(c: Canvas, r: Rect, b: Any, p: Palette) -> None:
    from app.services.copyfit import strip_citations

    items = b.items[:4]
    if len(items) == 1:
        _kpi_hero(c, r, items[0], p)
        return
    s = _SCALE[0]
    if r.w / len(items) < 2.4 * s:  # too narrow for tiles side by side: one row per figure
        for i, (item, row) in enumerate(zip(items, r.split_v(len(items), 0.15 * s))):
            c.rect(row, p.surface, f"decor:kpi-{i + 1}", rounded=True)
            value, label = row.inset(0.15 * s, 0.05 * s).take_left(row.w * 0.4, 0.15 * s)
            c.text(value, [BulletPara(_para(item.value, 30, p, bold=True, heading=True))],
                   TextStyle(p.accent, align="right", anchor="middle"), name=f"content:kpi-value-{i + 1}", min_pt=16)
            text = strip_citations(item.label) + (f" ({item.delta})" if item.delta else "") + _cite(item)
            c.text(label, [BulletPara(_para(text, 15, p))], TextStyle(p.text, anchor="middle"),
                   name=f"content:kpi-label-{i + 1}", min_pt=MIN_PT * 0.8)
        return
    tiles = r.split_h(len(items), 0.3) if len(items) > 1 else [r.inset(r.w * 0.25, 0)]
    for i, (item, t) in enumerate(zip(items, tiles)):
        tile = Rect(t.x, r.y + r.h * 0.15, t.w, r.h * 0.7)
        c.rect(tile, p.surface, f"decor:kpi-{i + 1}", rounded=True)
        c.rect(Rect(tile.x, tile.y, tile.w, 0.08), p.accent, f"decor:kpi-bar-{i + 1}")
        inner = tile.inset(0.2, 0.25)
        value, label = inner.take_top(inner.h * 0.5)
        c.text(value, [BulletPara(_para(item.value, 40, p, bold=True, heading=True))],
               TextStyle(p.accent, align="center", anchor="bottom"), name=f"content:kpi-value-{i + 1}", min_pt=20)
        text = strip_citations(item.label) + (f" ({item.delta})" if item.delta else "") + _cite(item)
        c.text(label, [BulletPara(_para(text, 15, p))], TextStyle(p.text, align="center"),
               name=f"content:kpi-label-{i + 1}", min_pt=MIN_PT)


def _kpi_hero(c: Canvas, r: Rect, item: Any, p: Palette) -> None:
    """One headline figure: the number large in the accent colour, an accent rule, its meaning beside it."""
    from app.services.copyfit import strip_citations

    s = _SCALE[0]
    if r.w < 6.0 * s:  # a narrow content column: the figure above its meaning
        h = min(r.h, 3.4 * s)
        band = Rect(r.x, r.y + (r.h - h) / 2, r.w, h)
        top, below = band.take_top(band.h * 0.5, 0.15 * s)
        c.text(top, [BulletPara(_para(item.value, 60, p, bold=True, heading=True))],
               TextStyle(p.accent, align="center", anchor="bottom"), name="content:kpi-value-1", min_pt=28)
        c.rect(Rect(band.x + band.w * 0.4, below.y - 0.1 * s, band.w * 0.2, 0.05 * s), p.accent, "decor:kpi-rule")
        text = strip_citations(item.label) + (f" ({item.delta})" if item.delta else "") + _cite(item)
        c.text(Rect(below.x, below.y + 0.1 * s, below.w, below.h - 0.1 * s), [BulletPara(_para(text, 20, p))],
               TextStyle(p.text, align="center"), name="content:kpi-label-1", min_pt=MIN_PT)
        return
    h = min(r.h, 3.2 * s)
    band = Rect(r.x, r.y + (r.h - h) / 2, r.w, h)
    left, right = band.take_left(band.w * 0.45, 0.5 * s)
    c.text(left, [BulletPara(_para(item.value, 66, p, bold=True, heading=True))],
           TextStyle(p.accent, align="right", anchor="middle"), name="content:kpi-value-1", min_pt=28)
    c.rect(Rect(right.x - 0.3 * s, band.y + band.h * 0.2, 0.06 * s, band.h * 0.6), p.accent, "decor:kpi-rule")
    text = strip_citations(item.label) + (f" ({item.delta})" if item.delta else "") + _cite(item)
    c.text(Rect(right.x, band.y, right.w, band.h), [BulletPara(_para(text, 22, p))],
           TextStyle(p.text, anchor="middle"), name="content:kpi-label-1", min_pt=MIN_PT)


def _timeline(c: Canvas, r: Rect, b: Any, p: Palette) -> None:
    events = b.events[:7]
    n = max(len(events), 1)
    if r.w / n < 1.7 * _SCALE[0]:  # too narrow for a horizontal axis: a vertical timeline (date, then event)
        _timeline_vertical(c, r, events, p)
        return
    y = r.y + r.h * 0.45
    c.line(r.x, y, r.right_x, y, p.muted, "decor:timeline-axis", width_pt=2.5)
    slot = r.w / n
    for i, e in enumerate(events):
        cx = r.x + slot * (i + 0.5)
        c.circle(cx, y, 0.32, p.accent, f"decor:timeline-dot-{i + 1}")
        top = Rect(cx - slot / 2 + 0.05, r.y, slot - 0.1, y - r.y - 0.25)
        if e.date.strip():
            c.text(top, [BulletPara(_para(e.date, 18, p, bold=True, heading=True))], TextStyle(p.accent, align="center", anchor="bottom"),
                   name=f"content:timeline-date-{i + 1}", min_pt=MIN_PT)
        body = Rect(cx - slot / 2 + 0.05, y + 0.25, slot - 0.1, r.bottom_y - y - 0.25)
        paras = [BulletPara(_para(e.label, 15, p, bold=True))] + ([BulletPara(_para(e.detail, 14, p))] if e.detail else [])
        c.text(body, paras, TextStyle(p.text, align="center"), name=f"content:timeline-label-{i + 1}", min_pt=MIN_PT)


def _process(c: Canvas, r: Rect, b: Any, p: Palette) -> None:
    steps = b.steps[:6]
    n = max(len(steps), 1)
    s_ = _SCALE[0]
    if r.w / n < 2.0 * s_:  # too narrow for chevrons in a row: numbered steps from top to bottom
        for i, (st, row) in enumerate(zip(steps, r.split_v(n, 0.12 * s_))):
            d = min(row.h * 0.8, 0.6 * s_)
            c.badge(row.x + d / 2, row.y + row.h / 2, d, str(i + 1), p.accent if i % 2 == 0 else p.accent2, p.on_accent,
                    p.heading_font or "Calibri", f"decor:step-{i + 1}")
            text = st.label + (f": {st.detail}" if st.detail else "")
            c.text(Rect(row.x + d + 0.2 * s_, row.y, row.w - d - 0.2 * s_, row.h), [BulletPara(_para(text, 15, p))],
                   TextStyle(p.text, anchor="middle"), name=f"content:step-label-{i + 1}", min_pt=MIN_PT * 0.8)
        return
    gap = 0.15
    cells = r.split_h(n, gap)
    for i, (s, cell) in enumerate(zip(steps, cells)):
        head = Rect(cell.x, cell.y, cell.w, min(0.9, cell.h * 0.3))
        shp = c.shape(MSO_SHAPE.CHEVRON if i else MSO_SHAPE.PENTAGON, head, fill=p.accent if i % 2 == 0 else p.accent2,
                      name=f"decor:step-{i + 1}")
        shp.text_frame.text = ""
        c.text(head.inset(0.35, 0.05), [BulletPara(_para(f"{i + 1}. {s.label}", 15, p, bold=True))],
               TextStyle(p.on_accent, align="center", anchor="middle"), name=f"content:step-label-{i + 1}", min_pt=MIN_PT)
        if s.detail:
            body = Rect(cell.x, head.bottom_y + 0.15, cell.w, cell.bottom_y - head.bottom_y - 0.15)
            c.text(body, [BulletPara(_para(s.detail, 14, p))], TextStyle(p.text), name=f"content:step-detail-{i + 1}", min_pt=MIN_PT)


def _timeline_vertical(c: Canvas, r: Rect, events: list[Any], p: Palette) -> None:
    s = _SCALE[0]
    rows = r.split_v(max(len(events), 1), 0.12 * s)
    x_line = r.x + r.w * 0.3
    c.line(x_line, rows[0].y + 0.1 * s, x_line, rows[-1].bottom_y - 0.1 * s, p.muted, "decor:timeline-axis", width_pt=2)
    for i, (e, row) in enumerate(zip(events, rows)):
        c.circle(x_line, row.y + row.h / 2, 0.22 * s, p.accent, f"decor:timeline-dot-{i + 1}")
        c.text(Rect(row.x, row.y, r.w * 0.3 - 0.3 * s, row.h), [BulletPara(_para(e.date, 18, p, bold=True, heading=True))],
               TextStyle(p.accent, align="right", anchor="middle"), name=f"content:timeline-date-{i + 1}", min_pt=MIN_PT)
        text = e.label + (f": {e.detail}" if e.detail else "")
        c.text(Rect(x_line + 0.3 * s, row.y, r.right_x - x_line - 0.3 * s, row.h), [BulletPara(_para(text, 16, p))],
               TextStyle(p.text, anchor="middle"), name=f"content:timeline-label-{i + 1}", min_pt=MIN_PT)


def _comparison(c: Canvas, r: Rect, b: Any, p: Palette) -> None:
    cols = b.columns[:3]
    if r.w / max(len(cols), 1) < 2.6 * _SCALE[0]:  # narrow area: options stacked as rows instead of columns
        r_rows = r.split_v(len(cols), 0.2 * _SCALE[0])
        for i, (col, row) in enumerate(zip(cols, r_rows)):
            c.rect(row, p.surface, f"decor:column-{i + 1}", rounded=True)
            head, body = row.take_top(min(0.5 * _SCALE[0], row.h * 0.35))
            c.rect(head, p.accent if i == 0 else p.accent2, f"decor:column-head-{i + 1}", rounded=True)
            c.text(head.inset(0.15, 0.03), [BulletPara(_para(col.heading, 16, p, bold=True, heading=True))],
                   TextStyle(p.on_accent, anchor="middle"), name=f"content:column-head-{i + 1}", min_pt=MIN_PT * 0.8)
            c.text(body.inset(0.15, 0.08), [BulletPara(_para(pt, 14, p), level=0) for pt in col.points],
                   TextStyle(p.text, anchor="middle", bullet_color=p.accent), name=f"content:column-{i + 1}", min_pt=MIN_PT)
        return
    cells = r.split_h(len(cols), 0.3)
    need = max(content_height([BulletPara(_para(pt, 15, p), level=0) for pt in col.points], cell.w - 0.4)
               for col, cell in zip(cols, cells)) + 0.7 + 0.5
    height = min(r.h, max(need, r.h * 0.45))  # cards as tall as their content, not the whole region
    cells = [Rect(cell.x, r.y, cell.w, height) for cell in cells]
    for i, (col, cell) in enumerate(zip(cols, cells)):
        c.rect(cell, p.surface, f"decor:column-{i + 1}", rounded=True)
        head, body = cell.take_top(0.7)
        c.rect(head, p.accent if i == 0 else p.accent2, f"decor:column-head-{i + 1}", rounded=True)
        c.text(head.inset(0.15, 0.05), [BulletPara(_para(col.heading, 18, p, bold=True, heading=True))],
               TextStyle(p.on_accent, align="center", anchor="middle"), name=f"content:column-head-{i + 1}", min_pt=MIN_PT)
        paras = [BulletPara(_para(pt, 15, p), level=0) for pt in col.points]
        c.text(body.inset(0.2, 0.15), paras, TextStyle(p.text, bullet_color=p.accent), name=f"content:column-{i + 1}", min_pt=MIN_PT)


def _quote(c: Canvas, r: Rect, b: Any, p: Palette) -> None:
    c.rect(Rect(r.x, r.y + r.h * 0.15, 0.12, r.h * 0.7), p.accent, "decor:quote-bar")
    body = Rect(r.x + 0.4, r.y, r.w - 0.4, r.h)
    paras = [BulletPara(_para(f"“{b.text}”", 28, p, heading=True), italic=True)]
    if b.attribution or b.source_ids:
        paras.append(BulletPara(_para(f"— {b.attribution}{_cite(b)}".strip(), 16, p), color=p.muted))
    c.text(body, paras, TextStyle(p.text, anchor="middle"), name="content:quote", min_pt=14)


def _callout(c: Canvas, r: Rect, b: Any, p: Palette) -> None:
    box = Rect(r.x, r.y + r.h * 0.25, r.w, r.h * 0.5)
    c.rect(box, p.surface, "decor:callout", rounded=True)
    c.text(box.inset(0.4, 0.3), [BulletPara(_para(b.text + _cite(b), 24, p, bold=True, heading=True))],
           TextStyle(p.accent, align="center", anchor="middle"), name="content:callout", min_pt=14)


DRAW = {"chart": _chart, "table": _table, "kpi": _kpi, "timeline": _timeline, "process": _process,
        "comparison": _comparison, "quote": _quote, "callout": _callout}
